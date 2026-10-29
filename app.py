import os
import re
import io
import json
import time
import base64
import html
import uuid
import random
import hashlib
import zipfile
import pandas as pd
from typing import List, Dict

import streamlit as st
import streamlit.components.v1 as components
from google import genai
from google.genai import types

try:
    from supabase import create_client
except ImportError:
    create_client = None

# -----------------------------
# 1. 기본 설정, 비밀번호 잠금 및 API 보안 로드
# -----------------------------
st.set_page_config(page_title="4pm 올인원 교재 시스템", layout="wide")

def get_secret(key: str, default: str = "") -> str:
    try:
        if key in st.secrets:
            return str(st.secrets[key]).strip()
    except Exception:
        pass
    return os.environ.get(key, default).strip()

APP_PASSWORD = get_secret("APP_PASSWORD", "")
if APP_PASSWORD:
    if "authenticated" not in st.session_state:
        st.session_state.authenticated = False
    if not st.session_state.authenticated:
        st.markdown("## 🔒 4pm 올인원 교재 시스템 (전용 웹 스튜디오)")
        pwd_input = st.text_input("접속 비밀번호를 입력하세요", type="password")
        if st.button("로그인", type="primary"):
            if pwd_input == APP_PASSWORD:
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("❌ 비밀번호가 일치하지 않습니다.")
        st.stop()

LOCAL_FALLBACK_KEY = ""  # 로컬 PC 단독 실행 시에만 입력
API_KEY = get_secret("GEMINI_API_KEY", LOCAL_FALLBACK_KEY)

if not API_KEY or API_KEY == "여기에_선생님의_API_키를_넣어주세요":
    st.error("API 키를 찾지 못했습니다. Streamlit Cloud의 Secrets에 GEMINI_API_KEY를 등록해주세요.")
    st.stop()

@st.cache_resource
def get_genai_client(api_key: str):
    return genai.Client(api_key=api_key)

client = get_genai_client(API_KEY)

# Supabase 클라우드 DB 연결 초기화 (URL 뒤에 경로가 붙어도 순수 도메인만 자동 추출)
raw_supabase_url = get_secret("SUPABASE_URL", "https://evwqyfmnqiklxcbwaqee.supabase.co")
m_url = re.match(r"(https?://[^/\s]+)", raw_supabase_url)
SUPABASE_URL = m_url.group(1) if m_url else raw_supabase_url
SUPABASE_KEY = get_secret("SUPABASE_KEY", "")

@st.cache_resource
def get_supabase_client(url: str, key: str):
    if url and key and create_client is not None:
        try:
            return create_client(url, key)
        except Exception as e:
            st.warning(f"⚠️ Supabase 연결 실패 (로컬 모드로 전환됩니다): {e}")
    return None

supabase = get_supabase_client(SUPABASE_URL, SUPABASE_KEY)
IS_CLOUD_DB = supabase is not None

# -----------------------------
# 2. 공통 유틸리티 및 하이브리드(로컬+클라우드) 데이터 관리
# -----------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
VICTORY_FILE = os.path.join(BASE_DIR, "victory_library.json")
ANATOMY_FILE = os.path.join(BASE_DIR, "anatomy_library.json")
CORE_FILE = os.path.join(BASE_DIR, "core_library.json")
VOCA_FILE = os.path.join(BASE_DIR, "voca_library.json")
PAST_VOCA_FILE = os.path.join(BASE_DIR, "past_voca_library.json")
MOCK_FILE = os.path.join(BASE_DIR, "mock_exam_library.json")
DNA_FILE = os.path.join(BASE_DIR, "exam_dna_rules.json")
IMAGE_DIR = os.path.join(BASE_DIR, "image_assets")
os.makedirs(IMAGE_DIR, exist_ok=True)

def filepath_to_doc_key(filepath: str) -> str:
    return os.path.splitext(os.path.basename(filepath))[0]

DEFAULT_EXAM_DNA = {
    "updated_at": "2026-10 (운양고 1·2·3학년 7개 시험지 & 제일고 2학년 1개 시험지 기출 DNA 탑재)",
    "rules_text": """[운양고 & 제일고 실제 내신 출제 DNA 핵심 규칙]
1. 상위 개념어(Meta-Vocabulary) 치환 원칙:
   - 지문의 구체적 서술을 주제/제목/요지/요약문에서는 반드시 학술적 상위 개념어로 압축할 것.
   - 실제 기출 예시: 'help a child in distress' -> 'altruistic behavior(이타적 행동)', 'start off hating' -> 'aversion(혐오)', 'associated with a reward' -> 'paired with positive reinforcement(긍정적 강화)', 'depriving them of personal choices' -> 'autonomy(자율성)', 'keep its location a secret' -> 'withhold its whereabouts(소재지를 감추다)', 'defensive, rigid state' -> 'emotionally guarded and unyielding posture'.
2. 요약문 빈칸 (A), (B) 출제 공식:
   - 지문 원문 단어를 그대로 빈칸으로 뚫지 말고, 원문의 핵심 인과관계/대조관계를 압축하는 학술 동사·명사·형용사·부사로 패러프레이징하여 (A), (B) 타겟으로 설정할 것.
   - 실제 기출 예시: trivialize -> downplay, defend -> safeguard, break down -> deconstruct / divide, capture energy -> harness, copy verbatim -> blindly / mindlessly copy, chronic -> persistent / continual, confuse -> malfunction, shaped by context -> conditioned / integrate.
3. 문맥상 어휘 추론(밑줄 변형) 및 함정 반의어 패턴:
   - 흐름을 정반대로 뒤집는 방향성 반의어 쌍을 적극 제시할 것 (예: facilitation <-> inhibition, impede <-> guarantee, lift burdens <-> impose burdens, permit <-> prevent, corridor <-> barrier, integration <-> fragmentation).
   - 철자나 형태가 유사해 착시를 일으키는 함정 어휘(Confusable words)를 주의 깊게 다룰 것 (예: priceless[매우 귀중한] vs valueless[가치 없는], replace[대체하다] vs misplace[잘못 두다], elicit[이끌어내다] vs eliminate[제거하다], feasible vs unfeasible).
4. 영문 선지 및 <보기> 진위형(T/F) 문장 스타일:
   - 요지/주장/주제/제목 및 <보기> 일치·불일치 문항은 모두 수준 높은 영문으로 출제되며, 원문의 인과관계를 역전시키거나 'only/solely/completely/independent of' 같은 극단적 한정어를 넣어 오답(False) 진술문을 구성함.""",
    "history": [
        {
            "title": "[기본 탑재] 운양고(고1·2·3) 7부 & 제일고(고2) 1부 기출 통합 분석",
            "summary": "상위 개념어 패러프레이징, 요약문 (A)/(B) 학술 어휘 치환, 방향성 역전 반의어 및 형태 유사 함정 어휘(priceless vs valueless 등), <보기> 영문 일치/불일치 변형 패턴 반영 완료"
        }
    ]
}

def load_exam_dna() -> Dict:
    doc_key = filepath_to_doc_key(DNA_FILE)
    if IS_CLOUD_DB:
        try:
            res = supabase.table("app_storage").select("content").eq("doc_key", doc_key).execute()
            if res.data and res.data[0].get("content"):
                return res.data[0]["content"]
        except Exception:
            pass
    if os.path.exists(DNA_FILE):
        try:
            with open(DNA_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            if IS_CLOUD_DB:
                save_exam_dna(data)
            return data
        except Exception:
            pass
    save_exam_dna(DEFAULT_EXAM_DNA)
    return DEFAULT_EXAM_DNA

def save_exam_dna(dna_data: Dict):
    doc_key = filepath_to_doc_key(DNA_FILE)
    if IS_CLOUD_DB:
        try:
            supabase.table("app_storage").upsert({"doc_key": doc_key, "content": dna_data}, on_conflict="doc_key").execute()
        except Exception as e:
            st.warning(f"⚠️ 클라우드 출제 DNA 저장 실패: {e}")
    try:
        with open(DNA_FILE, "w", encoding="utf-8") as f:
            json.dump(dna_data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

# 사이드바 로고·학원명 등 설정을 클라우드에 저장해 로그인할 때마다 다시 올리지 않도록 함
SETTINGS_FILE = os.path.join(BASE_DIR, "app_settings.json")
DEFAULT_LOGO_URL = "https://www.genspark.ai/api/files/s/Bvmm1owU"

def load_app_settings() -> Dict:
    doc_key = filepath_to_doc_key(SETTINGS_FILE)
    if IS_CLOUD_DB:
        try:
            res = supabase.table("app_storage").select("content").eq("doc_key", doc_key).execute()
            if res.data and isinstance(res.data[0].get("content"), dict):
                return res.data[0]["content"]
        except Exception:
            pass
    if os.path.exists(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def save_app_settings(settings: Dict):
    doc_key = filepath_to_doc_key(SETTINGS_FILE)
    if IS_CLOUD_DB:
        try:
            supabase.table("app_storage").upsert({"doc_key": doc_key, "content": settings}, on_conflict="doc_key").execute()
        except Exception as e:
            st.warning(f"⚠️ 설정 클라우드 저장 실패: {e}")
    try:
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

def logo_upload_to_data_url(uploaded, max_side: int = 400) -> str:
    """로고는 교재 머리말에 작게 쓰이므로 400px PNG로 줄여서 저장 (투명 배경 유지)"""
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(uploaded.getvalue()))
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, format="PNG", optimize=True)
        return f"data:image/png;base64,{base64.b64encode(buf.getvalue()).decode()}"
    except Exception:
        return logo_to_url(uploaded)

def clean_label(text) -> str:
    if not text:
        return ""
    text = str(text)
    text = re.sub(r'<[^<]+?>', '', text)
    if "color:#" in text or "font-size" in text:
        text = text.split(">")[-1].split("}")[-1].strip()
    return text.strip()

def natural_sort_key(item):
    lbl = item.get("label", "") if isinstance(item, dict) else str(item)
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'([0-9]+)', lbl)]

def extract_group_prefix(label: str) -> str:
    lbl = clean_label(label)
    if lbl.startswith("[") and "]" in lbl:
        in_bracket = lbl[1:lbl.find("]")].strip()
        parts = in_bracket.split()
        return parts[0] if parts else in_bracket
    prefix = re.sub(r'[\s_\-#]*(?:q|no\.?|문항|번)?\s*\d+\s*(?:번|문항)?$', '', lbl, flags=re.IGNORECASE).strip()
    return prefix if prefix else "기타"

def extract_question_number(label: str):
    nums = re.findall(r'\d+', clean_label(label))
    return int(nums[-1]) if nums else None

# -----------------------------
# 2-A. 기출 어휘 시험지 단위 자동 통합기 (기존 분할 데이터 + 신규 업로드 모두 시험지별로 통합)
# -----------------------------
def parse_past_exam_label(label: str) -> (str, str):
    """'[운양고 26년 1학기 중간 고1] 18번' -> ('[운양고 26년 1학기 중간 고1]', '18번') 분리"""
    lbl = clean_label(label)
    m = re.match(r'^(\[[^\]]+\])\s*(.*)$', lbl)
    if m:
        exam_base = m.group(1).strip()
        rem = re.sub(r'\s*기출\s*어휘.*$', '', m.group(2).strip()).strip()
        return exam_base, rem
    m2 = re.match(r'^(.*?)\s*((?:논술형|서술형|문항|Q|q)?\s*\d+\s*(?:번|문항)?(?:\s*~\s*\d+\s*번?)?)$', lbl)
    if m2 and m2.group(1).strip():
        return f"[{m2.group(1).strip()}]", m2.group(2).strip()
    return (lbl if lbl.startswith("[") else f"[{lbl}]"), ""

def extract_q_order_from_category(cat_str: str) -> (int, int):
    """단어의 출처 배지에서 문항 번호를 추출하여 1번 -> 2번 -> ... -> 논술형 순으로 정렬"""
    s = str(cat_str)
    m_sub = re.search(r'(?:논술형|서술형)\s*(\d+)', s)
    if m_sub:
        return (2, int(m_sub.group(1)))
    m_num = re.search(r'(\d+)\s*(?:번|문항)', s)
    if m_num:
        return (1, int(m_num.group(1)))
    return (3, 999)

def consolidate_past_voca_records(records: List[Dict]) -> (List[Dict], bool):
    """문항별로 쪼개진 기출 어휘 레코드들을 시험지 단위(예: [운양고 26년 1학기 중간 고1])로 자동 병합"""
    if not records:
        return [], False

    grouped = {}
    changed = False

    for p in records:
        orig_lbl = clean_label(p.get("label", ""))
        exam_base, q_str = parse_past_exam_label(orig_lbl)
        clean_exam_name = exam_base.strip("[]").strip()

        if orig_lbl != exam_base or exam_base in grouped:
            changed = True

        if exam_base not in grouped:
            grouped[exam_base] = {
                "id": p.get("id") or str(uuid.uuid4()),
                "label": exam_base,
                "data": {"voca_list": []},
                "body_text": f"{exam_base} 기출 통합 어휘",
                "academy_name": p.get("academy_name", "포피엠 입시 영어"),
                "school_name": p.get("school_name", clean_exam_name.split()[0] if clean_exam_name else "")
            }

        target_list = grouped[exam_base]["data"]["voca_list"]
        existing_keys = {
            (str(x.get("word", "")).strip().lower(), str(x.get("category", "")).strip())
            for x in target_list
        }

        for w in p.get("data", {}).get("voca_list", []):
            w_copy = dict(w)
            cat = str(w_copy.get("category", "")).strip()
            # 만약 단어 category에 문항 번호가 빠져 있다면 라벨에서 추출한 문항 번호를 자동 기입
            if q_str and (q_str not in cat):
                src_tag = cat if cat else "기출어휘"
                cat = f"{clean_exam_name} {q_str} [{src_tag}]"
                w_copy["category"] = cat
                changed = True

            dedup_key = (str(w_copy.get("word", "")).strip().lower(), cat)
            if dedup_key not in existing_keys:
                target_list.append(w_copy)
                existing_keys.add(dedup_key)

    consolidated_list = list(grouped.values())
    for item in consolidated_list:
        v_list = item["data"]["voca_list"]
        sorted_v = sorted(v_list, key=lambda x: extract_q_order_from_category(x.get("category", "")))
        if sorted_v != v_list:
            item["data"]["voca_list"] = sorted_v
            changed = True

    consolidated_list.sort(key=natural_sort_key)
    if len(consolidated_list) != len(records):
        changed = True

    return consolidated_list, changed

def save_base64_to_image_file(data_url: str, custom_filename: str = "") -> str:
    if not data_url or "," not in data_url:
        return ""
    try:
        _, b64_str = data_url.split(",", 1)
        img_bytes = base64.b64decode(b64_str)
        if custom_filename:
            filename = custom_filename if custom_filename.endswith(".png") else f"{custom_filename}.png"
        else:
            img_hash = hashlib.md5(img_bytes).hexdigest()[:14]
            filename = f"{img_hash}.png"
        img_path = os.path.join(IMAGE_DIR, filename)
        if not os.path.exists(img_path):
            with open(img_path, "wb") as img_f:
                img_f.write(img_bytes)
        if IS_CLOUD_DB:
            try:
                supabase.storage.from_("image_assets").upload(
                    path=filename,
                    file=img_bytes,
                    file_options={"content-type": "image/png", "upsert": "true"}
                )
            except Exception:
                try:
                    supabase.table("app_images").upsert({
                        "img_id": filename,
                        "b64_data": f"data:image/png;base64,{b64_str}"
                    }, on_conflict="img_id").execute()
                except Exception:
                    pass
        return img_path
    except Exception:
        return ""

def normalize_passage_records(data: List[Dict]) -> (List[Dict], bool):
    needs_save = False
    for p in data:
        if not p.get("id"):
            p["id"] = str(uuid.uuid4())
            needs_save = True
        cleaned = clean_label(p.get("label", ""))
        if p.get("label") != cleaned:
            p["label"] = cleaned
            needs_save = True
        if p.get("image_path"):
            base_fn = os.path.basename(str(p["image_path"]).replace("\\", "/"))
            norm_path = os.path.join(IMAGE_DIR, base_fn)
            if p["image_path"] != norm_path:
                p["image_path"] = norm_path
                needs_save = True
        if "image_data_url" in p and len(p.get("image_data_url", "")) > 1000:
            file_id = p.get("id", str(uuid.uuid4()))
            img_path = save_base64_to_image_file(p["image_data_url"], custom_filename=f"{file_id}.png")
            if img_path:
                p["image_path"] = img_path
                p["image_data_url"] = ""
                needs_save = True
    data.sort(key=natural_sort_key)
    return data, needs_save

def load_from_file(filepath: str) -> List[Dict]:
    doc_key = filepath_to_doc_key(filepath)
    data = None

    if IS_CLOUD_DB:
        try:
            res = supabase.table("app_storage").select("content").eq("doc_key", doc_key).execute()
            if res.data and isinstance(res.data[0].get("content"), list):
                data = res.data[0]["content"]
        except Exception as e:
            st.warning(f"⚠️ 클라우드 데이터 로드 중 오류 ({doc_key}): {e}")

    if data is None:
        if not os.path.exists(filepath):
            return []
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            st.error(f"⚠️ 파일 로드 중 에러 발생 ({os.path.basename(filepath)}): {e}")
            return []

    data, needs_save = normalize_passage_records(data)
    if filepath == PAST_VOCA_FILE:
        data, c_changed = consolidate_past_voca_records(data)
        if c_changed:
            needs_save = True

    if needs_save:
        save_to_file(filepath, data)
    return data

def save_to_file(filepath: str, passages: List[Dict]):
    doc_key = filepath_to_doc_key(filepath)
    passages.sort(key=natural_sort_key)
    if IS_CLOUD_DB:
        try:
            supabase.table("app_storage").upsert({
                "doc_key": doc_key,
                "content": passages
            }, on_conflict="doc_key").execute()
        except Exception as e:
            st.warning(f"⚠️ 클라우드 자동 저장 실패 ({doc_key}): {e}")
    try:
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(passages, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

# 통합 세션 스테이트 초기화
if "saved_v" not in st.session_state:
    st.session_state.saved_v = load_from_file(VICTORY_FILE)
if "v_selected" not in st.session_state:
    st.session_state.v_selected = []

if "saved_a" not in st.session_state:
    st.session_state.saved_a = load_from_file(ANATOMY_FILE)
if "a_selected" not in st.session_state:
    st.session_state.a_selected = []

if "saved_c" not in st.session_state:
    st.session_state.saved_c = load_from_file(CORE_FILE)
if "c_selected" not in st.session_state:
    st.session_state.c_selected = []

if "saved_w" not in st.session_state:
    st.session_state.saved_w = load_from_file(VOCA_FILE)
if "w_selected" not in st.session_state:
    st.session_state.w_selected = []

if "saved_past_voca" not in st.session_state:
    st.session_state.saved_past_voca = load_from_file(PAST_VOCA_FILE)
else:
    # 이미 세션에 로드된 기존 문항별 쪼개진 데이터도 즉시 시험지 단위로 자동 병합
    st.session_state.saved_past_voca, _migrated = consolidate_past_voca_records(st.session_state.saved_past_voca)
    if _migrated:
        save_to_file(PAST_VOCA_FILE, st.session_state.saved_past_voca)

if "past_v_selected" not in st.session_state:
    st.session_state.past_v_selected = []

if "saved_mock" not in st.session_state:
    st.session_state.saved_mock = load_from_file(MOCK_FILE)

if "app_settings" not in st.session_state:
    st.session_state.app_settings = load_app_settings()

if "exam_dna" not in st.session_state:
    st.session_state.exam_dna = load_exam_dna()

@st.cache_data(show_spinner=False, max_entries=30)
def cached_read_image_base64(img_filename: str) -> str:
    local_p = os.path.join(IMAGE_DIR, img_filename)
    if os.path.exists(local_p):
        mime = "image/jpeg" if img_filename.lower().endswith((".jpg", ".jpeg")) else "image/png"
        with open(local_p, "rb") as f:
            return f"data:{mime};base64,{base64.b64encode(f.read()).decode()}"
    if IS_CLOUD_DB:
        return f"{SUPABASE_URL}/storage/v1/object/public/image_assets/{img_filename}"
    return ""

def get_image_base64(p: Dict) -> str:
    path = p.get("image_path", "")
    if path:
        img_filename = os.path.basename(str(path).replace("\\", "/"))
        res = cached_read_image_base64(img_filename)
        if res:
            return res
    return p.get("image_data_url", "")

def logo_to_url(uploaded) -> str:
    if uploaded is None:
        return "https://www.genspark.ai/api/files/s/Bvmm1owU"
    mime = uploaded.type if hasattr(uploaded, 'type') and uploaded.type else 'image/png'
    return f"data:{mime};base64,{base64.b64encode(uploaded.getvalue()).decode()}"

def esc(t) -> str:
    return html.escape(str(t)) if t else ""

def nl2br(t) -> str:
    return esc(t).replace("\n", "<br>")

def md_safe(t) -> str:
    """st.markdown에서 '41~42번'의 ~가 취소선으로 바뀌지 않도록 이스케이프"""
    return str(t).replace("~", "\\~")


# -----------------------------
# 2-B. 단행본 표지(1P) + 공란(2P) + 저작권 뒤표지(마지막P) 빌더
# -----------------------------
BOOK_COVER_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;600;800;900&family=Noto+Sans+KR:wght@300;400;500;700;900&display=swap');
.book-special-page {
    width: 210mm; height: 297mm; box-sizing: border-box; margin: 10mm auto;
    background: #ffffff; box-shadow: 0 0 15px rgba(0,0,0,0.10);
    page-break-after: always; break-after: page; position: relative; overflow: hidden;
    font-family: 'Inter', 'Noto Sans KR', sans-serif;
}
.book-cover-inner {
    width: 100%; height: 100%; box-sizing: border-box; padding: 26mm 24mm 22mm 24mm;
    display: flex; flex-direction: column; justify-content: space-between; position: relative;
}
.cover-accent-bar {
    position: absolute; top: 0; left: 0; width: 100%; height: 10mm;
    background: linear-gradient(90deg, #0f172a 0%, #0047AB 65%, #38bdf8 100%);
}
.cover-side-line {
    position: absolute; top: 26mm; bottom: 22mm; left: 16mm; width: 1.5px; background: #e2e8f0;
}
.cover-top-meta {
    display: flex; justify-content: space-between; align-items: center;
    border-bottom: 1.5px solid #0f172a; padding-bottom: 14px;
}
.cover-series-tag {
    font-size: 11px; font-weight: 800; letter-spacing: 3.5px; color: #0047AB; text-transform: uppercase;
}
.cover-logo-circle {
    width: 46px; height: 46px; border-radius: 50%; border: 2px solid #0047AB;
    background-size: cover; background-position: center;
}
.cover-hero { margin: auto 0; padding: 15mm 0; }
.cover-edition-pill {
    display: inline-block; background: #0f172a; color: #38bdf8; font-size: 10.5px;
    font-weight: 800; letter-spacing: 2.5px; padding: 5px 12px; border-radius: 4px;
    margin-bottom: 18px; text-transform: uppercase;
}
.cover-main-title {
    font-size: 54px; font-weight: 900; color: #0f172a; letter-spacing: -1.5px;
    line-height: 1.08; margin: 0 0 14px 0;
}
.cover-main-title span { color: #0047AB; }
.cover-subtitle-en { font-size: 15px; font-weight: 600; color: #475569; letter-spacing: 0.5px; margin-bottom: 8px; }
.cover-subtitle-ko { font-size: 13.5px; font-weight: 500; color: #64748b; }
.cover-divider { width: 65px; height: 5px; background: #0047AB; margin: 28px 0; }
.cover-author-credit { font-size: 14px; font-weight: 700; color: #1e293b; letter-spacing: 1.2px; text-transform: uppercase; }
.cover-author-credit span { color: #0047AB; font-weight: 900; }
.cover-bottom-area {
    display: flex; justify-content: space-between; align-items: flex-end;
    border-top: 1px solid #cbd5e1; padding-top: 18px;
}
.cover-Creator-block { font-size: 11.5px; color: #475569; line-height: 1.6; }
.cover-student-box {
    width: 210px; border: 1.5px solid #0f172a; border-radius: 6px;
    padding: 10px 14px; font-size: 11.5px; color: #0f172a; background: #f8fafc;
}
.cover-student-row {
    display: flex; justify-content: space-between; border-bottom: 1px dashed #cbd5e1;
    padding: 5px 0; font-weight: 700;
}
.cover-student-row:last-child { border-bottom: none; }
.back-cover-inner {
    width: 100%; height: 100%; box-sizing: border-box; padding: 25mm 24mm 22mm 24mm;
    display: flex; flex-direction: column; justify-content: flex-end; align-items: flex-end;
}
.copyright-box-right { text-align: right; max-width: 125mm; border-right: 4px solid #0047AB; padding-right: 14px; }
.copyright-title-en { font-size: 12px; font-weight: 900; color: #0f172a; letter-spacing: 1.5px; margin-bottom: 6px; }
.copyright-ko-legal { font-size: 10.5px; font-weight: 700; color: #334155; line-height: 1.65; word-break: keep-all; margin-bottom: 6px; }
.copyright-en-sub { font-size: 9px; color: #94a3b8; letter-spacing: 0.4px; line-height: 1.4; }
@media print {
    .book-special-page {
        width: 210mm !important; height: 296mm !important; margin: 0 !important;
        box-shadow: none !important; border: none !important;
        page-break-after: always !important; break-after: page !important;
    }
    .book-last-page { page-break-after: auto !important; break-after: auto !important; }
}
"""

def build_booklet_front_and_back(book_type: str, h_img: str, f_img: str, academy: str, school: str, passage_count: int) -> (str, str):
    meta_map = {
        "victory": {
            "main_first": "Anatomy",
            "main_Highlight": "One-Shot",
            "sub_en": "Intensive Syntax Analysis & Direct Comprehension System",
            "sub_ko": "시험 직전 최후의 승자 — 문장별 구조 분석과 직독직해 한판승"
        },
        "anatomy": {
            "main_first": "Anatomy",
            "main_Highlight": "Self-Note",
            "sub_en": "Comprehensive Text Dissection, Grammar & Structure Workbook",
            "sub_ko": "내신과 수능 1등급을 위한 지문 해부 & 자기주도 필기·복습 노트"
        },
        "core": {
            "main_first": "Anatomy",
            "main_Highlight": "Core",
            "sub_en": "Essential Paraphrasing, Topics, Main Ideas & Summary Master",
            "sub_ko": "학교 내신 변형 1순위 — 다중 주제·제목·요지·요약문 패러프레이징"
        },
        "voca": {
            "main_first": "Anatomy",
            "main_Highlight": "Voca",
            "sub_en": "Contextual Vocabulary, Idioms, Synonyms & Antonyms Master",
            "sub_ko": "문맥 중심 중고등 필수 어휘 및 패러프레이징 유의어·반의어 마스터"
        },
        "past_voca": {
            "main_first": "Anatomy",
            "main_Highlight": "Voca",
            "sub_en": "Actual Past Exam Paper Vocabulary & Choices Master",
            "sub_ko": "학교별 실제 기출 시험지 본문·선지 출제 어휘 완벽 마스터"
        }
    }
    cfg = meta_map.get(book_type, meta_map["victory"])
    acad_esc = esc(academy)
    sch_esc = esc(school)

    front_pages_html = f"""
    <div class="book-special-page">
        <div class="cover-accent-bar"></div>
        <div class="cover-side-line"></div>
        <div class="book-cover-inner">
            <div class="cover-top-meta">
                <div class="cover-series-tag">PQRS YOUNG ENGLISH CURRICULUM SERIES</div>
                <div class="cover-logo-circle" style="background-image: url('{h_img}');"></div>
            </div>
            <div class="cover-hero">
                <div class="cover-edition-pill">PREMIUM EDITION &bull; {passage_count} SECTIONS INCLUDED</div>
                <h1 class="cover-main-title">{cfg['main_first']} <span>{cfg['main_Highlight']}</span></h1>
                <div class="cover-subtitle-en">{cfg['sub_en']}</div>
                <div class="cover-subtitle-ko">{cfg['sub_ko']}</div>
                <div class="cover-divider"></div>
                <div class="cover-author-credit">Created &amp; Designed by <span>PQRS Young</span></div>
            </div>
            <div class="cover-bottom-area">
                <div class="cover-Creator-block">
                    <div style="font-weight:900; color:#0f172a; font-size:13px; margin-bottom:2px;">PQRS Young &times; {acad_esc}</div>
                    <div>High School English Exam &amp; CSAT Preparation</div>
                    <div style="font-size:10.5px; color:#94a3b8;">Original Pedagogical System by PQRS Young</div>
                </div>
                <div class="cover-student-box">
                    <div class="cover-student-row"><span>SCHOOL</span><span style="color:#64748b; font-weight:500;">{sch_esc if sch_esc else ''}</span></div>
                    <div class="cover-student-row"><span>GRADE / CLASS</span><span></span></div>
                    <div class="cover-student-row"><span>STUDENT NAME</span><span></span></div>
                </div>
            </div>
        </div>
    </div>
    <div class="book-special-page"></div>
    """

    back_page_html = f"""
    <div class="book-special-page book-last-page">
        <div class="back-cover-inner">
            <div class="copyright-box-right">
                <div class="copyright-title-en">{cfg['main_first'].upper()} {cfg['main_Highlight'].upper()} &mdash; CREATED BY PQRS YOUNG</div>
                <div class="copyright-ko-legal">
                    본 교재의 모든 저작권은 제작자인 <b>PQRS영쌤</b>에게 있으며, 저작권자의 사전 서면 동의 없는 무단 복제, 배포, 전송, 스캔 및 도용 시 저작권법에 의거하여 민·형사상의 엄중한 법적 처벌을 받을 수 있습니다.
                </div>
                <div class="copyright-en-sub">
                    Copyright &copy; PQRS Young ({acad_esc}). All rights reserved.<br>
                    Unauthorized reproduction, distribution, or plagiarism of this material is strictly prohibited and subject to legal prosecution.
                </div>
            </div>
        </div>
    </div>
    """
    return front_pages_html, back_page_html


# -----------------------------
# 2-C. CORE 어휘 <-> VOCA 단어장 병합 유틸리티
# -----------------------------
def merge_core_vocab_into_voca_item(core_data: Dict, voca_data: Dict) -> int:
    if not core_data or not voca_data:
        return 0
    voca_list = voca_data.get("voca_list", [])
    existing_words = {str(item.get("word", "")).strip().lower() for item in voca_list}
    added_count = 0
    for cv in core_data.get("core_vocab", []):
        cw = str(cv.get("word", "")).strip()
        if not cw:
            continue
        if cw.lower() in existing_words:
            for item in voca_list:
                if str(item.get("word", "")).strip().lower() == cw.lower():
                    item["category"] = "★CORE 패러프레이징"
            continue
        voca_list.append({
            "word": cw,
            "pos": cv.get("pos", "n/v/a"),
            "pronunciation": cv.get("pronunciation", ""),
            "meaning": cv.get("meaning", ""),
            "category": "★CORE 패러프레이징",
            "synonyms": cv.get("synonyms", []),
            "antonyms": cv.get("antonyms", [])
        })
        existing_words.add(cw.lower())
        added_count += 1
    voca_data["voca_list"] = voca_list
    return added_count


# -----------------------------
# 2-D. 엑셀(.xlsx) 및 분리형 CSV 2종(표제어 통합 / 유의어·반의어 전용) 생성기
# -----------------------------
def generate_voca_main_csv_bytes(selected_passages: List[Dict]) -> bytes:
    """파일 1: 각 지문의 지문핵심 + ★CORE 패러프레이징(또는 기출 표제어)만 합친 CSV"""
    rows = []
    word_no = 1
    for p in selected_passages:
        lbl = p.get("label", "")
        voca_list = p.get("data", {}).get("voca_list", [])
        for item in voca_list:
            w = str(item.get("word", "")).strip()
            if not w:
                continue
            pos = str(item.get("pos", "")).strip()
            pron = str(item.get("pronunciation", "")).strip()
            mean = str(item.get("meaning", "")).strip()
            cat = str(item.get("category", "")).strip()
            cls_label = cat if ("본문" in cat or "선지" in cat or "번" in cat) else lbl
            full_mean = f"[{pos}] {mean}" if pos else mean

            rows.append({
                "분류": cls_label,
                "단어번호": word_no,
                "영어단어": w,
                "발음기호": str(item.get("ipa", "")).strip(),
                "뜻": full_mean,
                "한글발음": pron,
                "예문": str(item.get("example", "")).strip()
            })
            word_no += 1

    columns = ["분류", "단어번호", "영어단어", "발음기호", "뜻", "한글발음", "예문"]
    df = pd.DataFrame(rows, columns=columns) if rows else pd.DataFrame(columns=columns)
    return df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")


def generate_voca_rel_csv_bytes(selected_passages: List[Dict]) -> bytes:
    """파일 2: 각 지문의 유의어(Synonyms)와 반의어(Antonyms)만 따로 모아놓은 CSV"""
    rows = []
    word_no = 1
    for p in selected_passages:
        lbl = p.get("label", "")
        voca_list = p.get("data", {}).get("voca_list", [])
        for item in voca_list:
            w = str(item.get("word", "")).strip()
            if not w:
                continue
            cat = str(item.get("category", "")).strip()
            cls_label = cat if ("본문" in cat or "선지" in cat or "번" in cat) else lbl

            for syn in item.get("synonyms", []):
                sw = str(syn.get("word", "")).strip()
                if not sw:
                    continue
                rows.append({
                    "분류": f"{cls_label} [유의어]",
                    "단어번호": word_no,
                    "영어단어": sw,
                    "발음기호": "",
                    "뜻": f"[유의어: {w}] {str(syn.get('meaning', '')).strip()}",
                    "한글발음": str(syn.get("pronunciation", "")).strip(),
                    "예문": ""
                })
                word_no += 1

            for ant in item.get("antonyms", []):
                aw = str(ant.get("word", "")).strip()
                if not aw:
                    continue
                rows.append({
                    "분류": f"{cls_label} [반의어]",
                    "단어번호": word_no,
                    "영어단어": aw,
                    "발음기호": "",
                    "뜻": f"[반의어: {w}] {str(ant.get('meaning', '')).strip()}",
                    "한글발음": str(ant.get("pronunciation", "")).strip(),
                    "예문": ""
                })
                word_no += 1

    columns = ["분류", "단어번호", "영어단어", "발음기호", "뜻", "한글발음", "예문"]
    df = pd.DataFrame(rows, columns=columns) if rows else pd.DataFrame(columns=columns)
    return df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")


def generate_voca_excel_bytes(selected_passages: List[Dict]) -> bytes:
    rows_3col = []
    rows_detail = []
    rows_test = []
    row_types_3col = []

    for p in selected_passages:
        lbl = p.get("label", "")
        voca_list = p.get("data", {}).get("voca_list", [])
        if not voca_list:
            continue

        rows_3col.append({
            "영어 단어": f"■ {lbl} 핵심 어휘",
            "한글 발음": "",
            "한국어 뜻": ""
        })
        row_types_3col.append("passage_banner")

        for item in voca_list:
            w = str(item.get("word", "")).strip()
            pos = str(item.get("pos", "")).strip()
            pron = str(item.get("pronunciation", "")).strip()
            mean = str(item.get("meaning", "")).strip()
            cat = str(item.get("category", "지문핵심")).strip()
            is_highlight = any(k in cat.upper() for k in ["CORE", "패러", "요약", "주제", "선지"])

            tag_prefix = f"[{cat}] " if cat else ""
            full_mean = f"{tag_prefix}[{pos}] {mean}" if pos else f"{tag_prefix}{mean}"

            rows_3col.append({
                "영어 단어": w,
                "한글 발음": pron,
                "한국어 뜻": full_mean
            })
            row_types_3col.append("main_core" if is_highlight else "main")

            rows_detail.append({
                "시험지/지문": lbl,
                "상세 출처 (문항·위치)": cat,
                "구분": "★ 표제어",
                "품사": pos,
                "영어 단어": w,
                "한글 발음": pron,
                "한국어 뜻": mean
            })
            rows_test.append({
                "시험지/지문": lbl,
                "구분": f"★ 표제어({cat})" if cat else "★ 표제어",
                "영어 단어": w,
                "한글 발음": pron,
                "한국어 뜻 (시험빈칸)": ""
            })

            for syn in item.get("synonyms", []):
                sw = str(syn.get("word", "")).strip()
                spron = str(syn.get("pronunciation", "")).strip()
                smean = str(syn.get("meaning", "")).strip()
                if not sw:
                    continue
                rows_3col.append({
                    "영어 단어": f"  └ [유의어] {sw}",
                    "한글 발음": spron,
                    "한국어 뜻": smean
                })
                row_types_3col.append("syn")

                rows_detail.append({
                    "시험지/지문": lbl,
                    "상세 출처 (문항·위치)": cat,
                    "구분": "└ 유의어",
                    "품사": pos,
                    "영어 단어": sw,
                    "한글 발음": spron,
                    "한국어 뜻": smean
                })
                rows_test.append({
                    "시험지/지문": lbl,
                    "구분": "└ 유의어",
                    "영어 단어": sw,
                    "한글 발음": spron,
                    "한국어 뜻 (시험빈칸)": ""
                })

            for ant in item.get("antonyms", []):
                aw = str(ant.get("word", "")).strip()
                apron = str(ant.get("pronunciation", "")).strip()
                amean = str(ant.get("meaning", "")).strip()
                if not aw:
                    continue
                rows_3col.append({
                    "영어 단어": f"  └ [반의어/함정] {aw}",
                    "한글 발음": apron,
                    "한국어 뜻": amean
                })
                row_types_3col.append("ant")

                rows_detail.append({
                    "시험지/지문": lbl,
                    "상세 출처 (문항·위치)": cat,
                    "구분": "└ 반의어/함정",
                    "품사": pos,
                    "영어 단어": aw,
                    "한글 발음": apron,
                    "한국어 뜻": amean
                })
                rows_test.append({
                    "시험지/지문": lbl,
                    "구분": "└ 반의어/함정",
                    "영어 단어": aw,
                    "한글 발음": apron,
                    "한국어 뜻 (시험빈칸)": ""
                })

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_3col = pd.DataFrame(rows_3col) if rows_3col else pd.DataFrame(columns=["영어 단어", "한글 발음", "한국어 뜻"])
        df_detail = pd.DataFrame(rows_detail) if rows_detail else pd.DataFrame(columns=["시험지/지문", "상세 출처 (문항·위치)", "구분", "품사", "영어 단어", "한글 발음", "한국어 뜻"])
        df_test = pd.DataFrame(rows_test) if rows_test else pd.DataFrame(columns=["시험지/지문", "구분", "영어 단어", "한글 발음", "한국어 뜻 (시험빈칸)"])

        df_3col.to_excel(writer, sheet_name="3열_단어_발음_뜻", index=False)
        df_detail.to_excel(writer, sheet_name="상세_출처_트랙포함", index=False)
        df_test.to_excel(writer, sheet_name="학생용_단어테스트지", index=False)

        try:
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

            wb = writer.book
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            header_fill = PatternFill(start_color="0047AB", end_color="0047AB", fill_type="solid")
            header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
            banner_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            banner_font = Font(name="맑은 고딕", size=11, bold=True, color="38BDF8")
            main_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            main_font = Font(name="맑은 고딕", size=11, bold=True, color="0F172A")
            main_core_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            main_core_font = Font(name="맑은 고딕", size=11, bold=True, color="92400E")
            syn_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
            syn_font = Font(name="맑은 고딕", size=10, bold=False, color="1D4ED8")
            ant_fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
            ant_font = Font(name="맑은 고딕", size=10, bold=False, color="BE123C")

            ws1 = wb["3열_단어_발음_뜻"]
            ws1.column_dimensions["A"].width = 36
            ws1.column_dimensions["B"].width = 24
            ws1.column_dimensions["C"].width = 48
            for cell in ws1[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for idx, r_type in enumerate(row_types_3col, start=2):
                row_cells = ws1[idx]
                for c in row_cells:
                    c.border = thin_border
                    c.alignment = Alignment(vertical="center")
                if r_type == "passage_banner":
                    ws1.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=3)
                    row_cells[0].fill = banner_fill
                    row_cells[0].font = banner_font
                elif r_type == "main_core":
                    for c in row_cells:
                        c.fill = main_core_fill
                        c.font = main_core_font
                elif r_type == "main":
                    for c in row_cells:
                        c.fill = main_fill
                        c.font = main_font
                elif r_type == "syn":
                    for c in row_cells:
                        c.fill = syn_fill
                        c.font = syn_font
                elif r_type == "ant":
                    for c in row_cells:
                        c.fill = ant_fill
                        c.font = ant_font

            for sname, widths in [
                ("상세_출처_트랙포함", {"A": 26, "B": 30, "C": 14, "D": 10, "E": 28, "F": 22, "G": 40}),
                ("학생용_단어테스트지", {"A": 26, "B": 28, "C": 28, "D": 22, "E": 40}),
            ]:
                ws = wb[sname]
                for col_letter, w_val in widths.items():
                    ws.column_dimensions[col_letter].width = w_val
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                for row in ws.iter_rows(min_row=2):
                    row_str = " ".join([str(c.value or "") for c in row[:3]])
                    for c in row:
                        c.border = thin_border
                        c.alignment = Alignment(vertical="center")
                        if "유의어" in row_str:
                            c.fill, c.font = syn_fill, syn_font
                        elif "반의어" in row_str:
                            c.fill, c.font = ant_fill, ant_font
                        elif "CORE" in row_str or "선지" in row_str:
                            c.fill, c.font = main_core_fill, main_core_font
                        else:
                            c.fill, c.font = main_fill, main_font
        except Exception:
            pass

    return output.getvalue()


# -----------------------------
# 2-E. HTML 스마트 스튜디오 컨트롤 바
# -----------------------------
def get_studio_toolbar_html(mode_type: str = "standard") -> str:
    extra_btns = ""
    if mode_type in ["voca", "past_voca"]:
        extra_btns = """
        <button class="studio-btn" style="background:#0284c7; border-color:#38bdf8;" onclick="setVocaTestMode('study')">📖 학습용 표준</button>
        <button class="studio-btn" style="background:#0d9488; border-color:#2dd4bf;" onclick="setVocaTestMode('hide-meaning')">📝 뜻 시험지 전환</button>
        <button class="studio-btn" style="background:#7c3aed; border-color:#a78bfa;" onclick="setVocaTestMode('hide-word')">✍️ 영어 시험지 전환</button>
        """
    elif mode_type == "core":
        extra_btns = """
        <button class="studio-btn" style="background:#0284c7; border-color:#38bdf8;" onclick="setCoreTestMode('study')">📖 해설·정답 표시 모드</button>
        <button class="studio-btn" style="background:#0d9488; border-color:#2dd4bf;" onclick="setCoreTestMode('test')">📝 요약문(A)(B) 양자택일 & T/F 실전 시험지 모드</button>
        """

    return f"""
<style>
.studio-toolbar {{
    position: sticky; top: 0; z-index: 99999; background: #0f172a; color: white;
    padding: 10px 20px; display: flex; flex-wrap: wrap; justify-content: space-between;
    align-items: center; gap: 10px; box-shadow: 0 4px 12px rgba(0,0,0,0.25); font-family: 'Noto Sans KR', sans-serif;
}}
.studio-toolbar .st-title {{ font-weight: 900; font-size: 14px; color: #38bdf8; display: flex; align-items: center; gap: 8px; }}
.studio-toolbar .st-actions {{ display: flex; flex-wrap: wrap; gap: 6px; align-items: center; }}
.studio-btn {{
    background: #1e293b; color: #f8fafc; border: 1px solid #475569; padding: 6px 11px;
    border-radius: 6px; font-size: 12px; font-weight: 700; cursor: pointer; transition: all 0.15s; font-family: inherit;
}}
.studio-btn:hover {{ background: #334155; border-color: #94a3b8; }}
.studio-btn.primary {{ background: #2563eb; border-color: #3b82f6; }}
.studio-btn.active-edit {{ background: #dc2626; border-color: #f87171; }}
.studio-drawer {{
    display: none; background: #1e293b; color: #f1f5f9; padding: 12px 20px;
    border-bottom: 2px solid #334155; max-height: 240px; overflow-y: auto; font-family: 'Noto Sans KR', sans-serif;
}}
.studio-drawer.open {{ display: block; }}
.drawer-item {{
    display: flex; justify-content: space-between; align-items: center; padding: 6px 10px;
    margin-bottom: 4px; background: #0f172a; border-radius: 6px; font-size: 12px; border: 1px solid #334155;
}}
.drawer-item-btns button {{
    background: #334155; color: white; border: none; border-radius: 4px;
    padding: 3px 8px; margin-left: 4px; cursor: pointer; font-size: 11px;
}}
[contenteditable="true"] {{ outline: 1.5px dashed #ef4444 !important; background-color: rgba(254, 226, 226, 0.12) !important; }}
@media print {{
    .no-print, .studio-toolbar, .studio-drawer, .voca-del-btn {{ display: none !important; }}
    [contenteditable="true"] {{ outline: none !important; background-color: transparent !important; }}
}}
</style>
<nav class="studio-toolbar no-print">
    <div class="st-title"><span>⚡ 4pm 실시간 교재 스튜디오 (표지·공란·저작권 페이지 자동 탑재)</span></div>
    <div class="st-actions">
        {extra_btns}
        <button class="studio-btn" onclick="toggleOrderDrawer()">📋 지문/시험지 순서 변경 / 제외</button>
        <button id="btn-edit-mode" class="studio-btn" onclick="toggleEditMode()">✏️ 화면 텍스트 직접 수정: OFF</button>
        <button class="studio-btn" onclick="downloadEditedHtml()">💾 단행본 전체 HTML 저장</button>
        <button class="studio-btn primary" onclick="window.print()">🖨️ 즉시 인쇄 / PDF 저장</button>
    </div>
</nav>
<div id="studio-order-drawer" class="studio-drawer no-print">
    <div style="font-size:12px; font-weight:700; margin-bottom:8px; color:#38bdf8;">
        💡 화살표를 눌러 본문 순서를 실시간으로 바꾸거나 특정 항목을 제외할 수 있습니다. (앞표지·공란·뒤표지는 항상 고정됩니다)
    </div>
    <div id="drawer-list"></div>
</div>
<script>
let isEditMode = false;
function toggleEditMode() {{
    isEditMode = !isEditMode;
    const container = document.getElementById('book-wrapper-all');
    const btn = document.getElementById('btn-edit-mode');
    if (container) container.contentEditable = isEditMode ? "true" : "false";
    if (isEditMode) {{
        btn.innerText = "✏️ 화면 텍스트 직접 수정: ON (클릭해 수정)";
        btn.classList.add("active-edit");
    }} else {{
        btn.innerText = "✏️ 화면 텍스트 직접 수정: OFF";
        btn.classList.remove("active-edit");
        if (typeof fitAllPages === "function") setTimeout(fitAllPages, 100);
    }}
}}
function setVocaTestMode(mode) {{
    const container = document.getElementById('book-container');
    if (!container) return;
    container.classList.remove('mode-hide-meaning', 'mode-hide-word');
    if (mode === 'hide-meaning') container.classList.add('mode-hide-meaning');
    if (mode === 'hide-word') container.classList.add('mode-hide-word');
}}
function setCoreTestMode(mode) {{
    const container = document.getElementById('book-container');
    if (!container) return;
    if (mode === 'test') container.classList.add('core-test-mode');
    else container.classList.remove('core-test-mode');
}}
function removeVocaRow(btnEl) {{
    const tr = btnEl.closest('tr');
    if (tr) {{
        const tbody = tr.parentElement;
        tr.remove();
        if (tbody) {{
            Array.from(tbody.querySelectorAll('.voca-idx-num')).forEach((el, idx) => {{ el.innerText = (idx + 1); }});
        }}
    }}
}}
function getPassageGroups() {{
    const container = document.getElementById('book-container');
    if (!container) return [];
    const children = Array.from(container.children);
    const groups = [];
    const map = new Map();
    children.forEach(el => {{
        const pid = el.getAttribute('data-passage-id') || 'unknown';
        const plabel = el.getAttribute('data-passage-label') || pid;
        if (!map.has(pid)) {{
            const g = {{ id: pid, label: plabel, elements: [] }};
            map.set(pid, g);
            groups.push(g);
        }}
        map.get(pid).elements.push(el);
    }});
    return groups;
}}
function renderDrawerList() {{
    const listEl = document.getElementById('drawer-list');
    if (!listEl) return;
    const groups = getPassageGroups();
    listEl.innerHTML = groups.map((g, idx) => `
        <div class="drawer-item">
            <span><b>${{idx + 1}}.</b> ${{g.label}}</span>
            <div class="drawer-item-btns">
                <button onclick="movePassageGroup(${{idx}}, -1)" ${{idx === 0 ? 'disabled style="opacity:0.3"' : ''}}>⬆️ 위로</button>
                <button onclick="movePassageGroup(${{idx}}, 1)" ${{idx === groups.length - 1 ? 'disabled style="opacity:0.3"' : ''}}>⬇️ 아래로</button>
                <button onclick="removePassageGroup('${{g.id}}')" style="background:#991b1b;">❌ 제외</button>
            </div>
        </div>
    `).join('');
}}
function toggleOrderDrawer() {{
    const drawer = document.getElementById('studio-order-drawer');
    drawer.classList.toggle('open');
    if (drawer.classList.contains('open')) renderDrawerList();
}}
function movePassageGroup(idx, dir) {{
    const groups = getPassageGroups();
    const targetIdx = idx + dir;
    if (targetIdx < 0 || targetIdx >= groups.length) return;
    const temp = groups[idx];
    groups[idx] = groups[targetIdx];
    groups[targetIdx] = temp;
    const container = document.getElementById('book-container');
    groups.forEach(g => g.elements.forEach(el => container.appendChild(el)));
    renderDrawerList();
}}
function removePassageGroup(pid) {{
    const groups = getPassageGroups();
    const target = groups.find(g => g.id === pid);
    if (target) {{
        target.elements.forEach(el => el.remove());
        renderDrawerList();
    }}
}}
function downloadEditedHtml() {{
    if (isEditMode) toggleEditMode();
    const docClone = document.documentElement.cloneNode(true);
    const blob = new Blob(["<!DOCTYPE html>\\n" + docClone.outerHTML], {{type: "text/html;charset=utf-8"}});
    const a = document.createElement("a");
    a.href = URL.createObjectURL(blob);
    a.download = document.title ? document.title.replace(/\\s+/g, '_') + "_Booklet.html" : "Anatomy_Booklet.html";
    a.click();
}}
</script>
"""


# -----------------------------
# 3-A. [한판승: Anatomy One-Shot] 서식 및 렌더링
# -----------------------------
def build_css_victory(h_img, f_img, e_img):
    return f"""<style>
    @import url('https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@300;400;700;900&display=swap');
    @page {{ size: A4; margin: 12mm 15mm; }} 
    body {{ font-family: 'Noto Sans KR', sans-serif; background: #f0f2f5; margin: 0; padding: 0; -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }}
    {BOOK_COVER_CSS}
    .passage-wrapper {{ width: 210mm; box-sizing: border-box; padding: 15mm 18mm; margin: 10mm auto; background: white; box-shadow: 0 0 10px rgba(0,0,0,0.1); page-break-after: always; }}
    @media print {{ body {{ background: white; }} .passage-wrapper {{ width: 100%; margin: 0; padding: 0; box-shadow: none; border: none; }} }}
    .master-table {{ width: 100%; border-collapse: collapse; }}
    .master-table thead {{ display: table-header-group; }}
    .master-table tfoot {{ display: table-footer-group; }}
    .master-table > tbody > tr > td {{ padding: 0; }}
    .header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 3px solid #0047AB; padding-bottom: 8px; margin-bottom: 12px; }}
    .logo-area {{ display: flex; align-items: center; gap: 10px; font-weight: 900; font-size: 22px; color: #0047AB; }}
    .h-img {{ width: 40px; height: 40px; background: url('{h_img}') center/cover; border-radius: 50%; border: 2px solid #0047AB; }}
    .footer {{ text-align: center; font-size: 10px; color: #999; border-top: 1px solid #eee; padding-top: 8px; margin-top: 15px; display: flex; justify-content: center; align-items: center; gap: 8px; }}
    .f-img {{ width: 22px; height: 22px; background: url('{f_img}') center/cover; border-radius: 50%; }}
    .top-layout {{ display: flex; gap: 15px; margin-bottom: 12px; }}
    .visual-box {{ flex: 0 0 220px; height: 160px; border-radius: 10px; background: #f8f9fa; display: flex; align-items: center; justify-content: center; overflow: hidden; border: 1px solid #eee; }}
    .visual-box img {{ max-width: 100%; max-height: 100%; object-fit: contain; }}
    .summary-box {{ flex: 1; background: #0047AB; color: white; padding: 15px; border-radius: 10px; display: flex; align-items: center; font-weight: 700; font-size: 14.5px; line-height: 1.55; }}
    .keyword-section {{ background: #fff4e5; border: 1px solid #f5c27a; border-radius: 10px; padding: 8px 12px; margin-bottom: 12px; }}
    .keyword-title {{ color: #9a5a00; font-weight: 900; font-size: 12px; margin-bottom: 5px; }}
    .keyword-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 5px 10px; }}
    .keyword-item {{ font-size: 10.5px; color: #333; }}
    .keyword-word {{ font-weight: 800; color: #0047AB; }}
    .analysis-table {{ width: 100%; border-collapse: collapse; }} 
    .analysis-table tr {{ page-break-inside: avoid; break-inside: avoid; }} 
    .analysis-table td {{ padding-bottom: 15px; }} 
    .eng-box {{ padding: 10px 14px 15px 14px; font-weight: 700; color: #1e293b; background: #fdfdfd; border-top: 1px solid #dee2e6; border-left: 1px solid #dee2e6; border-right: 1px solid #dee2e6; border-top-left-radius: 8px; border-top-right-radius: 8px; border-bottom: 1px dashed #e2e8f0; font-size: 13px; line-height: 2.3; word-break: keep-all; }}
    ruby.syntax-ruby {{ ruby-position: under; ruby-align: center; margin: 0 2px; }}
    ruby.syntax-ruby rt {{ font-weight: 900; font-size: 0.7em; line-height: 1; margin-top: 2px; }}
    .st-S span.syntax-text {{ border-bottom: 2px solid #e11d48; padding-bottom: 1px; }} .st-S rt {{ color: #e11d48; }}
    .st-V span.syntax-text {{ border-bottom: 2px solid #2563eb; padding-bottom: 1px; }} .st-V rt {{ color: #2563eb; }}
    .st-O span.syntax-text {{ border-bottom: 2px solid #16a34a; padding-bottom: 1px; }} .st-O rt {{ color: #16a34a; }}
    .st-C span.syntax-text {{ border-bottom: 2px solid #ea580c; padding-bottom: 1px; }} .st-C rt {{ color: #ea580c; }}
    .st-M span.syntax-text {{ border-bottom: 2px dashed #9333ea; padding-bottom: 1px; }} .st-M rt {{ color: #9333ea; }}
    .inline-vocab-box {{ padding: 10px 14px; background: #fafafa; border-left: 1px solid #dee2e6; border-right: 1px solid #dee2e6; border-bottom: 1px dashed #e2e8f0; font-size: 11.5px; line-height: 1.5; color: #444; }}
    .trans-box {{ padding: 10px 14px; color: #475569; background: #ffffff; border-left: 1px solid #dee2e6; border-right: 1px solid #dee2e6; border-bottom: 1px dashed #e2e8f0; font-size: 12.5px; line-height: 1.6; font-style: normal; word-break: keep-all; }}
    .inline-grammar-box {{ padding: 10px 14px; background: #fefce8; color: #854d0e; border-left: 1px solid #dee2e6; border-right: 1px solid #dee2e6; border-bottom: 1px dashed #e2e8f0; font-size: 12px; line-height: 1.5; }}
    .ys-box {{ padding: 10px 14px; background: #eef4ff; color: #0047AB; font-weight: 500; border-left: 4px solid #FF8C00 !important; border-right: 1px solid #dee2e6; border-bottom: 1px solid #dee2e6; border-bottom-left-radius: 8px; border-bottom-right-radius: 8px; font-size: 12.5px; line-height: 1.6; word-break: keep-all; }}
    .ys-header {{ display: flex; align-items: center; gap: 6px; margin-bottom: 6px; font-weight: 900; color: #FF8C00; font-size: 11px; }}
    .ys-img {{ width: 22px; height: 22px; background: url('{e_img}') center/cover; border-radius: 50%; border: 1px solid #FF8C00; }}
    </style>"""

def highlight_syntax(text):
    text = esc(text)
    pattern = r'\[(.*?)\]\s*\(([SVOCM])\)'
    replacement = r'<ruby class="syntax-ruby st-\2"><span class="syntax-text">\1</span><rt>(\2)</rt></ruby>'
    return re.sub(pattern, replacement, text, flags=re.DOTALL)

def render_victory(p):
    data = p["data"]
    analysis_list = data.get("analysis", [])
    img_url = get_image_base64(p)
    img_html = f'<img src="{img_url}">' if img_url else '<div style="color:#ccc;">No Image</div>'
    kw_html = "".join([
        f'<div class="keyword-item"><span class="keyword-word">{esc(kw.get("word",""))}</span> ({esc(kw.get("meaning",""))}) : {esc(kw.get("synonyms",""))}</div>'
        for kw in data.get("keywords", [])[:8]
    ])
    top_section = f"""
    <div class="top-layout">
        <div class="visual-box">{img_html}</div>
        <div class="summary-box">🎯 한판승 핵심 요약:<br>{esc(data.get('summary', ''))}</div>
    </div>
    <div class="keyword-section">
        <div class="keyword-title">📍 반드시 알아야 할 핵심 키워드</div>
        <div class="keyword-grid">{kw_html}</div>
    </div>
    """
    rows_html = ""
    for row in analysis_list:
        eng_highlighted = highlight_syntax(row.get('eng', ''))
        vocab_list = row.get("sentence_vocab", [])
        vocab_html = ""
        if vocab_list:
            v_items = [
                f"<span style='white-space: nowrap;'>• <b>{esc(v.get('word', ''))}</b> ({esc(v.get('pronunciation', ''))}) [{esc(v.get('pos', ''))}] {esc(v.get('meaning', ''))}</span>"
                for v in vocab_list
            ]
            vocab_html = f"""<div class="inline-vocab-box">
                <div style="font-weight: 700; margin-bottom: 5px; color: #333;">📖 Sentence Vocab</div>
                <div style="display: flex; flex-wrap: wrap; gap: 4px 15px;">{"".join(v_items)}</div>
            </div>"""
        grammar_html = ""
        if row.get("grammar"):
            grammar_html = f'<div class="inline-grammar-box"><b>💡 Grammar Point</b> : {esc(row.get("grammar"))}</div>'

        rows_html += f"""
        <tr>
            <td>
                <div class="eng-box"><span style="color:#0047AB; font-weight:900; margin-right:5px; vertical-align: top;">E.</span>{eng_highlighted}</div>
                {vocab_html}
                <div class="trans-box"><span style="color:#64748b; font-weight:900; margin-right:5px;">K.</span>{esc(row.get('kor', ''))}</div>
                {grammar_html}
                <div class="ys-box"><div class="ys-header"><div class="ys-img"></div>영쌤의 한판승 코칭</div>{esc(row.get('ys', ''))}</div>
            </td>
        </tr>"""
    label_text = esc(p['label'])
    pid = esc(p.get('id', label_text))
    academy = esc(p.get('academy_name', '4pm 입시영어'))

    return f"""
    <div class="passage-wrapper" data-passage-id="{pid}" data-passage-label="{label_text}">
        <table class="master-table">
            <thead>
                <tr><td>
                    <div class="header">
                        <div class="logo-area"><div class="h-img"></div>ANATOMY 한판승</div>
                        <div style="text-align:right; font-size:11px;"><b>{academy}</b><br>{label_text}</div>
                    </div>
                </td></tr>
            </thead>
            <tfoot>
                <tr><td>
                    <div class="footer">
                        <div class="f-img"></div>{academy} | 시험 직전 최후의 승자, 한판승
                    </div>
                </td></tr>
            </tfoot>
            <tbody>
                <tr><td>
                    {top_section}
                    <table class="analysis-table">
                        <tbody>{rows_html}</tbody>
                    </table>
                </td></tr>
            </tbody>
        </table>
    </div>"""

def wrap_full_html_victory(body_html, css, front_html="", back_html=""):
    return f"""<!DOCTYPE html><html lang="ko"><head><meta charset="UTF-8"><title>Anatomy_One_Shot</title>{css}</head><body>{get_studio_toolbar_html('standard')}<div id="book-wrapper-all">{front_html}<main id="book-container">{body_html}</main>{back_html}</div></body></html>"""


# -----------------------------
# 3-B. [ANATOMY: Anatomy Self-Note] 서식 및 렌더링
# -----------------------------
def build_css_anatomy(h_logo, f_logo, e_logo):
    return f"""<style>
    @import url('https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@300;400;700;900&display=swap');
    @page {{ size: A4; margin: 0; }}
    body {{ font-family: 'Noto Sans KR', sans-serif; margin: 0; padding: 0; background: #f4f7fa; color: #333; -webkit-print-color-adjust: exact; print-color-adjust: exact; counter-reset: pagenum; }}
    {BOOK_COVER_CSS}
    #book-container {{ counter-reset: pagenum; }}
    .page {{ --scale: 1; width: 210mm; height: 297mm; padding: calc(15mm * var(--scale)) calc(18mm * var(--scale)) calc(12mm * var(--scale)) calc(18mm * var(--scale)); margin: 10mm auto; background: white; box-shadow: 0 0 15px rgba(0,0,0,0.10); box-sizing: border-box; overflow: hidden; display: flex; flex-direction: column; page-break-after: always; position: relative; counter-increment: pagenum; }}
    .header {{ display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 4px solid #0047AB; padding-bottom: calc(12px * var(--scale)); margin-bottom: calc(18px * var(--scale)); flex: 0 0 auto; }}
    .logo-area {{ color: #0047AB; font-weight: 900; font-size: calc(30px * var(--scale)); display: flex; align-items: center; gap: calc(12px * var(--scale)); }}
    .header-char {{ width: calc(50px * var(--scale)); height: calc(50px * var(--scale)); background: url('{h_logo}') center/cover; border-radius: 50%; border: 2px solid #0047AB; flex-shrink: 0; }}
    .info-area {{ text-align: right; font-size: calc(14px * var(--scale)); color: #555; line-height: 1.5; }}
    .page-content {{ flex: 1 1 auto; min-height: 0; overflow: hidden; display: flex; flex-direction: column; gap: calc(8px * var(--scale)); }}
    .footer {{ flex: 0 0 auto; margin-top: calc(8px * var(--scale)); text-align: center; font-size: calc(11px * var(--scale)); color: #999; border-top: 1px solid #eee; padding-top: calc(8px * var(--scale)); display: flex; justify-content: center; align-items: center; gap: calc(8px * var(--scale)); }}
    .footer-char {{ width: calc(25px * var(--scale)); height: calc(25px * var(--scale)); background: url('{f_logo}') center/cover; border-radius: 50%; flex-shrink: 0; }}
    .section-title {{ font-size: calc(16px * var(--scale)); font-weight: bold; color: #0047AB; border-left: 6px solid #0047AB; padding-left: calc(10px * var(--scale)); margin: calc(4px * var(--scale)) 0; flex: 0 0 auto; }}
    .visual-container {{ width: 100%; min-height: calc(250px * var(--scale)); border: 2px dashed #0047AB; border-radius: calc(15px * var(--scale)); display: flex; align-items: center; justify-content: center; background: #fcfcfc; padding: calc(12px * var(--scale)); box-sizing: border-box; overflow: hidden; flex: 0 0 auto; }}
    .visual-image {{ width: 100%; height: 100%; max-width: 100%; max-height: calc(220px * var(--scale)); object-fit: contain; border-radius: calc(12px * var(--scale)); box-shadow: 0 4px 14px rgba(0,0,0,0.08); background: white; display: block; }}
    .title-area {{ text-align: center; margin-bottom: calc(4px * var(--scale)); flex: 0 0 auto; }}
    .title-en {{ font-size: calc(26px * var(--scale)); font-weight: 900; line-height: 1.3; }}
    .title-ko {{ font-size: calc(18px * var(--scale)); color: #666; line-height: 1.4; margin-top: calc(4px * var(--scale)); }}
    .topic-box {{ background: #0047AB; color: white; padding: calc(18px * var(--scale)); border-radius: calc(12px * var(--scale)); font-weight: bold; line-height: 1.65; font-size: calc(14.5px * var(--scale)); flex: 0 0 auto; }}
    .vocab-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: calc(8px * var(--scale)); flex: 1 1 auto; min-height: 0; overflow: hidden; }}
    .vocab-item {{ background: #f8f9fa; border: 1px solid #dee2e6; border-radius: calc(8px * var(--scale)); padding: calc(10px * var(--scale)) calc(12px * var(--scale)); overflow: hidden; min-height: 0; display: flex; align-items: stretch; justify-content: stretch; }}
    .vocab-fit {{ width: 100%; height: 100%; overflow: hidden; font-size: calc(12.5px * var(--scale)); line-height: 1.45; }}
    .vocab-word {{ font-weight: 900; color: #0047AB; display: block; margin-bottom: calc(2px * var(--scale)); font-size: 1em; line-height: 1.35; word-break: break-word; }}
    .vocab-syn {{ font-size: 0.88em; color: #666; display: block; line-height: 1.4; word-break: break-word; }}
    .passage-box {{ background: #fff; border: 2px solid #555; padding: calc(22px * var(--scale)); border-radius: calc(12px * var(--scale)); text-align: left; overflow: hidden; flex: 1 1 auto; min-height: 0; }}
    .passage-fit {{ width: 100%; height: 100%; overflow: hidden; font-size: calc(12px * var(--scale)); font-weight: 800; line-height: 2.8 !important; word-break: keep-all; }}
    .passage-meta-badge {{ display: inline-block; margin-bottom: calc(5px * var(--scale)); padding: calc(2px * var(--scale)) calc(4px * var(--scale)); background: #fff4e5; border: 1px solid #f5c27a; border-radius: 999px; color: #9a5a00; font-weight: 900; font-size: calc(5px * var(--scale)); letter-spacing: 0; }}
    .trans-content {{ font-size: calc(8px * var(--scale)); line-height: 1.75; color: #444; border-left: 5px solid #ddd; padding-left: calc(15px * var(--scale)); white-space: pre-line; overflow: hidden; flex: 1 1 auto; min-height: 0; }}
    .grammar-box {{ background: #fff9c4; padding: calc(18px * var(--scale)); border-radius: calc(12px * var(--scale)); border: 1px solid #f9e79f; font-size: calc(8px * var(--scale)); line-height: 1.75; overflow: hidden; flex: 0 0 auto; }}
    .easy-box {{ display: flex; gap: calc(10px * var(--scale)); background: #e8f0fe; padding: calc(10px * var(--scale)); border-radius: calc(12px * var(--scale)); border: 1px solid #b3d1ff; flex: 0 0 auto; }}
    .easy-char {{ width: calc(55px * var(--scale)); height: calc(55px * var(--scale)); background: url('{e_logo}') center/cover; border-radius: 12px; flex-shrink: 0; border: 2px solid white; }}
    .easy-text {{ font-size: calc(13px * var(--scale)); color: #0047AB; line-height: 1.55; overflow: hidden; }}
    .structure-box {{ text-align: center; padding: calc(10px * var(--scale)); border: 2px dashed #0047AB; border-radius: calc(15px * var(--scale)); font-weight: bold; color: #0047AB; background: #f9fbff; line-height: 1.6; white-space: pre-line; font-size: calc(12.5px * var(--scale)); overflow: hidden; flex: 0 0 auto; }}
    .quiz-container {{ background: #f8f9fa; padding: calc(16px * var(--scale)); border-radius: calc(12px * var(--scale)); border: 1px solid #dee2e6; overflow: hidden; flex: 1 1 auto; min-height: 0; }}
    .quiz-row {{ display: flex; justify-content: space-between; align-items: stretch; border-bottom: 1px solid #eee; padding: calc(6px * var(--scale)) 0; gap: calc(12px * var(--scale)); overflow: hidden; min-height: 0; }}
    .quiz-text {{ flex: 1; min-width: 0; overflow: hidden; font-size: calc(12px * var(--scale)); line-height: 1.45; }}
    .quiz-en {{ font-weight: 700; color: #222; margin-bottom: calc(4px * var(--scale)); line-height: 1.45; font-size: 1em; word-break: break-word; }}
    .quiz-ko {{ color: #4a4a4a; font-size: 0.92em; line-height: 1.45; word-break: break-word; }}
    .quiz-choice {{ white-space: nowrap; font-weight: bold; color: #0047AB; margin-left: calc(10px * var(--scale)); font-size: calc(12px * var(--scale)); letter-spacing: 1px; flex-shrink: 0; align-self: center; }}
    .quiz-answer-key {{ margin-top: calc(8px * var(--scale)); text-align: right; font-size: calc(10px * var(--scale)); color: #bbb; letter-spacing: 1px; flex: 0 0 auto; }}
    .auto-fit-box {{ overflow: hidden; }}
    .auto-fit-target {{ display: block; width: 100%; }}
    @media print {{ .page {{ margin: 0; box-shadow: none; border: none; page-break-after: always; }} }}
    .page-num {{ position: absolute; bottom: 12mm; font-size: 11px; font-weight: 900; color: #999; }}
    .page-num::after {{ content: "- " counter(pagenum) " -"; }}
    #book-container > .page:nth-child(odd) .page-num {{ right: 18mm; }}
    #book-container > .page:nth-child(even) .page-num {{ left: 18mm; }}
</style>"""

FIT_SCRIPT = """
<script>
function fitTextBox(box) {
    const target = box.querySelector('.auto-fit-target');
    if (!target) return;
    const max = parseFloat(target.dataset.max || "12.5");
    const min = parseFloat(target.dataset.min || "8");
    let size = max, guard = 0;
    target.style.fontSize = size + "px";
    target.style.lineHeight = (size * 1.40) + "px";
    while ((target.scrollHeight > box.clientHeight || target.scrollWidth > box.clientWidth) && size > min && guard < 120) {
        size -= 0.2; target.style.fontSize = size + "px"; target.style.lineHeight = (size * 1.36) + "px"; guard++;
    }
}
function fitOnePage(page) {
    let scale = 1.00;
    const content = page.querySelector('.page-content');
    if (!content) return;
    page.style.setProperty('--scale', scale.toFixed(2));
    let guard = 0;
    while (content.scrollHeight > content.clientHeight && scale > 0.72 && guard < 20) {
        scale -= 0.03; page.style.setProperty('--scale', scale.toFixed(2)); guard++;
    }
    while (content.scrollHeight > content.clientHeight && scale > 0.64 && guard < 40) {
        scale -= 0.01; page.style.setProperty('--scale', scale.toFixed(2)); guard++;
    }
}
function fitAllPages() {
    const pages = document.querySelectorAll('.fit-page');
    pages.forEach(p => p.style.setProperty('--scale', '1'));
    document.querySelectorAll('.auto-fit-box').forEach(b => fitTextBox(b));
    pages.forEach(p => fitOnePage(p));
    document.querySelectorAll('.auto-fit-box').forEach(b => fitTextBox(b));
    pages.forEach(p => fitOnePage(p));
}
window.addEventListener('load', function() {
    setTimeout(fitAllPages, 120);
    if (document.fonts && document.fonts.ready) { document.fonts.ready.then(() => { setTimeout(fitAllPages, 120); setTimeout(fitAllPages, 260); setTimeout(fitAllPages, 420); }); }
});
window.addEventListener('resize', () => setTimeout(fitAllPages, 120));
</script>
"""

def wrap_full_html_anatomy(body_html, css, front_html="", back_html=""):
    return f"""<!DOCTYPE html><html lang="ko"><head><meta charset="UTF-8"><title>Anatomy_Self_Note</title>{css}</head><body>{get_studio_toolbar_html('standard')}<div id="book-wrapper-all">{front_html}<main id="book-container">{body_html}</main>{back_html}</div>{FIT_SCRIPT}</body></html>"""

def format_passage_for_display(body_text: str) -> str:
    if not body_text:
        return ""
    raw = str(body_text).strip().replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [re.sub(r'\s+', ' ', para).strip() for para in re.split(r'\n\s*\n', raw) if para.strip()]
    return "<br><br>".join([esc(p) for p in paragraphs])

def render_anatomy(p):
    data = p["data"]
    img_url = get_image_base64(p)
    img_html = f'<img class="visual-image" src="{img_url}">' if img_url else '<div class="visual-empty"></div>'
    academy = esc(p.get('academy_name', '4pm 입시영어'))
    school = esc(p.get('school_name', ''))
    label_text = esc(p["label"])
    pid = esc(p.get("id", label_text))
    meta_html = f'<div class="passage-meta-badge">{label_text}</div>'
    body_html = format_passage_for_display(p.get('body_text', ''))

    vocab_html = "".join([
        f"""<div class="vocab-item auto-fit-box"><div class="vocab-fit auto-fit-target" data-max="12.5" data-min="8.2"><span class="vocab-word">{esc(v.get("word",""))} ({esc(v.get("meaning",""))})</span><span class="vocab-syn">{esc(v.get("synonyms_text",""))}</span></div></div>"""
        for v in data.get("vocab", [])
    ])
    grammar_html = "<br>".join([
        f"<b>{i+1}. {esc(g.get('title',''))}</b>: {esc(g.get('explanation',''))}"
        for i, g in enumerate(data.get("grammar_points", []))
    ])
    quiz_html = "".join([
        f"""<div class="quiz-row auto-fit-box"><div class="quiz-text auto-fit-target" data-max="12" data-min="8.0"><div class="quiz-en">{esc(q.get("english_statement",""))}</div><div class="quiz-ko">{esc(q.get("korean_statement",""))}</div></div><span class="quiz-choice">( O / X )</span></div>"""
        for q in data.get("quizzes", [])
    ])
    answers = "ANS: " + " ".join([f"{i+1}.{q.get('answer','')}" for i, q in enumerate(data.get("quizzes", []))])
    footer_left  = f"{academy} | 내신과 수능 1등급을 위한 필수선택 영어는 영목쌤"
    footer_right = f"{academy} | 영목쌤"
    data_attr = f'data-passage-id="{pid}" data-passage-label="{label_text}"'

    p1 = f"""<div class="page fit-page" {data_attr}><div class="header"><div class="logo-area"><div class="header-char"></div>ANATOMY PREVIEW</div><div class="info-area"><b>{academy}</b><br>{label_text}</div></div><div class="page-content"><div class="visual-container">{img_html}</div><div class="title-area"><div class="title-en">{esc(data.get("english_title",""))}</div><div class="title-ko">{esc(data.get("korean_title",""))}</div></div><div class="section-title">Topic</div><div class="topic-box">{esc(data.get("topic",""))}</div><div class="section-title">Vocabulary Focus (동의어 뜻 병기)</div><div class="vocab-grid">{vocab_html}</div></div><div class="footer"><div class="footer-char"></div>{footer_left}</div><div class="page-num"></div></div>"""
    p2 = f"""<div class="page fit-page" {data_attr}><div class="header"><div class="logo-area"><div class="header-char"></div>ANATOMY SELF-NOTE</div><div class="info-area"><b>{academy}</b><br>셀프노트: 너만의 손필기</div></div><div class="page-content"><div class="section-title">■ Original Text (강의용 필기 노트)</div><div class="passage-box auto-fit-box"><div class="passage-fit auto-fit-target" data-max="12.0" data-min="8.0">{meta_html}<div class="passage-body">{body_html}</div></div></div><div class="section-title">■ 지문 해석</div><div class="trans-content">{nl2br(data.get("translation",""))}</div></div><div class="footer"><div class="footer-char"></div>{footer_right}</div><div class="page-num"></div></div>"""
    p3 = f"""<div class="page fit-page" {data_attr}><div class="header"><div class="logo-area"><div class="header-char"></div>ANATOMY REVIEW-NOTE</div><div class="info-area"><b>{academy}</b><br>리뷰노트: 따라 쓰며 복습!</div></div><div class="page-content"><div class="section-title">■ Original Text (너만의 복습 노트)</div><div class="passage-box auto-fit-box" style="border-color:#0047AB;"><div class="passage-fit auto-fit-target" data-max="12.0" data-min="8.0">{meta_html}<div class="passage-body">{body_html}</div></div></div><div class="section-title">■ 주요 문법 Point 3 정리</div><div class="grammar-box">{grammar_html}</div></div><div class="footer"><div class="footer-char"></div>{footer_left}</div><div class="page-num"></div></div>"""
    p4 = f"""<div class="page fit-page" {data_attr}><div class="header"><div class="logo-area"><div class="header-char"></div>STRUCTURE & QUIZ</div><div class="info-area"><b>{academy}</b><br>최종 확인 | {school}</div></div><div class="page-content"><div class="section-title">■ 영쌤의 쉬운 지문 설명</div><div class="easy-box"><div class="easy-char"></div><div class="easy-text">{nl2br(data.get('easy_explanation',''))}</div></div><div class="section-title">■ 글의 흐름 (Structure)</div><div class="structure-box">{nl2br(data.get("structure_flow",""))}</div><div class="section-title">■ ANATOMY OX Quiz (5문항)</div><div class="quiz-container">{quiz_html}</div><div class="quiz-answer-key">{esc(answers)}</div></div><div class="footer"><div class="footer-char"></div>{footer_right}</div><div class="page-num"></div></div>"""
    return p1 + p2 + p3 + p4


# -----------------------------
# 3-C. [💎 ANATOMY CORE] 서식 및 렌더링
# -----------------------------
def build_css_core(h_img, f_img, e_img):
    return f"""<style>
    @import url('https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@300;400;500;700;900&display=swap');
    @page {{ size: A4; margin: 12mm 14mm; }}
    body {{ font-family: 'Noto Sans KR', sans-serif; background: #f0f2f5; margin: 0; padding: 0; color: #1e293b; -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }}
    {BOOK_COVER_CSS}
    .core-wrapper {{ width: 210mm; box-sizing: border-box; padding: 13mm 15mm; margin: 10mm auto; background: white; box-shadow: 0 0 12px rgba(0,0,0,0.1); page-break-after: always; }}
    @media print {{ body {{ background: white; }} .core-wrapper {{ width: 100%; margin: 0; padding: 0; box-shadow: none; border: none; }} }}
    .core-master-table {{ width: 100%; border-collapse: collapse; }}
    .core-master-table thead {{ display: table-header-group; }}
    .core-master-table tfoot {{ display: table-footer-group; }}
    .core-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 3px solid #0047AB; padding-bottom: 8px; margin-bottom: 10px; }}
    .core-logo-area {{ display: flex; align-items: center; gap: 10px; font-weight: 900; font-size: 22px; color: #0047AB; }}
    .core-h-img {{ width: 40px; height: 40px; background: url('{h_img}') center/cover; border-radius: 50%; border: 2px solid #0047AB; }}
    .core-footer {{ text-align: center; font-size: 10px; color: #94a3b8; border-top: 1px solid #e2e8f0; padding-top: 8px; margin-top: 12px; display: flex; justify-content: center; align-items: center; gap: 8px; }}
    .core-f-img {{ width: 22px; height: 22px; background: url('{f_img}') center/cover; border-radius: 50%; }}
    .core-sec-title {{ font-size: 13px; font-weight: 900; color: #0047AB; background: #eff6ff; border-left: 5px solid #0047AB; padding: 5px 10px; margin: 10px 0 6px 0; border-radius: 0 6px 6px 0; display: flex; justify-content: space-between; align-items: center; }}
    .core-grid-2 {{ display: grid; grid-template-columns: 1fr 1fr; gap: 8px; margin-bottom: 8px; }}
    .core-card {{ border: 1px solid #cbd5e1; border-radius: 8px; padding: 8px 10px; background: #ffffff; page-break-inside: avoid; break-inside: avoid; }}
    .core-card-title {{ font-size: 11px; font-weight: 900; color: #1e40af; margin-bottom: 5px; border-bottom: 1px dashed #cbd5e1; padding-bottom: 3px; }}
    .core-item {{ margin-bottom: 6px; font-size: 11px; line-height: 1.45; }}
    .core-item:last-child {{ margin-bottom: 0; }}
    .core-en {{ font-weight: 800; color: #0f172a; font-size: 11.5px; }}
    .core-ko {{ color: #475569; font-size: 10.5px; margin-top: 1px; }}
    .core-badge {{ display: inline-block; background: #dbeafe; color: #1d4ed8; font-weight: 800; font-size: 9.5px; padding: 1px 5px; border-radius: 4px; margin-right: 4px; }}
    .summary-card {{ border: 1.5px solid #93c5fd; background: #f8fafc; border-radius: 8px; padding: 9px 12px; margin-bottom: 7px; page-break-inside: avoid; break-inside: avoid; }}
    .summary-en {{ font-size: 12px; font-weight: 800; color: #0f172a; line-height: 1.65; margin-bottom: 4px; }}
    .summary-ko {{ font-size: 11px; color: #475569; margin-bottom: 6px; }}
    .ab-choice-pill {{ display: inline-block; background: #eff6ff; border: 1.5px solid #3b82f6; color: #1e293b; font-weight: 800; padding: 0px 7px; border-radius: 5px; margin: 0 2px; }}
    .ab-correct-opt {{ color: #0047AB; font-weight: 900; text-decoration: underline; text-underline-offset: 3px; }}
    .ab-wrong-opt {{ color: #64748b; font-weight: 700; }}
    .summary-ans-bar {{ display: flex; flex-wrap: wrap; gap: 6px 14px; background: #fffbeb; border: 1px dashed #fcd34d; border-radius: 6px; padding: 5px 9px; font-size: 10.8px; color: #78350f; }}
    .tf-row {{ border: 1px solid #e2e8f0; border-radius: 6px; padding: 6px 10px; margin-bottom: 5px; display: flex; justify-content: space-between; align-items: center; gap: 10px; background: #fcfcfd; page-break-inside: avoid; break-inside: avoid; }}
    .tf-ans-badge {{ font-weight: 900; font-size: 11px; padding: 2px 8px; border-radius: 999px; white-space: nowrap; }}
    .tf-O {{ background: #dcfce7; color: #15803d; border: 1px solid #86efac; }}
    .tf-X {{ background: #fee2e2; color: #b91c1c; border: 1px solid #fca5a5; }}
    .core-voca-strip {{ display: grid; grid-template-columns: 1fr 1fr; gap: 5px 10px; background: #fff4e5; border: 1px solid #f5c27a; border-radius: 8px; padding: 8px 10px; font-size: 10.5px; page-break-inside: avoid; break-inside: avoid; }}
    #book-container.core-test-mode .ab-choice-pill {{ background: #ffffff !important; border-color: #475569 !important; }}
    #book-container.core-test-mode .ab-correct-opt,
    #book-container.core-test-mode .ab-wrong-opt {{ color: #0f172a !important; font-weight: 800 !important; text-decoration: none !important; }}
    #book-container.core-test-mode .summary-ans-bar {{ display: none !important; }}
    #book-container.core-test-mode .tf-ans-reveal {{ display: none !important; }}
    #book-container.core-test-mode .tf-ans-badge.tf-O,
    #book-container.core-test-mode .tf-ans-badge.tf-X {{
        background: #f8fafc !important; color: #334155 !important; border: 1px solid #cbd5e1 !important;
    }}
    </style>"""

def extract_trap_word_fallback(syn_str: str, default_word: str = "opposite") -> str:
    if not syn_str:
        return default_word
    m = re.search(r'(?:\[함정\]|반의어\s*[:\-]?|/)\s*([A-Za-z\-]+)', str(syn_str))
    if m:
        return m.group(1).strip()
    words = re.findall(r'[A-Za-z\-]{3,}', str(syn_str))
    return words[-1] if words else default_word

def format_summary_dual_choice(eng_text: str, blank_a: str, dist_a: str, blank_b: str, dist_b: str, seed_key: str) -> str:
    escaped = esc(eng_text)
    rng = random.Random(seed_key)

    def make_pill(tag: str, ans_w: str, trap_w: str) -> str:
        ans_clean = esc(ans_w.strip())
        trap_clean = esc(trap_w.strip()) if trap_w and trap_w.strip() else "opposite"
        ans_span = f'<span class="ab-correct-opt">{ans_clean}</span>'
        trap_span = f'<span class="ab-wrong-opt">{trap_clean}</span>'
        opts = [ans_span, trap_span]
        rng.shuffle(opts)
        return f'<span class="ab-choice-pill">({tag}) [ {opts[0]} / {opts[1]} ]</span>'

    if blank_a:
        pill_a = make_pill("A", blank_a, dist_a)
        pat_a = r'\(\s*A\s*\)\s*' + re.escape(esc(blank_a))
        if re.search(pat_a, escaped, flags=re.IGNORECASE):
            escaped = re.sub(pat_a, pill_a, escaped, count=1, flags=re.IGNORECASE)
        else:
            escaped = re.sub(r'\(\s*A\s*\)', pill_a, escaped, count=1)

    if blank_b:
        pill_b = make_pill("B", blank_b, dist_b)
        pat_b = r'\(\s*B\s*\)\s*' + re.escape(esc(blank_b))
        if re.search(pat_b, escaped, flags=re.IGNORECASE):
            escaped = re.sub(pat_b, pill_b, escaped, count=1, flags=re.IGNORECASE)
        else:
            escaped = re.sub(r'\(\s*B\s*\)', pill_b, escaped, count=1)

    return escaped

def render_core(p: Dict) -> str:
    data = p.get("data", {})
    label_text = esc(p.get("label", ""))
    pid = esc(p.get("id", label_text))
    academy = esc(p.get("academy_name", "포피엠 입시 영어"))
    school = esc(p.get("school_name", ""))

    topics_html = "".join([
        f'<div class="core-item"><div class="core-en"><span class="core-badge">Topic {i}</span>{esc(t.get("eng", ""))}</div><div class="core-ko">↳ {esc(t.get("kor", ""))}</div></div>'
        for i, t in enumerate(data.get("topics", []), start=1)
    ])
    titles_html = "".join([
        f'<div class="core-item"><div class="core-en"><span class="core-badge">{esc(t.get("style", f"Title {i}"))}</span>{esc(t.get("eng", ""))}</div><div class="core-ko">↳ {esc(t.get("kor", ""))}</div></div>'
        for i, t in enumerate(data.get("titles", []), start=1)
    ])
    claims_html = "".join([
        f'<div class="core-item" style="border-bottom:1px dashed #e2e8f0; padding-bottom:5px;"><div class="core-en"><span class="core-badge" style="background:#e0e7ff; color:#3730a3;">요지·주장 {i}</span>{esc(c.get("eng", ""))}</div><div class="core-ko">↳ {esc(c.get("kor", ""))}</div></div>'
        for i, c in enumerate(data.get("main_ideas", []), start=1)
    ])

    summaries_html = ""
    for i, s in enumerate(data.get("summaries", []), start=1):
        ba = s.get("blank_a", "")
        bb = s.get("blank_b", "")
        sa_syn = s.get("synonyms_a", "")
        sb_syn = s.get("synonyms_b", "")
        dist_a = s.get("distractor_a") or extract_trap_word_fallback(sa_syn, "superficial")
        dist_b = s.get("distractor_b") or extract_trap_word_fallback(sb_syn, "passivity")

        eng_hl = format_summary_dual_choice(s.get("eng", ""), ba, dist_a, bb, dist_b, seed_key=f"{pid}_sum_{i}")
        summaries_html += f"""
        <div class="summary-card">
            <div class="summary-en"><span class="core-badge" style="background:#0047AB; color:white;">Summary {i}</span> {eng_hl}</div>
            <div class="summary-ko">해석: {esc(s.get('kor', ''))}</div>
            <div class="summary-ans-bar">
                <div>🎯 <b>(A) 정답: {esc(ba)}</b> vs 오답: {esc(dist_a)} <span style="color:#475569;">({esc(sa_syn)})</span></div>
                <div>🎯 <b>(B) 정답: {esc(bb)}</b> vs 오답: {esc(dist_b)} <span style="color:#475569;">({esc(sb_syn)})</span></div>
            </div>
        </div>"""

    tf_list = list(data.get("paraphrased_tf", []))
    rng_tf = random.Random(f"{pid}_tf_shuffle")
    rng_tf.shuffle(tf_list)

    tf_html = ""
    tf_ans_summary = []
    for i, tf in enumerate(tf_list, start=1):
        ans = tf.get("answer", "O")
        badge_cls = "tf-O" if ans == "O" else "tf-X"
        tf_ans_summary.append(f"{i}.{ans}")
        tf_html += f"""
        <div class="tf-row">
            <div style="flex:1;">
                <div class="core-en"><b>{i}.</b> {esc(tf.get('eng', ''))}</div>
                <div class="core-ko">↳ {esc(tf.get('kor', ''))} <span class="tf-ans-reveal" style="color:#b45309; font-weight:700;">| 💡 출제포인트: {esc(tf.get('point', ''))}</span></div>
            </div>
            <div class="tf-ans-badge {badge_cls}"><span class="tf-ans-reveal">정답: {esc(ans)} | </span>( O / X )</div>
        </div>"""

    cv_items = []
    for cv in data.get("core_vocab", [])[:10]:
        syn_txt = ", ".join([f"{s.get('word','')}({s.get('meaning','')})" for s in cv.get("synonyms", [])[:2] if s.get("word")])
        ant_txt = ", ".join([f"{a.get('word','')}({a.get('meaning','')})" for a in cv.get("antonyms", [])[:1] if a.get("word")])
        rel_str = f"유의: {syn_txt}" + (f" | 반의: {ant_txt}" if ant_txt else "")
        cv_items.append(
            f"<div>• <b style='color:#0047AB;'>{esc(cv.get('word',''))}</b> [{esc(cv.get('pronunciation',''))}] "
            f"{esc(cv.get('meaning',''))} <span style='color:#64748b; font-size:9.8px;'>({esc(rel_str)})</span></div>"
        )
    cv_html = f'<div class="core-voca-strip">{"".join(cv_items)}</div>' if cv_items else ""

    return f"""
    <div class="core-wrapper" data-passage-id="{pid}" data-passage-label="{label_text}">
        <table class="core-master-table">
            <thead>
                <tr><td>
                    <div class="core-header">
                        <div class="core-logo-area"><div class="core-h-img"></div>PQRS영쌤 | ANATOMY CORE</div>
                        <div style="text-align:right; font-size:11px;"><b>{academy}</b> {('| ' + school) if school else ''}<br><span style="color:#0047AB; font-weight:900;">{label_text}</span></div>
                    </div>
                </td></tr>
            </thead>
            <tfoot>
                <tr><td>
                    <div class="core-footer">
                        <div class="core-f-img"></div>{academy} | 내신 1등급 결정타: 주제·제목·요지·요약문 패러프레이징 완벽 마스터
                    </div>
                </td></tr>
            </tfoot>
            <tbody>
                <tr><td>
                    <div class="core-grid-2">
                        <div class="core-card">
                            <div class="core-card-title">📌 출제 1순위 영문 주제 (Topic Paraphrasing 3선)</div>
                            {topics_html}
                        </div>
                        <div class="core-card">
                            <div class="core-card-title">🏷️ 내신 빈출 영문 제목 (Title Variations 3선)</div>
                            {titles_html}
                        </div>
                    </div>
                    <div class="core-sec-title">
                        <span>💡 출제 유력 영문 요지 & 주장문 (Main Idea & Claim Paraphrasing)</span>
                        <span style="font-size:10.5px; font-weight:500; color:#475569;">원문 문장을 상위 개념어로 재구성한 선지 대비용 문장</span>
                    </div>
                    <div class="core-card" style="margin-bottom:8px;">{claims_html}</div>
                    <div class="core-sec-title">
                        <span>🎯 실전 요약문 (Summary 3선) & 빈칸 (A), (B) 양자택일 어휘 추론</span>
                        <span style="font-size:10.5px; font-weight:500; color:#475569;">문맥에 맞는 알맞은 패러프레이징 어휘를 고르시오</span>
                    </div>
                    {summaries_html}
                    <div class="core-sec-title">
                        <span>🔍 학교 내신 &lt;보기&gt;형 패러프레이징 진위(T/F) 선지 분석 (순서 랜덤 배치)</span>
                        <span class="tf-ans-reveal" style="font-size:10.5px; font-weight:700; color:#1e40af;">정답: {' / '.join(tf_ans_summary)}</span>
                    </div>
                    <div style="margin-bottom:8px;">{tf_html}</div>
                    <div class="core-sec-title" style="background:#fffbeb; border-left-color:#d97706; color:#92400e;">
                        <span>✨ CORE 패러프레이징 핵심 어휘 (ANATOMY VOCA 자동 연동 어휘)</span>
                    </div>
                    {cv_html}
                </td></tr>
            </tbody>
        </table>
    </div>"""

def wrap_full_html_core(body_html, css, front_html="", back_html=""):
    return f"""<!DOCTYPE html><html lang="ko"><head><meta charset="UTF-8"><title>Anatomy_Core</title>{css}</head><body>{get_studio_toolbar_html('core')}<div id="book-wrapper-all">{front_html}<main id="book-container">{body_html}</main>{back_html}</div></body></html>"""


# -----------------------------
# 3-D. [📖 ANATOMY VOCA & 기출 VOCA] 서식 및 시험지별 연속 렌더링
# -----------------------------
def build_css_voca(h_img, f_img):
    return f"""<style>
    @import url('https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@300;400;500;700;900&display=swap');
    @page {{ size: A4; margin: 12mm 14mm; }}
    body {{ font-family: 'Noto Sans KR', sans-serif; background: #f0f2f5; margin: 0; padding: 0; color: #1e293b; -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }}
    {BOOK_COVER_CSS}
    .voca-wrapper {{ width: 210mm; box-sizing: border-box; padding: 14mm 16mm; margin: 10mm auto; background: white; box-shadow: 0 0 12px rgba(0,0,0,0.1); page-break-after: always; break-after: page; }}
    @media print {{ body {{ background: white; }} .voca-wrapper {{ width: 100%; margin: 0; padding: 0; box-shadow: none; border: none; page-break-after: always; break-after: page; }} }}
    .voca-master-table {{ width: 100%; border-collapse: collapse; }}
    .voca-master-table thead {{ display: table-header-group; }}
    .voca-master-table tfoot {{ display: table-footer-group; }}
    .voca-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 3px solid #0047AB; padding-bottom: 8px; margin-bottom: 12px; }}
    .voca-logo-area {{ display: flex; align-items: center; gap: 10px; font-weight: 900; font-size: 23px; color: #0047AB; letter-spacing: -0.5px; }}
    .voca-h-img {{ width: 40px; height: 40px; background: url('{h_img}') center/cover; border-radius: 50%; border: 2px solid #0047AB; }}
    .voca-footer {{ text-align: center; font-size: 10px; color: #94a3b8; border-top: 1px solid #e2e8f0; padding-top: 8px; margin-top: 14px; display: flex; justify-content: center; align-items: center; gap: 8px; }}
    .voca-f-img {{ width: 22px; height: 22px; background: url('{f_img}') center/cover; border-radius: 50%; }}
    .voca-subbar {{ display: flex; justify-content: space-between; align-items: center; background: #f8fafc; border: 1px solid #cbd5e1; border-left: 5px solid #0047AB; border-radius: 6px; padding: 8px 14px; margin-bottom: 12px; font-size: 12px; }}
    .voca-table {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
    .voca-table thead {{ display: table-header-group; }}
    .voca-table th {{ background: #0047AB; color: white; font-size: 11.5px; font-weight: 800; padding: 7px 8px; border: 1px solid #003580; text-align: center; }}
    .voca-table tr {{ page-break-inside: avoid; break-inside: avoid; }}
    .voca-table td {{ border: 1px solid #cbd5e1; padding: 7px 9px; vertical-align: middle; font-size: 11.5px; line-height: 1.45; word-break: keep-all; }}
    .voca-table tbody tr:nth-child(even) {{ background: #f8fafc; }}
    .voca-table tbody tr.core-para-row {{ background: #fffbeb; }}
    .voca-idx-cell {{ text-align: center; font-weight: 900; color: #64748b; font-size: 11px; position: relative; }}
    .voca-word-main {{ font-weight: 900; font-size: 14px; color: #0f172a; display: block; margin-bottom: 2px; }}
    .voca-pron-badge {{ display: inline-block; font-size: 10.5px; color: #0369a1; background: #e0f2fe; padding: 1px 6px; border-radius: 4px; font-weight: 700; margin-right: 4px; }}
    .voca-pos-badge {{ display: inline-block; font-size: 10px; color: #475569; background: #e2e8f0; padding: 1px 5px; border-radius: 4px; font-weight: 800; margin-right: 4px; }}
    .voca-cat-core {{ display: inline-block; font-size: 9.5px; color: #92400e; background: #fef3c7; border: 1px solid #f59e0b; padding: 0px 5px; border-radius: 4px; font-weight: 900; margin-top: 2px; }}
    .voca-cat-text {{ display: inline-block; font-size: 9.5px; color: #1e40af; background: #dbeafe; padding: 0px 5px; border-radius: 4px; font-weight: 800; margin-top: 2px; }}
    .voca-meaning-cell {{ font-weight: 700; color: #1e293b; font-size: 12.5px; }}
    .syn-row-chip {{ display: block; background: #eff6ff; border: 1px solid #bfdbfe; color: #1e40af; border-radius: 5px; padding: 3px 7px; margin-bottom: 3px; font-size: 10.8px; }}
    .ant-row-chip {{ display: block; background: #fff1f2; border: 1px solid #fecdd3; color: #9f1239; border-radius: 5px; padding: 3px 7px; margin-bottom: 3px; font-size: 10.8px; }}
    .chip-tag-syn {{ font-weight: 900; color: #2563eb; margin-right: 4px; }}
    .chip-tag-ant {{ font-weight: 900; color: #e11d48; margin-right: 4px; }}
    .chip-pron {{ color: #64748b; font-size: 10px; margin: 0 3px; }}
    .voca-del-btn {{ background: #fee2e2; color: #dc2626; border: 1px solid #fca5a5; border-radius: 4px; font-size: 9px; padding: 1px 4px; cursor: pointer; margin-top: 3px; }}
    .voca-del-btn:hover {{ background: #ef4444; color: white; }}
    #book-container.mode-hide-meaning .voca-meaning-text,
    #book-container.mode-hide-meaning .chip-meaning-text {{ visibility: hidden; }}
    #book-container.mode-hide-word .voca-word-main,
    #book-container.mode-hide-word .chip-word-text {{ visibility: hidden; }}
    </style>"""

def render_voca(p):
    data = p.get("data", {})
    voca_list = data.get("voca_list", [])
    label_text = esc(p.get("label", ""))
    pid = esc(p.get("id", label_text))
    academy = esc(p.get("academy_name", "포피엠 입시 영어"))
    school = esc(p.get("school_name", ""))

    rows_html = ""
    for idx, item in enumerate(voca_list, start=1):
        w = esc(item.get("word", ""))
        pos = esc(item.get("pos", ""))
        pron = esc(item.get("pronunciation", ""))
        mean = esc(item.get("meaning", ""))
        cat = str(item.get("category", "지문핵심")).strip()
        is_core = any(k in cat.upper() for k in ["CORE", "패러", "요약", "주제", "선지"])
        cat_badge = f'<span class="voca-cat-core">{esc(cat)}</span>' if is_core else f'<span class="voca-cat-text">{esc(cat)}</span>'
        row_cls = 'class="core-para-row"' if is_core else ""

        syn_chips = []
        for s in item.get("synonyms", []):
            sw = esc(s.get("word", ""))
            spron = esc(s.get("pronunciation", ""))
            smean = esc(s.get("meaning", ""))
            if sw:
                pron_part = f'<span class="chip-pron">[{spron}]</span>' if spron else ""
                syn_chips.append(
                    f'<div class="syn-row-chip"><span class="chip-tag-syn">[유의어/대체]</span>'
                    f'<b class="chip-word-text">{sw}</b>{pron_part} : '
                    f'<span class="chip-meaning-text">{smean}</span></div>'
                )

        ant_chips = []
        for a in item.get("antonyms", []):
            aw = esc(a.get("word", ""))
            apron = esc(a.get("pronunciation", ""))
            amean = esc(a.get("meaning", ""))
            if aw:
                pron_part = f'<span class="chip-pron">[{apron}]</span>' if apron else ""
                ant_chips.append(
                    f'<div class="ant-row-chip"><span class="chip-tag-ant">[반의어/함정]</span>'
                    f'<b class="chip-word-text">{aw}</b>{pron_part} : '
                    f'<span class="chip-meaning-text">{amean}</span></div>'
                )

        rel_html = "".join(syn_chips + ant_chips)
        if not rel_html:
            rel_html = '<span style="color:#cbd5e1; font-size:10.5px;">-</span>'

        rows_html += f"""
        <tr {row_cls}>
            <td class="voca-idx-cell">
                <div class="voca-idx-num">{idx}</div>
                <button class="voca-del-btn no-print" onclick="removeVocaRow(this)" title="이 단어 제외">✕</button>
            </td>
            <td>
                <span class="voca-word-main">{w}</span>
                <span class="voca-pron-badge">[{pron}]</span>
                <span class="voca-pos-badge">{pos}</span><br>
                {cat_badge}
            </td>
            <td class="voca-meaning-cell">
                <span class="voca-meaning-text">{mean}</span>
            </td>
            <td>{rel_html}</td>
            <td style="text-align:center; color:#cbd5e1; font-size:13px;">□ &nbsp; □</td>
        </tr>"""

    return f"""
    <div class="voca-wrapper" data-passage-id="{pid}" data-passage-label="{label_text}">
        <table class="voca-master-table">
            <thead>
                <tr><td>
                    <div class="voca-header">
                        <div class="voca-logo-area"><div class="voca-h-img"></div>PQRS영쌤 | ANATOMY VOCA</div>
                        <div style="text-align:right; font-size:11.5px;"><b>{academy}</b> {('| ' + school) if school else ''}<br><span style="color:#0047AB; font-weight:800;">{label_text}</span></div>
                    </div>
                </td></tr>
            </thead>
            <tfoot>
                <tr><td>
                    <div class="voca-footer">
                        <div class="voca-f-img"></div>포피엠 입시 영어 | 지문·선지 핵심 어휘 + 패러프레이징 유의어·반의어 통합 단어장
                    </div>
                </td></tr>
            </tfoot>
            <tbody>
                <tr><td>
                    <div class="voca-subbar">
                        <div>📌 <b>출처 시험지/지문:</b> {label_text} (총 <b>{len(voca_list)}</b>개 핵심 어휘 연속 수록)</div>
                        <div>이름: ________________ &nbsp;&nbsp;|&nbsp;&nbsp; 점수: ______ / {len(voca_list)}</div>
                    </div>
                    <table class="voca-table">
                        <colgroup>
                            <col style="width: 6%;">
                            <col style="width: 28%;">
                            <col style="width: 22%;">
                            <col style="width: 36%;">
                            <col style="width: 8%;">
                        </colgroup>
                        <thead>
                            <tr>
                                <th>No.</th>
                                <th>단어 / 숙어 [한글발음] &amp; 문항출처</th>
                                <th>문맥 및 패러프레이징 핵심 의미</th>
                                <th>문맥 유의어 (Synonyms) &amp; 내신 함정 반의어 (Antonyms)</th>
                                <th>암기체크</th>
                            </tr>
                        </thead>
                        <tbody>{rows_html}</tbody>
                    </table>
                </td></tr>
            </tbody>
        </table>
    </div>"""

def render_past_voca_grouped_by_exam(sel_passages: List[Dict]) -> str:
    """기출 어휘 렌더링 시 동일한 시험지에 속한 항목들은 한 시험지 테이블에 연달아 출력하고, 다음 시험지로 넘어갈 때만 페이지를 분리"""
    consolidated, _ = consolidate_past_voca_records(sel_passages)
    return "".join([render_voca(p) for p in consolidated])

def wrap_full_html_voca(body_html, css, front_html="", back_html=""):
    return f"""<!DOCTYPE html><html lang="ko"><head><meta charset="UTF-8"><title>Anatomy_Voca</title>{css}</head><body>{get_studio_toolbar_html('voca')}<div id="book-wrapper-all">{front_html}<main id="book-container">{body_html}</main>{back_html}</div></body></html>"""


# -----------------------------
# 4. AI 분석 엔진 (한판승 & ANATOMY & ANATOMY CORE & ANATOMY VOCA)
# -----------------------------
SCHEMA_VICTORY = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "keywords": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {"word": {"type": "string"}, "meaning": {"type": "string"}, "synonyms": {"type": "string"}},
                "required": ["word", "meaning", "synonyms"]
            }
        },
        "analysis": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "eng": {"type": "string"},
                    "sentence_vocab": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {"word": {"type": "string"}, "pronunciation": {"type": "string"}, "pos": {"type": "string"}, "meaning": {"type": "string"}}
                        }
                    },
                    "kor": {"type": "string"},
                    "grammar": {"type": "string"},
                    "ys": {"type": "string"}
                },
                "required": ["eng", "kor", "ys"]
            }
        }
    },
    "required": ["summary", "keywords", "analysis"]
}

SCHEMA_ANATOMY = {
    "type": "object",
    "properties": {
        "english_title": {"type": "string"}, "korean_title": {"type": "string"}, "topic": {"type": "string"},
        "vocab": {"type": "array", "items": {"type": "object", "properties": {"word": {"type": "string"}, "meaning": {"type": "string"}, "synonyms_text": {"type": "string"}}, "required": ["word", "meaning", "synonyms_text"]}},
        "translation": {"type": "string"},
        "grammar_points": {"type": "array", "items": {"type": "object", "properties": {"title": {"type": "string"}, "explanation": {"type": "string"}}}},
        "easy_explanation": {"type": "string"}, "structure_flow": {"type": "string"},
        "quizzes": {"type": "array", "items": {"type": "object", "properties": {"english_statement": {"type": "string"}, "korean_statement": {"type": "string"}, "answer": {"type": "string", "enum": ["O", "X"]}}}},
        "visual_guide": {"type": "string"}
    },
    "required": ["english_title", "korean_title", "topic", "vocab", "translation", "grammar_points", "easy_explanation", "structure_flow", "quizzes", "visual_guide"]
}

SCHEMA_CORE = {
    "type": "object",
    "properties": {
        "topics": {"type": "array", "items": {"type": "object", "properties": {"eng": {"type": "string"}, "kor": {"type": "string"}}, "required": ["eng", "kor"]}},
        "titles": {"type": "array", "items": {"type": "object", "properties": {"style": {"type": "string"}, "eng": {"type": "string"}, "kor": {"type": "string"}}, "required": ["style", "eng", "kor"]}},
        "main_ideas": {"type": "array", "items": {"type": "object", "properties": {"eng": {"type": "string"}, "kor": {"type": "string"}}, "required": ["eng", "kor"]}},
        "summaries": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "eng": {"type": "string"}, "kor": {"type": "string"},
                    "blank_a": {"type": "string"}, "distractor_a": {"type": "string"},
                    "blank_b": {"type": "string"}, "distractor_b": {"type": "string"},
                    "synonyms_a": {"type": "string"}, "synonyms_b": {"type": "string"}
                },
                "required": ["eng", "kor", "blank_a", "distractor_a", "blank_b", "distractor_b", "synonyms_a", "synonyms_b"]
            }
        },
        "paraphrased_tf": {"type": "array", "items": {"type": "object", "properties": {"eng": {"type": "string"}, "kor": {"type": "string"}, "answer": {"type": "string", "enum": ["O", "X"]}, "point": {"type": "string"}}, "required": ["eng", "kor", "answer", "point"]}},
        "core_vocab": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "word": {"type": "string"}, "pos": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"},
                    "synonyms": {"type": "array", "items": {"type": "object", "properties": {"word": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"}}, "required": ["word", "pronunciation", "meaning"]}},
                    "antonyms": {"type": "array", "items": {"type": "object", "properties": {"word": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"}}, "required": ["word", "pronunciation", "meaning"]}}
                },
                "required": ["word", "pos", "pronunciation", "meaning", "synonyms", "antonyms"]
            }
        }
    },
    "required": ["topics", "titles", "main_ideas", "summaries", "paraphrased_tf", "core_vocab"]
}

SCHEMA_VOCA = {
    "type": "object",
    "properties": {
        "voca_list": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "word": {"type": "string"}, "pos": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"}, "category": {"type": "string"},
                    "synonyms": {"type": "array", "items": {"type": "object", "properties": {"word": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"}}, "required": ["word", "pronunciation", "meaning"]}},
                    "antonyms": {"type": "array", "items": {"type": "object", "properties": {"word": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"}}, "required": ["word", "pronunciation", "meaning"]}}
                },
                "required": ["word", "pos", "pronunciation", "meaning", "category", "synonyms", "antonyms"]
            }
        }
    },
    "required": ["voca_list"]
}

SCHEMA_PAST_PDF_VOCA = {
    "type": "object",
    "properties": {
        "school_name": {"type": "string", "description": "학교명 (예: 운양고, 제일고)"},
        "exam_year_term": {"type": "string", "description": "연도 및 학기 (예: 26년 1학기 중간, 26년 1학기 기말)"},
        "grade": {"type": "string", "description": "학년 (예: 고1, 고2, 고3)"},
        "questions": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "question_number": {"type": "string", "description": "문항 번호 (예: 1번, 18번, 21번, 논술형1 등)"},
                    "words": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "word": {"type": "string"},
                                "pos": {"type": "string"},
                                "pronunciation": {"type": "string"},
                                "meaning": {"type": "string"},
                                "source_part": {"type": "string", "description": "'기출본문' 또는 '기출선지'"},
                                "synonyms": {"type": "array", "items": {"type": "object", "properties": {"word": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"}}, "required": ["word", "pronunciation", "meaning"]}},
                                "antonyms": {"type": "array", "items": {"type": "object", "properties": {"word": {"type": "string"}, "pronunciation": {"type": "string"}, "meaning": {"type": "string"}}, "required": ["word", "pronunciation", "meaning"]}}
                            },
                            "required": ["word", "pos", "pronunciation", "meaning", "source_part", "synonyms", "antonyms"]
                        }
                    }
                },
                "required": ["question_number", "words"]
            }
        }
    },
    "required": ["school_name", "exam_year_term", "grade", "questions"]
}

def analyze_passage(label, text, model_name, target_type="victory", core_context: Dict = None, max_retries=5):
    dna_rules = st.session_state.exam_dna.get("rules_text", DEFAULT_EXAM_DNA["rules_text"])

    if target_type == "victory":
        prompt = f"""너는 한국 고등학생 영어 내신용 '한판승' 교재 편집자다. 지문을 분석해 JSON만 반환하라.
1. summary: 반드시 '한글(한국어)'로 작성된 핵심 요약 2문장.
2. keywords: 중요 단어 8개. synonyms는 반드시 "동의어1(뜻), 동의어2(뜻)" 형식으로 2개 작성.
3. analysis: 모든 문장 분석. 
   - eng: 문장에 끊어 읽기 표시(/)를 삽입해라. 단, 종속절의 자잘한 주어/동사가 아닌 **문장 전체의 메인 주어(S), 본동사(V), 메인 목적어(O), 메인 보어(C)**만 표시하고, 관계대명사절 등 나머지는 묶어서 수식어구(M)로 처리해라.
   ★(중요) 표시 형태는 반드시 해당 어구를 대괄호로 묶고 바로 뒤에 괄호로 기호를 붙여라.
   (예시: [Factors](S) / [such as team chemistry, which is related to how well people get along,](M) / [will likely remain](V) / [difficult to measure](C).)
   - sentence_vocab: 해당 문장에서 사용된 주요 단어들을 추출. 유의어 없이 딱 문장에 쓰인 의미만 작성. (word: 영어단어, pronunciation: 한글발음기호 예:나이스, pos: 품사 약어(n, v, a, ad 등), meaning: 뜻)
   - kor: 직독직해.
   - grammar: 해당 문장에서 중요한 문법이나 구문 설명 (없으면 빈 문자열).
   - ys: 친근한 구어체 설명(~야, ~지!, ~거거든!), 해석된 한글 문장을 쉽고 간략한 예를 들어서 이해하기 쉽게 다른말로 설명해줘.
지문: \"\"\"{text}\"\"\""""
        schema = SCHEMA_VICTORY

    elif target_type == "anatomy":
        prompt = f"""너는 한국 고등학생용 영어 내신 교재를 만드는 최상급 편집자다. 지문을 분석해 반드시 JSON만 반환하라.
지문 정보: {label}
[반환 규칙] 
1. vocab: 정확히 10개 추출. synonyms_text는 반드시 "동의어1(뜻), 동의어2(뜻)" 형식으로 작성.
2. translation: 지문 전체의 자연스러운 한국어 해석.
3. grammar_points: 핵심 문법 3개 (title, explanation).
4. structure_flow: 반드시 '➔'를 사용한 도식형으로 작성.
5. quizzes: 5개의 OX퀴즈. answer는 'O' 또는 'X'만 사용.
지문: \"\"\"{text}\"\"\""""
        schema = SCHEMA_ANATOMY

    elif target_type == "core":
        prompt = f"""너는 대한민국 고등학교 영어 내신 변형 문제 및 패러프레이징 최고 전문가다.
아래의 [학교 실제 기출 출제 DNA]를 철저히 적용하여, 이 지문에서 출제될 가능성이 가장 높은 다중 주제문, 제목, 요지/주장문, 요약문(빈칸 A/B 양자택일 포함), <보기> 진위형 선지, 그리고 여기에 사용된 패러프레이징 핵심 어휘를 JSON으로 반환하라.

{dna_rules}

[세부 생성 규칙]
1. topics (영문 주제문 3선): 원문 표현을 그대로 복사하지 말고 상위 개념어(Meta-vocabulary)와 명사구 구문을 활용해 서로 다른 각도에서 3개 작성 (eng, kor).
2. titles (영문 제목 3선): 1번(style="콜론(:) 핵심형"), 2번(style="의문문(?) 탐구형"), 3번(style="대조/비유 압축형").
3. main_ideas (영문 요지·주장문 3선): 완결된 학술 영문 평서문 3개로 패러프레이징 작성 (eng, kor).
4. summaries (실전 요약문 3선 & 양자택일 어휘):
   - 서로 다른 문장 구조로 지문 전체를 요약한 영문 문장 3개 작성.
   - 반드시 영문 요약문(eng) 안에 `(A) 정답단어` 와 `(B) 정답단어` 형태로 표기할 것!
   - blank_a, blank_b에는 각각 (A), (B)의 정확한 정답 영어 단어를 적고, distractor_a, distractor_b에는 각각 (A), (B) 자리에 들어갈 매력적인 오답 영어 단어(반의어 또는 철자 유사 함정 어휘 1단어, 품사 일치 필수)를 적을 것.
   - synonyms_a, synonyms_b에는 각각 동의어 2개와 오답 단어의 뜻을 "동의어1(뜻), 동의어2(뜻) / [함정]오답단어(뜻)" 형태로 작성할 것.
5. paraphrased_tf (내신 <보기>형 패러프레이징 진위 문장 4선):
   - 2개는 O(True: 참 진술), 2개는 X(False: 매력적인 오답 진술). 단, O와 X의 배치 순서는 예측할 수 없게 섞어서 작성하라.
6. core_vocab (위 주제/제목/요지/요약문/진위문에서 사용된 핵심 패러프레이징 어휘 8~10개):
   - pronunciation은 반드시 '한글 발음'으로만 적고, synonyms(2개)와 antonyms(1~2개)도 각각 word, pronunciation(한글발음), meaning(뜻)을 완비할 것.

지문 정보: {label}
지문: \"\"\"{text}\"\"\""""
        schema = SCHEMA_CORE

    else:
        core_ref_str = ""
        if core_context:
            core_ref_str = f"""
[★중요 연동 지시: 이 지문의 ANATOMY CORE에서 추출된 핵심 패러프레이징 어휘 목록]
{json.dumps(core_context.get('core_vocab', []), ensure_ascii=False)}
위 ANATOMY CORE 어휘들을 누락 없이 단어장에 포함시키고, 해당 단어들의 category 필드는 반드시 "★CORE 패러프레이징"으로 설정하라!
"""
        prompt = f"""너는 대한민국 고등학교 내신 및 수능 영어 어휘 전문 편집자다.
아래의 [학교 실제 기출 출제 DNA]를 반영하여, 지문 속 내신 변형 1순위 어휘/숙어와 주제·요약문 패러프레이징 핵심 어휘를 총 15~22개 엄선해 JSON으로 반환하라.

{dna_rules}
{core_ref_str}

[절대 준수 규칙]
1. 초등 기초 단어 제외: be동사, 관사(a, an, the), 대명사, 단순 전치사, beautiful, happy, make, good 같은 기초 단어 제외.
2. 2-Track 어휘 구성:
   - Track 1 (`category`: "지문핵심"): 지문에 등장한 중고등 핵심 어휘 및 숙어 10~14개.
   - Track 2 (`category`: "★CORE 패러프레이징"): 주제문, 제목, 요지문, 요약문 빈칸(A)/(B)용 상위 개념어 및 핵심 요약 어휘 5~8개.
3. 각 항목 세부 규칙:
   - word: 영어 단어 또는 숙어 / pos: 품사 약어 / pronunciation: 한글 발음(IPA 금지, 한글로만 표기!) / meaning: 한국어 뜻
   - synonyms: 유의어/대체 표현 2~3개 (word, pronunciation, meaning 필수).
   - antonyms: 반의어 또는 철자 유사 혼동어 1~2개 (word, pronunciation, meaning 필수).

지문 정보: {label}
지문: \"\"\"{text}\"\"\""""
        schema = SCHEMA_VOCA

    for attempt in range(max_retries):
        try:
            res = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=0.3,
                    response_mime_type="application/json",
                    response_json_schema=schema
                )
            )
            return json.loads(res.text.strip())
        except Exception as e:
            error_str = str(e).lower()
            if any(k in error_str for k in ["503", "high demand", "429", "unavailable", "quota", "overloaded"]):
                if attempt < max_retries - 1:
                    wait_time = 10 + (attempt * 10)
                    st.toast(f"⏳ 서버 대기 감지 ({label} - {target_type}). {wait_time}초 후 재시도... ({attempt + 1}/{max_retries})")
                    time.sleep(wait_time)
                    continue
            st.error(f"❌ 오류 발생 ({label} - {target_type}): {e}")
            return None
    return None


# -----------------------------
# 4-B. 원본 JSON 영구 수정 에디터 (5개 라이브러리 전체 지원)
# -----------------------------
def parse_rel_string(rel_str: str) -> List[Dict]:
    res_list = []
    if not rel_str or pd.isna(rel_str):
        return res_list
    chunks = [c.strip() for c in re.split(r',\s*(?![^\(\[]*[\)\]])', str(rel_str)) if c.strip()]
    for ch in chunks:
        m = re.match(r'^([^\[\(]+?)(?:\[(.*?)\])?\s*(?:\((.*?)\))?$', ch.strip())
        if m:
            res_list.append({
                "word": (m.group(1) or "").strip(),
                "pronunciation": (m.group(2) or "").strip(),
                "meaning": (m.group(3) or "").strip()
            })
        else:
            res_list.append({"word": ch, "pronunciation": "", "meaning": ""})
    return res_list

def render_permanent_json_editor(prefix: str, target_type: str, file_path: str, selected_passages: List[Dict], all_saved: List[Dict]):
    if not selected_passages:
        return

    with st.expander("🛠️ 선택 지문/시험지 원본 데이터 영구 수정 (수정 후 클라우드/JSON 파일에 영구 반영)", expanded=False):
        passage_map = {f"{p['label']} (ID: {p['id'][:6]})": p for p in selected_passages}
        chosen_key = st.selectbox("✏️ 수정할 항목을 선택하세요", options=list(passage_map.keys()), key=f"{prefix}_edit_select")
        target_p = passage_map[chosen_key]
        pid = target_p["id"]
        pdata = target_p.get("data", {})

        new_label = st.text_input("📌 시험지/지문 라벨명", value=target_p.get("label", ""), key=f"{prefix}_elbl_{pid}")

        if target_type in ["voca", "past_voca"]:
            st.markdown("##### 📖 단어장 목록 편집 (셀 더블클릭 수정, 하단 `+`로 행 추가, 좌측 체크 후 삭제 가능)")
            st.caption("💡 유의어와 반의어는 `단어[한글발음](뜻), 단어2[한글발음](뜻)` 형식으로 입력하면 구조가 그대로 보존됩니다.")

            rows_for_editor = []
            for item in pdata.get("voca_list", []):
                syn_str = ", ".join([
                    f"{s.get('word','')}[{s.get('pronunciation','')}]({s.get('meaning','')})"
                    for s in item.get("synonyms", []) if s.get("word")
                ])
                ant_str = ", ".join([
                    f"{a.get('word','')}[{a.get('pronunciation','')}]({a.get('meaning','')})"
                    for a in item.get("antonyms", []) if a.get("word")
                ])
                rows_for_editor.append({
                    "출제 트랙/문항출처": item.get("category", "지문핵심"),
                    "영어 단어/숙어": item.get("word", ""),
                    "품사": item.get("pos", ""),
                    "한글 발음": item.get("pronunciation", ""),
                    "한국어 뜻": item.get("meaning", ""),
                    "유의어 (단어[발음](뜻))": syn_str,
                    "반의어 (단어[발음](뜻))": ant_str,
                })

            edited_df = st.data_editor(
                pd.DataFrame(rows_for_editor) if rows_for_editor else pd.DataFrame(columns=["출제 트랙/문항출처", "영어 단어/숙어", "품사", "한글 발음", "한국어 뜻", "유의어 (단어[발음](뜻))", "반의어 (단어[발음](뜻))"]),
                num_rows="dynamic",
                use_container_width=True,
                key=f"{prefix}_voca_df_{pid}"
            )

            if st.button("💾 단어장 수정사항 영구 저장", type="primary", key=f"{prefix}_save_json_{pid}"):
                new_voca_list = []
                for _, r in edited_df.iterrows():
                    w_val = str(r.get("영어 단어/숙어", "")).strip()
                    if not w_val or w_val == "nan":
                        continue
                    new_voca_list.append({
                        "category": str(r.get("출제 트랙/문항출처", "지문핵심")).strip(),
                        "word": w_val,
                        "pos": str(r.get("품사", "")).strip(),
                        "pronunciation": str(r.get("한글 발음", "")).strip(),
                        "meaning": str(r.get("한국어 뜻", "")).strip(),
                        "synonyms": parse_rel_string(r.get("유의어 (단어[발음](뜻))", "")),
                        "antonyms": parse_rel_string(r.get("반의어 (단어[발음](뜻))", ""))
                    })
                target_p["label"] = clean_label(new_label)
                target_p["data"]["voca_list"] = new_voca_list
                save_to_file(file_path, all_saved)
                st.success(f"✅ [{target_p['label']}] 단어장 데이터가 영구 업데이트되었습니다!")

        elif target_type == "core":
            st.markdown("##### 📌 영문 주제 (Topics) 편집")
            top_df = st.data_editor(pd.DataFrame(pdata.get("topics", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_ctop_{pid}")
            st.markdown("##### 🏷️ 영문 제목 (Titles) 편집")
            tit_df = st.data_editor(pd.DataFrame(pdata.get("titles", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_ctit_{pid}")
            st.markdown("##### 💡 영문 요지·주장 (Main Ideas) 편집")
            mi_df = st.data_editor(pd.DataFrame(pdata.get("main_ideas", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_cmi_{pid}")
            st.markdown("##### 🎯 실전 요약문 (Summaries & 빈칸 A/B 양자택일) 편집")
            sum_df = st.data_editor(pd.DataFrame(pdata.get("summaries", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_csum_{pid}")
            st.markdown("##### 🔍 <보기>형 진위 진술문 (Paraphrased T/F) 편집")
            tf_df = st.data_editor(pd.DataFrame(pdata.get("paraphrased_tf", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_ctf_{pid}")

            if st.button("💾 ANATOMY CORE 수정사항 영구 저장", type="primary", key=f"{prefix}_save_json_{pid}"):
                target_p["label"] = clean_label(new_label)
                pdata["topics"] = top_df.to_dict(orient="records")
                pdata["titles"] = tit_df.to_dict(orient="records")
                pdata["main_ideas"] = mi_df.to_dict(orient="records")
                pdata["summaries"] = sum_df.to_dict(orient="records")
                pdata["paraphrased_tf"] = tf_df.to_dict(orient="records")
                save_to_file(file_path, all_saved)
                st.success(f"✅ [{target_p['label']}] ANATOMY CORE 데이터가 영구 업데이트되었습니다!")

        elif target_type == "victory":
            new_sum = st.text_area("🎯 한판승 핵심 요약", value=pdata.get("summary", ""), height=80, key=f"{prefix}_esum_{pid}")
            kw_df = st.data_editor(pd.DataFrame(pdata.get("keywords", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_ekw_{pid}")
            st.markdown("##### 📝 문장별 구문/해석/코칭 수정")
            ana_rows = [{"eng": row.get("eng", ""), "kor": row.get("kor", ""), "grammar": row.get("grammar", ""), "ys": row.get("ys", "")} for row in pdata.get("analysis", [])]
            ana_df = st.data_editor(pd.DataFrame(ana_rows), use_container_width=True, key=f"{prefix}_eana_{pid}")

            if st.button("💾 한판승 수정사항 영구 저장", type="primary", key=f"{prefix}_save_json_{pid}"):
                target_p["label"] = clean_label(new_label)
                pdata["summary"] = new_sum
                pdata["keywords"] = kw_df.to_dict(orient="records")
                orig_analysis = pdata.get("analysis", [])
                updated_analysis = []
                for idx, r in ana_df.iterrows():
                    sc_vocab = orig_analysis[idx].get("sentence_vocab", []) if idx < len(orig_analysis) else []
                    updated_analysis.append({
                        "eng": str(r.get("eng", "")), "kor": str(r.get("kor", "")),
                        "grammar": str(r.get("grammar", "")), "ys": str(r.get("ys", "")),
                        "sentence_vocab": sc_vocab
                    })
                pdata["analysis"] = updated_analysis
                save_to_file(file_path, all_saved)
                st.success(f"✅ [{target_p['label']}] 한판승 데이터가 영구 업데이트되었습니다!")

        elif target_type == "anatomy":
            c1, c2 = st.columns(2)
            new_en_t = c1.text_input("영문 제목", value=pdata.get("english_title", ""), key=f"{prefix}_eent_{pid}")
            new_ko_t = c2.text_input("한글 제목", value=pdata.get("korean_title", ""), key=f"{prefix}_ekot_{pid}")
            new_topic = st.text_area("Topic", value=pdata.get("topic", ""), height=65, key=f"{prefix}_etop_{pid}")
            new_trans = st.text_area("지문 해석", value=pdata.get("translation", ""), height=130, key=f"{prefix}_etr_{pid}")
            new_easy = st.text_area("영쌤의 쉬운 지문 설명", value=pdata.get("easy_explanation", ""), height=90, key=f"{prefix}_eeasy_{pid}")
            new_flow = st.text_area("글의 흐름 (Structure)", value=pdata.get("structure_flow", ""), height=70, key=f"{prefix}_eflow_{pid}")
            v_df = st.data_editor(pd.DataFrame(pdata.get("vocab", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_evoc_{pid}")
            q_df = st.data_editor(pd.DataFrame(pdata.get("quizzes", [])), num_rows="dynamic", use_container_width=True, key=f"{prefix}_eq_{pid}")

            if st.button("💾 ANATOMY 수정사항 영구 저장", type="primary", key=f"{prefix}_save_json_{pid}"):
                target_p["label"] = clean_label(new_label)
                pdata["english_title"], pdata["korean_title"] = new_en_t, new_ko_t
                pdata["topic"], pdata["translation"] = new_topic, new_trans
                pdata["easy_explanation"], pdata["structure_flow"] = new_easy, new_flow
                pdata["vocab"] = v_df.to_dict(orient="records")
                pdata["quizzes"] = q_df.to_dict(orient="records")
                save_to_file(file_path, all_saved)
                st.success(f"✅ [{target_p['label']}] ANATOMY 데이터가 영구 업데이트되었습니다!")


# -----------------------------
# 5. 초고속 라이브러리 관리 컴포넌트 (@st.fragment 적용)
# -----------------------------
def parse_number_query(query_str: str) -> List[int]:
    result = []
    if not query_str.strip():
        return result
    parts = re.split(r'[,;\s]+', query_str.strip())
    for part in parts:
        if not part:
            continue
        if "-" in part or "~" in part:
            bounds = re.split(r'[\-~]', part)
            if len(bounds) == 2 and bounds[0].isdigit() and bounds[1].isdigit():
                start, end = int(bounds[0]), int(bounds[1])
                step = 1 if start <= end else -1
                result.extend(range(start, end + step, step))
        elif part.isdigit():
            result.append(int(part))
    return result

@st.fragment
def render_library_manager(prefix: str, title_name: str, file_path: str, css_str: str, target_type: str, text_model: str, api_delay: int, academy_name: str, school_name: str, h_url: str, f_url: str):
    saved_key = {"victory": "saved_v", "anatomy": "saved_a", "core": "saved_c", "voca": "saved_w", "past_voca": "saved_past_voca"}[target_type]
    sel_key = {"victory": "v_selected", "anatomy": "a_selected", "core": "c_selected", "voca": "w_selected", "past_voca": "past_v_selected"}[target_type]
    chk_prefix = {"victory": "vchk_", "anatomy": "achk_", "core": "cchk_", "voca": "wchk_", "past_voca": "past_vchk_"}[target_type]

    saved_list: List[Dict] = st.session_state[saved_key]
    selected_ids: List[str] = st.session_state[sel_key]

    if target_type in ["core", "voca"]:
        existing_labels = {p["label"] for p in saved_list}
        source_pool = {}
        for p in st.session_state.saved_v + st.session_state.saved_a + st.session_state.saved_c:
            lbl = p.get("label", "")
            if lbl and lbl not in existing_labels and p.get("body_text"):
                source_pool[lbl] = p

        core_by_label = {p["label"]: p["data"] for p in st.session_state.saved_c if p.get("label") and p.get("data")}

        if source_pool or (target_type == "voca" and core_by_label):
            with st.container(border=True):
                col_sync1, col_sync2 = st.columns([3.5, 2])
                if source_pool:
                    col_sync1.markdown(
                        f"💡 **엑셀 재업로드 불필요!** 기존 저장된 지문 중 **{title_name} 데이터가 아직 없는 지문이 `{len(source_pool)}개`** 있습니다."
                    )
                    if col_sync2.button(f"🔄 누락된 {len(source_pool)}개 지문 {title_name} 자동 생성", type="primary", use_container_width=True, key=f"sync_{prefix}_btn"):
                        p_bar = st.progress(0)
                        p_txt = st.empty()
                        pool_items = list(source_pool.values())
                        for idx, sp in enumerate(pool_items):
                            lbl = sp["label"]
                            p_txt.info(f"⏳ [{title_name}] 기출 DNA 기반 생성 중: {lbl} ({idx + 1}/{len(pool_items)})")
                            time.sleep(api_delay)
                            c_ctx = core_by_label.get(lbl) if target_type == "voca" else None
                            gen_data = analyze_passage(lbl, sp["body_text"], text_model, target_type, core_context=c_ctx)
                            if gen_data:
                                st.session_state[saved_key].append({
                                    "id": str(uuid.uuid4()),
                                    "label": lbl,
                                    "data": gen_data,
                                    "body_text": sp["body_text"],
                                    "academy_name": sp.get("academy_name", academy_name),
                                    "school_name": sp.get("school_name", school_name)
                                })
                                save_to_file(file_path, st.session_state[saved_key])
                                if target_type == "core":
                                    for vp in st.session_state.saved_w:
                                        if vp.get("label") == lbl:
                                            merge_core_vocab_into_voca_item(gen_data, vp.get("data", {}))
                                    save_to_file(VOCA_FILE, st.session_state.saved_w)
                            p_bar.progress((idx + 1) / len(pool_items))
                        p_txt.success(f"✅ {title_name} 자동 생성 및 어휘 연동 완료!")
                        st.rerun()

                if target_type == "voca" and core_by_label and saved_list:
                    if st.button("🔗 이미 생성된 [ANATOMY CORE]의 주제·요약 패러프레이징 어휘들을 현재 VOCA 단어장에 일괄 병합", use_container_width=True, key="merge_core_to_voca_btn"):
                        total_merged = 0
                        for vp in saved_list:
                            lbl = vp.get("label", "")
                            if lbl in core_by_label:
                                total_merged += merge_core_vocab_into_voca_item(core_by_label[lbl], vp.get("data", {}))
                        save_to_file(VOCA_FILE, saved_list)
                        st.success(f"✅ 총 {total_merged}개의 ANATOMY CORE 패러프레이징 핵심 어휘가 단어장에 추가·병합되었습니다!")

    if not saved_list:
        st.info(f"저장된 {title_name} 데이터가 없습니다. 상단에서 파일을 업로드해 생성해주세요.")
        return

    id_to_item = {p["id"]: p for p in saved_list}
    selected_ids[:] = [sid for sid in selected_ids if sid in id_to_item]
    selected_set = set(selected_ids)

    if target_type == "past_voca":
        q_c1, q_c2, q_c3 = st.columns(3)
        if q_c1.button("🏫 운양고 기출 시험지 전체 일괄 선택", use_container_width=True, key="past_sel_unyang"):
            selected_ids.clear()
            selected_set.clear()
            for p in saved_list:
                is_unyang = "운양" in p.get("label", "") or "운양" in p.get("school_name", "")
                st.session_state[f"{chk_prefix}{p['id']}"] = is_unyang
                if is_unyang:
                    selected_ids.append(p["id"])
                    selected_set.add(p["id"])
        if q_c2.button("🏫 제일고 기출 시험지 전체 일괄 선택", use_container_width=True, key="past_sel_jeil"):
            selected_ids.clear()
            selected_set.clear()
            for p in saved_list:
                is_jeil = "제일" in p.get("label", "") or "제일" in p.get("school_name", "")
                st.session_state[f"{chk_prefix}{p['id']}"] = is_jeil
                if is_jeil:
                    selected_ids.append(p["id"])
                    selected_set.add(p["id"])
        if q_c3.button("📚 저장된 모든 학교 기출 시험지 전체 선택", use_container_width=True, key="past_sel_all_schools"):
            selected_ids[:] = [p["id"] for p in saved_list]
            selected_set = set(selected_ids)
            for p in saved_list:
                st.session_state[f"{chk_prefix}{p['id']}"] = True

    groups = sorted(list({extract_group_prefix(p["label"]) for p in saved_list}), key=natural_sort_key)
    f_col1, f_col2, f_col3, f_col4 = st.columns([2, 2.5, 2.5, 1.5])
    
    with f_col1:
        selected_group = st.selectbox("📁 학교/출처 그룹 필터", options=["전체 보기"] + groups, key=f"{prefix}_group_filter")
    with f_col2:
        search_kw = st.text_input(f"🔍 {title_name} 키워드 검색 (띄어쓰기로 다중 검색)", placeholder="예: 운양고 26년 1학기", key=f"{prefix}_search")
    with f_col3:
        num_query = st.text_input("⚡ 번호 퀵 선택", placeholder="예: 1, 2 (또는 문항 번호)", key=f"{prefix}_num_query")
    with f_col4:
        page_limit_opt = st.selectbox("📄 화면 표시 개수", options=[60, 120, 240, "전체"], index=0, key=f"{prefix}_page_limit")

    keywords = [k.lower() for k in search_kw.split() if k.strip()]
    display_items = []
    for p in saved_list:
        if selected_group != "전체 보기" and extract_group_prefix(p["label"]) != selected_group:
            continue
        lbl_lower = p["label"].lower()
        if keywords and not all(k in lbl_lower for k in keywords):
            continue
        display_items.append(p)

    b_col1, b_col2, b_col3, b_col4, b_col5 = st.columns([1.5, 1.5, 1.5, 1.5, 1.5])

    if b_col1.button("✅ 검색 결과 전체 선택", use_container_width=True, key=f"{prefix}_sel_all"):
        for p in display_items:
            pid = p["id"]
            if pid not in selected_set:
                selected_ids.append(pid)
                selected_set.add(pid)
            st.session_state[f"{chk_prefix}{pid}"] = True

    if b_col2.button("⚡ 입력한 번호만 선택 추가", use_container_width=True, key=f"{prefix}_apply_nums"):
        target_nums = parse_number_query(num_query)
        if target_nums:
            for t_num in target_nums:
                for p in display_items:
                    if extract_question_number(p["label"]) == t_num:
                        pid = p["id"]
                        if pid not in selected_set:
                            selected_ids.append(pid)
                            selected_set.add(pid)
                        st.session_state[f"{chk_prefix}{pid}"] = True
        else:
            st.toast("💡 먼저 '번호 퀵 선택' 칸에 숫자를 입력해주세요.")

    if b_col3.button("🔢 선택 항목 이름/번호순 정렬", use_container_width=True, key=f"{prefix}_sort_sel"):
        selected_ids.sort(key=lambda sid: natural_sort_key(id_to_item[sid]))

    if b_col4.button("⬜ 전체 선택 해제", use_container_width=True, key=f"{prefix}_desel_all"):
        for sid in list(selected_ids):
            st.session_state[f"{chk_prefix}{sid}"] = False
        selected_ids.clear()
        selected_set.clear()

    if b_col5.button("🗑️ 선택 항목 영구 삭제", type="primary", use_container_width=True, key=f"{prefix}_del"):
        if selected_ids:
            st.session_state[saved_key] = [p for p in saved_list if p["id"] not in selected_set]
            for sid in list(selected_ids):
                if f"{chk_prefix}{sid}" in st.session_state:
                    del st.session_state[f"{chk_prefix}{sid}"]
            selected_ids.clear()
            save_to_file(file_path, st.session_state[saved_key])
            st.rerun()

    total_filtered = len(display_items)
    if page_limit_opt == "전체" or total_filtered <= int(page_limit_opt):
        paged_items = display_items
        cur_page_info = f"총 {len(saved_list)}개 중 {total_filtered}개 표시"
    else:
        limit = int(page_limit_opt)
        total_pages = max(1, (total_filtered + limit - 1) // limit)
        pg_col1, pg_col2 = st.columns([4, 1])
        with pg_col2:
            cur_page = st.number_input("페이지 이동", min_value=1, max_value=total_pages, value=1, step=1, key=f"{prefix}_pg")
        start_idx = (cur_page - 1) * limit
        paged_items = display_items[start_idx:start_idx + limit]
        cur_page_info = f"총 {len(saved_list)}개 중 {total_filtered}개 검색됨 ({cur_page}/{total_pages} 페이지 표시 중)"

    st.caption(f"📊 {cur_page_info} | 🎯 **현재 선택된 시험지/지문: {len(selected_ids)}개**")

    def on_card_toggle(pid: str):
        chk_val = st.session_state.get(f"{chk_prefix}{pid}", False)
        if chk_val and pid not in st.session_state[sel_key]:
            st.session_state[sel_key].append(pid)
        elif not chk_val and pid in st.session_state[sel_key]:
            st.session_state[sel_key].remove(pid)

    order_map = {sid: idx + 1 for idx, sid in enumerate(selected_ids)}

    with st.container(height=400, border=True):
        cols = st.columns(6)
        for i, p in enumerate(paged_items):
            pid = p["id"]
            is_sel = pid in order_map
            lbl = p["label"]
            order_idx = order_map.get(pid, 0)
            
            chk_key = f"{chk_prefix}{pid}"
            if chk_key not in st.session_state or st.session_state[chk_key] != is_sel:
                st.session_state[chk_key] = is_sel

            with cols[i % 6]:
                badge = (
                    f'<div style="position:absolute; top:-8px; right:-8px; background:#FF4B4B; color:white; '
                    f'border-radius:50%; width:22px; height:22px; font-size:11px; display:flex; '
                    f'align-items:center; justify-content:center; font-weight:bold; box-shadow:0 2px 4px rgba(0,0,0,0.2);">{order_idx}</div>'
                    if is_sel else ""
                )
                sub_info = ""
                if target_type in ["voca", "past_voca"]:
                    v_list = p.get("data", {}).get("voca_list", [])
                    sub_info = f'<div style="font-size:9.5px; color:#64748b; margin-top:3px;">(총 {len(v_list)}단어 수록)</div>'
                elif target_type == "core":
                    s_cnt = len(p.get("data", {}).get("summaries", []))
                    sub_info = f'<div style="font-size:9.5px; color:#0284c7; margin-top:3px;">(요약 {s_cnt}선·주제·T/F)</div>'

                st.markdown(
                    f'<div style="position:relative; background:{"#e8f0fe" if is_sel else "white"}; '
                    f'border:2px solid {"#0047AB" if is_sel else "#ddd"}; border-radius:10px; '
                    f'padding:10px 8px; text-align:center; font-weight:bold; color:#0047AB; '
                    f'min-height:45px; font-size:11px; display:flex; flex-direction:column; align-items:center; justify-content:center;">'
                    f'{badge}<div>{esc(lbl)}</div>{sub_info}</div>',
                    unsafe_allow_html=True
                )
                st.checkbox("선택", key=chk_key, on_change=on_card_toggle, args=(pid,))

    if not selected_ids:
        st.info("👆 위 목록에서 출력하거나 다운로드할 시험지/지문을 선택해주세요.")
    else:
        sel_passages = [id_to_item[sid] for sid in selected_ids if sid in id_to_item]
        preview_limit = 15
        order_preview = " ➔ ".join([f"**[{i+1}]** {md_safe(p['label'])}" for i, p in enumerate(sel_passages[:preview_limit])])
        if len(sel_passages) > preview_limit:
            order_preview += f" ... 외 {len(sel_passages) - preview_limit}개"
        st.markdown(f"🛒 **출력 순서 (총 {len(sel_passages)}개):** {order_preview}")

        render_permanent_json_editor(prefix, target_type, file_path, sel_passages, saved_list)

        if target_type in ["voca", "past_voca"]:
            r_col1, r_col2, r_col3, r_col4, r_col5 = st.columns([1.8, 1.8, 2.1, 2.1, 2.2])
        else:
            r_col1, r_col2 = st.columns([2.4, 5.6])
            r_col3, r_col4, r_col5 = None, None, None

        render_clicked = r_col1.button(
            f"👁️ 선택한 {title_name} 단행본 렌더링 ({len(sel_passages)}개)",
            type="primary",
            use_container_width=True,
            key=f"{prefix}_render_btn"
        )

        if target_type in ["voca", "past_voca"] and r_col2 is not None and r_col3 is not None and r_col4 is not None:
            excel_bytes = generate_voca_excel_bytes(sel_passages)
            r_col2.download_button(
                label="📊 선택 단어장 엑셀(.xlsx)",
                data=excel_bytes,
                file_name=f"{title_name.replace(' ', '_')}_통합단어장.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key=f"{prefix}_excel_dl_btn"
            )
            csv_main_bytes = generate_voca_main_csv_bytes(sel_passages)
            r_col3.download_button(
                label="📄 표제어(핵심+CORE) CSV",
                data=csv_main_bytes,
                file_name=f"{title_name.replace(' ', '_')}_표제어_통합.csv",
                mime="text/csv",
                use_container_width=True,
                key=f"{prefix}_csv_main_dl_btn"
            )
            csv_rel_bytes = generate_voca_rel_csv_bytes(sel_passages)
            r_col4.download_button(
                label="📄 유의어·반의어 전용 CSV",
                data=csv_rel_bytes,
                file_name=f"{title_name.replace(' ', '_')}_유의어_반의어.csv",
                mime="text/csv",
                use_container_width=True,
                key=f"{prefix}_csv_rel_dl_btn"
            )

        if render_clicked:
            with st.spinner(f"{title_name} 단행본(표지+공란+본문+저작권) 교재를 렌더링 중입니다..."):
                front_html, back_html = build_booklet_front_and_back(
                    target_type, h_url, f_url, academy_name, school_name, len(sel_passages)
                )
                if target_type == "victory":
                    body_str = "".join([render_victory(p) for p in sel_passages])
                    final_html = wrap_full_html_victory(body_str, css_str, front_html, back_html)
                    file_name = "Anatomy_One-Shot_Booklet.html"
                elif target_type == "anatomy":
                    body_str = "".join([render_anatomy(p) for p in sel_passages])
                    final_html = wrap_full_html_anatomy(body_str, css_str, front_html, back_html)
                    file_name = "Anatomy_Self-Note_Booklet.html"
                elif target_type == "core":
                    body_str = "".join([render_core(p) for p in sel_passages])
                    final_html = wrap_full_html_core(body_str, css_str, front_html, back_html)
                    file_name = "Anatomy_Core_Booklet.html"
                elif target_type == "past_voca":
                    # 시험지 단위로 연속 출력되고 다음 시험지로 바뀔 때만 페이지 넘김
                    body_str = render_past_voca_grouped_by_exam(sel_passages)
                    final_html = wrap_full_html_voca(body_str, css_str, front_html, back_html)
                    file_name = f"{title_name.replace(' ', '_')}_Booklet.html"
                else:
                    body_str = "".join([render_voca(p) for p in sel_passages])
                    final_html = wrap_full_html_voca(body_str, css_str, front_html, back_html)
                    file_name = f"{title_name.replace(' ', '_')}_Booklet.html"

                dl_target_col = r_col5 if (target_type in ["voca", "past_voca"] and r_col5 is not None) else r_col2
                with dl_target_col:
                    st.download_button(
                        f"📥 단행본 HTML 다운로드 ({title_name})",
                        final_html,
                        file_name,
                        "text/html",
                        use_container_width=True,
                        key=f"{prefix}_dl_btn"
                    )
                components.html(final_html, height=850, scrolling=True)


# -----------------------------
# 5-A. 올인원 통합 생성 엔진 (엑셀 탭 & 모의고사 탭 공용)
# -----------------------------
def run_all_in_one_generation(rows: List[Dict], img_dict: Dict, do_v: bool, do_a: bool, do_c: bool, do_w: bool, text_model: str, api_delay: int, academy_name: str, school_name: str):
    """rows: [{"label", "text", "school", "img_n"}] 형식의 지문 목록으로 4종 교재를 생성 (이미 생성된 라벨은 건너뜀)"""
    v_existing_labels = {p["label"] for p in st.session_state.saved_v}
    a_existing_labels = {p["label"] for p in st.session_state.saved_a}
    c_existing_map = {p["label"]: p["data"] for p in st.session_state.saved_c}
    w_existing_labels = {p["label"] for p in st.session_state.saved_w}

    progress_text = st.empty()
    progress_bar = st.progress(0)

    for idx, r in enumerate(rows):
        lbl = r['label']
        need_v = do_v and (lbl not in v_existing_labels)
        need_a = do_a and (lbl not in a_existing_labels)
        need_c = do_c and (lbl not in c_existing_map)
        need_w = do_w and (lbl not in w_existing_labels)

        if not (need_v or need_a or need_c or need_w):
            st.toast(f"⏭️ 이미 모두 생성됨 (건너뜀): {lbl}")
            progress_bar.progress((idx + 1) / max(1, len(rows)))
            continue

        img_data_url = img_dict.get(r['img_n'], "")
        excel_school = r['school']
        final_school = excel_school if pd.notna(excel_school) and str(excel_school).strip() not in ["", "nan"] else school_name

        # 모의고사 탭처럼 이미 저장된 이미지 경로가 있으면 그대로 사용
        shared_img_path = r.get("img_path") or (save_base64_to_image_file(img_data_url) if img_data_url else "")

        if need_v:
            progress_text.info(f"⏳ [Anatomy One-Shot] 분석 중: {lbl} ({idx + 1}/{len(rows)})")
            time.sleep(api_delay)
            v_data = analyze_passage(lbl, r['text'], text_model, "victory")
            if v_data:
                st.session_state.saved_v.append({
                    "id": str(uuid.uuid4()), "label": lbl, "data": v_data,
                    "body_text": r['text'], "image_data_url": "",
                    "image_path": shared_img_path, "academy_name": academy_name
                })
                v_existing_labels.add(lbl)
                save_to_file(VICTORY_FILE, st.session_state.saved_v)

        if need_a:
            progress_text.info(f"⏳ [Anatomy Self-Note] 분석 중: {lbl} ({idx + 1}/{len(rows)})")
            time.sleep(api_delay)
            a_data = analyze_passage(lbl, r['text'], text_model, "anatomy")
            if a_data:
                st.session_state.saved_a.append({
                    "id": str(uuid.uuid4()), "label": lbl, "data": a_data,
                    "body_text": r['text'], "image_data_url": "",
                    "image_path": shared_img_path, "academy_name": academy_name, "school_name": final_school
                })
                a_existing_labels.add(lbl)
                save_to_file(ANATOMY_FILE, st.session_state.saved_a)

        current_core_data = c_existing_map.get(lbl)
        if need_c:
            progress_text.info(f"⏳ [Anatomy Core] 다중 주제·요지·요약문 및 패러프레이징 생성 중: {lbl} ({idx + 1}/{len(rows)})")
            time.sleep(api_delay)
            c_data = analyze_passage(lbl, r['text'], text_model, "core")
            if c_data:
                current_core_data = c_data
                st.session_state.saved_c.append({
                    "id": str(uuid.uuid4()), "label": lbl, "data": c_data,
                    "body_text": r['text'], "academy_name": academy_name, "school_name": final_school
                })
                c_existing_map[lbl] = c_data
                save_to_file(CORE_FILE, st.session_state.saved_c)

        if need_w:
            progress_text.info(f"⏳ [Anatomy Voca] 지문 핵심어 + CORE 패러프레이징 통합 단어장 생성 중: {lbl} ({idx + 1}/{len(rows)})")
            time.sleep(api_delay)
            w_data = analyze_passage(lbl, r['text'], text_model, "voca", core_context=current_core_data)
            if w_data:
                if current_core_data:
                    merge_core_vocab_into_voca_item(current_core_data, w_data)
                st.session_state.saved_w.append({
                    "id": str(uuid.uuid4()), "label": lbl, "data": w_data,
                    "body_text": r['text'], "academy_name": academy_name, "school_name": final_school
                })
                w_existing_labels.add(lbl)
                save_to_file(VOCA_FILE, st.session_state.saved_w)

        progress_bar.progress((idx + 1) / max(1, len(rows)))

    progress_text.success("✅ 모든 교재(Anatomy One-Shot, Self-Note, Core, Voca) 통합 생성 완료!")
    progress_bar.empty()


# -----------------------------
# 5-B. 모의고사 PDF -> 정답 반영 원문 복원 & 년도/월/번호별 원문 라이브러리
# -----------------------------
SCHEMA_MOCK_RESTORE = {
    "type": "object",
    "properties": {
        "exam_year": {"type": "string", "description": "시행 연도 4자리 (예: 2026)"},
        "exam_month": {"type": "string", "description": "시행 월 숫자 (예: 3, 6, 9, 11)"},
        "grade": {"type": "string", "description": "학년 (예: 고1, 고2, 고3)"},
        "answer_key_found": {"type": "boolean", "description": "업로드 파일에서 공식 정답표를 찾았는지 여부"},
        "passages": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "number": {"type": "string", "description": "문항 번호 (예: 18, 31, 41~42, 43~45)"},
                    "q_type": {"type": "string", "description": "문제 유형 (예: 글의 목적, 빈칸 추론, 글의 순서, 문장 삽입)"},
                    "answer": {"type": "string", "description": "정답 번호 (예: ③, 41~42는 '41:② 42:④')"},
                    "restored_text": {"type": "string", "description": "정답이 반영된 원래의 영어 원문"},
                    "needs_review": {"type": "boolean", "description": "정답표가 없거나 복원이 불확실해 사람 확인이 필요하면 true"},
                    "note": {"type": "string", "description": "복원 시 적용한 내용 요약 (예: 빈칸에 ③ 표현 삽입, (B)-(A)-(C) 순서로 배열)"}
                },
                "required": ["number", "q_type", "answer", "restored_text", "needs_review", "note"]
            }
        }
    },
    "required": ["exam_year", "exam_month", "grade", "answer_key_found", "passages"]
}

MOCK_RESTORE_PROMPT = """너는 대한민국 고등학교 영어 모의고사(학력평가·수능 모의평가) 원문 복원 전문가다.
업로드된 모의고사 시험지 PDF{answer_note}에서 **독해 문항(18번~45번)** 의 지문을 꺼내, 문제로 변형되기 전의 **정답이 모두 반영된 원래 영어 원문**으로 복원해 JSON으로 반환하라. 듣기 문항(1~17번)은 제외한다.

[정답 출처]
- 별도 정답표 파일이 있으면 그 정답을 최우선으로 사용하라.
- 정답표 파일이 없으면 시험지 PDF 안(보통 마지막 페이지)의 정답표를 찾아 사용하라.
- 어디에도 정답표가 없으면 직접 풀어서 복원하되, 해당 문항의 needs_review를 true로 하고 answer_key_found는 false로 하라.

[유형별 복원 규칙]
- 빈칸 추론(31~34번 등): 빈칸 자리에 정답 선지의 표현을 문법적으로 자연스럽게 넣는다.
- 무관한 문장(35번): 정답에 해당하는 흐름과 무관한 문장을 삭제한다.
- 글의 순서(36~37번): 주어진 글 뒤에 (A)(B)(C)를 정답 순서대로 이어 붙인다.
- 문장 삽입(38~39번): 주어진 문장을 정답 위치(①~⑤)에 넣는다.
- 요약문(40번): 본문만 원문으로 남기고 요약문 문장은 넣지 않는다. (A)(B) 정답은 note에 적는다.
- 어법(29번): 틀린 밑줄 표현을 올바른 형태로 고친다. 나머지 밑줄 표현은 그대로 둔다.
- 어휘(30번, 42번 등): 문맥상 부적절한 낱말을 문맥에 맞는 낱말로 바꾼다.
- 함축 의미(21번) 등 밑줄 문항: 밑줄 표시만 없애고 문장은 그대로 둔다.
- 장문 41~42번: 한 지문이므로 number를 "41~42"로 하나만 만든다.
- 장문 43~45번: (A) 다음에 (B)(C)(D)를 정답 순서대로 배열하고, number를 "43~45"로 하나만 만든다.
- 도표·안내문·내용 일치 문항(25~28번 등)도 지문 텍스트를 그대로 복원한다.

[공통 규칙]
- 발문(문제 지시문), 선택지, 문항 번호, 원문자(①~⑤), (A)(B)(C)/(a)~(e) 표시, 밑줄, 하단 어휘 각주(*표시)는 모두 제거한다.
- 원문 문장을 절대 요약·의역·생략하지 말고 철자와 구두점을 시험지 그대로 옮긴다.
- 문단 구분은 빈 줄로 표시한다.
- note에는 어떤 정답을 어떻게 반영했는지 한국어로 짧게 적는다.
- exam_year, exam_month, grade는 시험지 표지/머리말에서 파악한다."""

def normalize_mock_meta(year: str, month: str, grade: str) -> (str, str, str):
    y = "".join(re.findall(r'\d+', str(year)))
    if len(y) == 2:
        y = "20" + y
    m = "".join(re.findall(r'\d+', str(month))).lstrip("0")
    g_digits = re.findall(r'\d', str(grade))
    g = f"고{g_digits[0]}" if g_digits else str(grade).strip()
    return y, m, g

def mock_exam_key(year: str, month: str, grade: str) -> str:
    return f"[{year}년 {month}월 {grade}]"

def extract_mock_exam_passages(pdf_file, answer_file, model_name: str) -> Dict:
    contents = [types.Part.from_bytes(data=pdf_file.getvalue(), mime_type="application/pdf")]
    answer_note = ""
    if answer_file is not None:
        ans_mime = answer_file.type if getattr(answer_file, "type", None) else "application/pdf"
        contents.append(types.Part.from_bytes(data=answer_file.getvalue(), mime_type=ans_mime))
        answer_note = f"(파일명: {pdf_file.name})과 별도 정답표 파일(파일명: {answer_file.name})"
    else:
        answer_note = f"(파일명: {pdf_file.name})"
    contents.append(MOCK_RESTORE_PROMPT.format(answer_note=answer_note))

    for attempt in range(4):
        try:
            res = client.models.generate_content(
                model=model_name,
                contents=contents,
                config=types.GenerateContentConfig(
                    temperature=0.1,
                    response_mime_type="application/json",
                    response_json_schema=SCHEMA_MOCK_RESTORE
                )
            )
            return json.loads(res.text.strip())
        except Exception as e:
            err_s = str(e).lower()
            if attempt < 3 and any(k in err_s for k in ["503", "429", "overloaded", "unavailable", "quota", "high demand"]):
                wait_t = 10 + attempt * 10
                st.toast(f"⏳ 서버 대기 감지 ({pdf_file.name}). {wait_t}초 후 재시도...")
                time.sleep(wait_t)
                continue
            st.error(f"❌ '{pdf_file.name}' 원문 복원 중 오류: {e}")
            return None
    return None

def mock_number_sort_key(num: str):
    nums = re.findall(r'\d+', str(num))
    return int(nums[0]) if nums else 999

def build_mock_excel_bytes(records: List[Dict]) -> bytes:
    """기존 '엑셀 원클릭 통합 생성' 탭에 그대로 올릴 수 있는 형식(라벨/지문/학교명/이미지명)"""
    rows = [{"라벨": p["label"], "지문": p.get("body_text", ""), "학교명": "", "이미지명": ""} for p in records]
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(rows, columns=["라벨", "지문", "학교명", "이미지명"]).to_excel(writer, sheet_name="모의고사_원문", index=False)
        try:
            ws = writer.book["모의고사_원문"]
            ws.column_dimensions["A"].width = 28
            ws.column_dimensions["B"].width = 120
        except Exception:
            pass
    return output.getvalue()

# ---- 5-C. 지문 이미지 생성 (기존 교재 이미지 스타일 참고) ----
# 3.1 flash: 장당 약 20초, 교재용으로 충분한 품질에 저렴 / pro: 더 정교하지만 비싸고 느림
IMAGE_MODELS = ["gemini-3.1-flash-image", "gemini-3-pro-image"]
IMAGE_PARALLEL = 4
BOOKLET_IMG_MAX_SIDE = 1280

MOCK_IMAGE_PROMPT = """Create one illustration for a Korean high school English textbook page, based on the English passage below.

[Style - match the attached reference images exactly]
- Cinematic, photorealistic scene, like a still from a movie, 16:9 landscape.
- Warm, moody practical lighting (lamps, window light), rich environmental detail and props that tell the story.
- Real people and objects drawn from the passage's situation or a concrete metaphor of its main idea.
- Include 2-4 short English text elements inside the scene (a poster, sign, note, board, book cover) that state the passage's key idea or keywords. Spell every word correctly. No Korean text.
- Do NOT copy the people, places or text of the reference images; use them only for visual style.
- No watermark, no logo, no border, no caption outside the scene.

[Passage]
\"\"\"{passage}\"\"\""""

def mock_image_src(p: Dict) -> str:
    """st.image에 넘길 이미지 위치 (로컬 파일 우선, 없으면 Supabase 공개 URL)"""
    path = p.get("image_path", "")
    if not path:
        return ""
    fn = os.path.basename(str(path).replace("\\", "/"))
    local_p = os.path.join(IMAGE_DIR, fn)
    if os.path.exists(local_p):
        return local_p
    if IS_CLOUD_DB:
        return f"{SUPABASE_URL}/storage/v1/object/public/image_assets/{fn}"
    return ""

def shrink_image_bytes(raw: bytes, max_side: int = 1024) -> bytes:
    """스타일 참고용 이미지를 작게 줄여 요청 용량을 줄임 (기존 교재 이미지는 장당 7~9MB)"""
    from PIL import Image
    im = Image.open(io.BytesIO(raw)).convert("RGB")
    im.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=85)
    return buf.getvalue()

@st.cache_data(show_spinner=False, max_entries=10)
def load_style_reference_bytes(img_filename: str) -> bytes:
    local_p = os.path.join(IMAGE_DIR, img_filename)
    try:
        if os.path.exists(local_p):
            with open(local_p, "rb") as f:
                raw = f.read()
        elif IS_CLOUD_DB:
            import urllib.request
            raw = urllib.request.urlopen(f"{SUPABASE_URL}/storage/v1/object/public/image_assets/{img_filename}", timeout=60).read()
        else:
            return b""
        return shrink_image_bytes(raw)
    except Exception:
        return b""

def default_style_reference_files(limit: int = 2) -> List[str]:
    """한판승 교재에 쓰인 이미지 중 앞에서부터 참고 이미지로 사용"""
    files = []
    for p in st.session_state.saved_v:
        if p.get("image_path"):
            fn = os.path.basename(str(p["image_path"]).replace("\\", "/"))
            if fn not in files:
                files.append(fn)
        if len(files) >= limit:
            break
    return files

def generate_passage_image(passage: str, style_refs: List[bytes], model_name: str) -> (bytes, str):
    """(이미지 bytes, 오류 메시지) 반환. 동시 생성 스레드에서 호출되므로 st.* 를 쓰지 않는다."""
    contents = [types.Part.from_bytes(data=b, mime_type="image/jpeg") for b in style_refs if b]
    contents.append(MOCK_IMAGE_PROMPT.format(passage=passage))
    last_err = ""
    for attempt in range(3):
        try:
            res = client.models.generate_content(
                model=model_name,
                contents=contents,
                config=types.GenerateContentConfig(
                    response_modalities=["IMAGE"],
                    image_config=types.ImageConfig(aspect_ratio="16:9")
                )
            )
            for cand in (res.candidates or []):
                for part in (cand.content.parts if cand.content else []) or []:
                    if part.inline_data and part.inline_data.data:
                        return part.inline_data.data, ""
            return b"", "모델이 이미지를 반환하지 않았습니다."
        except Exception as e:
            last_err = str(e)
            if attempt < 2 and any(k in last_err.lower() for k in ["503", "429", "overloaded", "unavailable", "quota", "high demand"]):
                time.sleep(10 + attempt * 10)
                continue
            return b"", last_err
    return b"", last_err

def save_booklet_image(raw: bytes, name_prefix: str) -> str:
    """교재용으로 1280px JPEG로 줄여 로컬 + Supabase Storage에 저장 (장당 약 200~300KB)"""
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(raw)).convert("RGB")
        im.thumbnail((BOOKLET_IMG_MAX_SIDE, BOOKLET_IMG_MAX_SIDE))
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=82, optimize=True)
        jpg = buf.getvalue()
    except Exception:
        return ""
    # 다시 만들 때 캐시에 옛 이미지가 남지 않도록 매번 새 파일명 사용
    filename = f"{name_prefix}_{uuid.uuid4().hex[:6]}.jpg"
    img_path = os.path.join(IMAGE_DIR, filename)
    with open(img_path, "wb") as f:
        f.write(jpg)
    if IS_CLOUD_DB:
        try:
            supabase.storage.from_("image_assets").upload(
                path=filename, file=jpg,
                file_options={"content-type": "image/jpeg", "upsert": "true"}
            )
        except Exception as e:
            st.warning(f"⚠️ 이미지 클라우드 업로드 실패 ({filename}): {e}")
    return img_path

def set_mock_image(p: Dict, img_path: str):
    """모의고사 지문에 이미지를 연결하고, 같은 라벨의 한판승·Self-Note 교재에도 반영
    (교재 이미지가 비어 있거나 이 지문의 이전 이미지였던 경우에만 교체)"""
    old_fn = os.path.basename(str(p.get("image_path", "")).replace("\\", "/"))
    p["image_path"] = img_path
    for saved_key, fpath in [("saved_v", VICTORY_FILE), ("saved_a", ANATOMY_FILE)]:
        hit = False
        for bp in st.session_state[saved_key]:
            if bp.get("label") != p["label"]:
                continue
            bp_fn = os.path.basename(str(bp.get("image_path", "")).replace("\\", "/"))
            if not bp_fn or (old_fn and bp_fn == old_fn):
                bp["image_path"] = img_path
                hit = True
        if hit:
            save_to_file(fpath, st.session_state[saved_key])

def render_mock_image_section(selected: List[Dict]):
    with st.container(border=True):
        st.markdown("##### 🎨 선택 지문 이미지 생성 (기존 교재 이미지 스타일 그대로)")
        st.caption("💡 기존 한판승 교재 이미지 2장을 스타일 참고로 자동 사용합니다. 다른 스타일을 원하면 참고 이미지를 직접 올려주세요. 만든 이미지는 이 지문으로 생성하는 한판승·Self-Note 교재에 자동으로 들어갑니다.")
        i_col1, i_col2 = st.columns([2, 1])
        ref_uploads = i_col1.file_uploader("🖼️ 스타일 참고 이미지 (선택, 여러 장 가능)", type=["png", "jpg", "jpeg"], accept_multiple_files=True, key="mock_style_refs")
        img_model = i_col2.selectbox("🎨 이미지 모델", IMAGE_MODELS, key="mock_img_model", help="gemini-3.1-flash-image: 빠르고 저렴(교재용 권장) / gemini-3-pro-image: 더 정교하지만 비싸고 느림")
        redo = i_col2.checkbox("이미 있는 이미지도 다시 만들기", value=False, key="mock_img_redo")

        if ref_uploads:
            style_refs = [shrink_image_bytes(u.getvalue()) for u in ref_uploads]
        else:
            style_refs = [load_style_reference_bytes(fn) for fn in default_style_reference_files()]
        style_refs = [b for b in style_refs if b]
        if style_refs:
            st.image(style_refs, width=160, caption=[f"참고 {i+1}" for i in range(len(style_refs))])
        else:
            st.warning("⚠️ 스타일 참고 이미지를 불러오지 못했습니다. 참고 이미지를 직접 올려주세요.")

        targets = [p for p in selected if redo or not p.get("image_path")]
        if st.button(f"🎨 이미지 생성 ({len(targets)}개, {IMAGE_PARALLEL}장씩 동시 생성)", type="primary", use_container_width=True, disabled=not targets or not style_refs, key="mock_img_gen_btn"):
            from concurrent.futures import ThreadPoolExecutor, as_completed
            prog = st.progress(0)
            status = st.empty()
            status.info(f"⏳ 이미지 {len(targets)}개 생성 중... ({IMAGE_PARALLEL}장씩 동시에 진행, 장당 약 20초)")
            done, errors = 0, []
            with ThreadPoolExecutor(max_workers=IMAGE_PARALLEL) as pool:
                futures = {pool.submit(generate_passage_image, p.get("body_text", ""), style_refs, img_model): p for p in targets}
                for n, fut in enumerate(as_completed(futures), start=1):
                    p = futures[fut]
                    img_bytes, err = fut.result()
                    img_path = save_booklet_image(img_bytes, f"mock_{p['id'][:8]}") if img_bytes else ""
                    if img_path:
                        set_mock_image(p, img_path)
                        done += 1
                    else:
                        errors.append(f"{p['label']}: {err or '이미지 저장 실패'}")
                    prog.progress(n / len(targets))
                    status.info(f"⏳ {n}/{len(targets)} 완료: {p['label']}")
            save_to_file(MOCK_FILE, st.session_state.saved_mock)
            prog.empty()
            status.success(f"✅ 이미지 {done}개를 만들었습니다.")
            for e in errors:
                st.error(f"❌ {e}")

        st.markdown(f"**🖼️ 선택 지문 이미지 ({sum(1 for p in selected if mock_image_src(p))}/{len(selected)}개)** · 마음에 안 드는 이미지는 '📤 내 이미지로 교체'로 직접 올릴 수 있습니다.")
        g_cols = st.columns(3)
        for i, p in enumerate(selected):
            with g_cols[i % 3]:
                src = mock_image_src(p)
                if src:
                    st.image(src, caption=p["label"], width="stretch")
                else:
                    st.caption(f"🖼️ {p['label']} · 이미지 없음")
                with st.popover("📤 내 이미지로 교체", use_container_width=True):
                    up = st.file_uploader("이미지 선택", type=["png", "jpg", "jpeg", "webp"], key=f"mock_own_img_{p['id']}")
                    if up is not None and st.button("이 이미지로 저장", type="primary", key=f"mock_own_save_{p['id']}"):
                        img_path = save_booklet_image(up.getvalue(), f"mock_{p['id'][:8]}")
                        if img_path:
                            set_mock_image(p, img_path)
                            save_to_file(MOCK_FILE, st.session_state.saved_mock)
                            st.rerun()
                        else:
                            st.error("❌ 이미지를 읽지 못했습니다.")

def render_mock_exam_tab(text_model: str, api_delay: int, academy_name: str, school_name: str):
    st.markdown("### 📝 모의고사 원문 라이브러리 (PDF 업로드 → 정답 반영 원문 복원 → 번호 클릭으로 교재 생성)")
    st.info(
        "💡 모의고사 시험지 PDF와 정답표를 올리면 AI가 **18~45번 지문을 정답이 반영된 원래 원문**으로 복원합니다. "
        "(빈칸 채움, 순서 재배열, 삽입 문장 반영, 무관한 문장 삭제, 어법·어휘 수정) "
        "복원 결과를 표에서 확인·수정한 뒤 저장하면, 아래에서 **년도·월·학년별로 번호를 클릭해 원하는 지문만 교재로 생성**할 수 있습니다."
    )

    # ---- 1) 업로드 & 복원 ----
    with st.container(border=True):
        st.markdown("##### 1️⃣ 모의고사 PDF 업로드 & 원문 복원")
        u_col1, u_col2 = st.columns(2)
        mock_pdf = u_col1.file_uploader("📄 모의고사 시험지 PDF", type=["pdf"], key="mock_pdf_uploader")
        mock_ans = u_col2.file_uploader(
            "✅ 정답표 (PDF 또는 이미지, 시험지 PDF 안에 정답표가 있으면 생략 가능)",
            type=["pdf", "png", "jpg", "jpeg"], key="mock_ans_uploader"
        )
        m_col1, m_col2, m_col3, m_col4 = st.columns([1, 1, 1, 2])
        in_year = m_col1.text_input("연도 (비우면 자동 인식)", placeholder="예: 2026", key="mock_in_year")
        in_month = m_col2.text_input("월 (비우면 자동 인식)", placeholder="예: 3", key="mock_in_month")
        in_grade = m_col3.selectbox("학년", ["자동 인식", "고1", "고2", "고3"], key="mock_in_grade")

        if mock_pdf and m_col4.button("🧠 정답 반영 원문 복원 시작", type="primary", use_container_width=True, key="mock_extract_btn"):
            with st.spinner(f"'{mock_pdf.name}'의 18~45번 지문을 정답 반영 원문으로 복원 중입니다... (1~2분 소요)"):
                parsed = extract_mock_exam_passages(mock_pdf, mock_ans, text_model)
            if parsed:
                y, m, g = normalize_mock_meta(
                    in_year.strip() or parsed.get("exam_year", ""),
                    in_month.strip() or parsed.get("exam_month", ""),
                    in_grade if in_grade != "자동 인식" else parsed.get("grade", "")
                )
                draft_rows = []
                for ps in sorted(parsed.get("passages", []), key=lambda x: mock_number_sort_key(x.get("number", ""))):
                    num = re.sub(r'\s*번\s*$', '', str(ps.get("number", "")).strip()).replace("-", "~")
                    if not num:
                        continue
                    draft_rows.append({
                        "번호": num,
                        "유형": ps.get("q_type", ""),
                        "정답": ps.get("answer", ""),
                        "확인필요": bool(ps.get("needs_review", False)) or not parsed.get("answer_key_found", False),
                        "복원 원문": ps.get("restored_text", "").strip(),
                        "복원 메모": ps.get("note", ""),
                    })
                for k in ["mock_draft_editor", "mock_draft_year", "mock_draft_month", "mock_draft_grade"]:
                    st.session_state.pop(k, None)
                st.session_state.mock_draft = {
                    "year": y, "month": m, "grade": g,
                    "answer_key_found": parsed.get("answer_key_found", False),
                    "source": mock_pdf.name,
                    "rows": draft_rows
                }
                st.rerun()

    # ---- 2) 검수 & 저장 ----
    draft = st.session_state.get("mock_draft")
    if draft:
        with st.container(border=True):
            st.markdown("##### 2️⃣ 복원 결과 검수 (셀을 더블클릭해 수정 → 저장)")
            d_col1, d_col2, d_col3 = st.columns(3)
            draft["year"] = d_col1.text_input("연도", value=draft["year"], key="mock_draft_year")
            draft["month"] = d_col2.text_input("월", value=draft["month"], key="mock_draft_month")
            draft["grade"] = d_col3.text_input("학년", value=draft["grade"], key="mock_draft_grade")
            y, m, g = normalize_mock_meta(draft["year"], draft["month"], draft["grade"])
            ekey = mock_exam_key(y, m, g)

            if not draft.get("answer_key_found"):
                st.warning("⚠️ 정답표를 찾지 못해 AI가 직접 푼 정답으로 복원했습니다. **확인필요** 표시된 문항을 꼭 검토해주세요.")
            review_cnt = sum(1 for r in draft["rows"] if r.get("확인필요"))
            st.caption(f"📌 저장될 이름: **{ekey}** · 지문 {len(draft['rows'])}개 · 확인필요 {review_cnt}개 · 원본 파일: {draft.get('source', '')}")

            exists = any(p.get("exam_key") == ekey for p in st.session_state.saved_mock)
            if exists:
                st.info(f"ℹ️ {ekey} 시험이 이미 라이브러리에 있습니다. 저장하면 같은 번호의 지문은 새 내용으로 교체됩니다.")

            edited = st.data_editor(
                pd.DataFrame(draft["rows"], columns=["번호", "유형", "정답", "확인필요", "복원 원문", "복원 메모"]),
                num_rows="dynamic",
                use_container_width=True,
                height=420,
                column_config={
                    "복원 원문": st.column_config.TextColumn(width="large"),
                    "확인필요": st.column_config.CheckboxColumn(),
                },
                key="mock_draft_editor"
            )

            with st.expander("🔍 지문 하나씩 크게 보기 (표에서 수정한 내용 반영)", expanded=False):
                nums = [str(n) for n in edited["번호"].tolist() if str(n).strip() and str(n) != "nan"]
                if nums:
                    v_num = st.selectbox("번호", nums, key="mock_draft_view_num")
                    v_row = edited[edited["번호"].astype(str) == v_num].iloc[0]
                    st.caption(f"유형: {v_row.get('유형', '')} · 정답: {v_row.get('정답', '')} · 메모: {v_row.get('복원 메모', '')}")
                    with st.container(border=True):
                        st.markdown(
                            f'<div style="white-space:pre-wrap; line-height:1.7; font-size:14px;">{esc(v_row.get("복원 원문", ""))}</div>',
                            unsafe_allow_html=True
                        )

            s_col1, s_col2 = st.columns([3, 1])
            if s_col1.button(f"💾 {ekey} 원문 라이브러리에 저장", type="primary", use_container_width=True, key="mock_draft_save"):
                if not (y and m and g):
                    st.error("연도·월·학년을 모두 입력해주세요.")
                else:
                    new_records = []
                    for _, r in edited.iterrows():
                        num = str(r.get("번호", "")).strip()
                        body = str(r.get("복원 원문", "")).strip()
                        if not num or num == "nan" or not body or body == "nan":
                            continue
                        new_records.append({
                            "id": str(uuid.uuid4()),
                            "label": f"{ekey} {num}번",
                            "exam_key": ekey,
                            "year": y, "month": m, "grade": g,
                            "number": num,
                            "q_type": str(r.get("유형", "") or ""),
                            "answer": str(r.get("정답", "") or ""),
                            "needs_review": bool(r.get("확인필요", False)),
                            "note": str(r.get("복원 메모", "") or ""),
                            "body_text": body,
                        })
                    new_labels = {p["label"] for p in new_records}
                    st.session_state.saved_mock = [p for p in st.session_state.saved_mock if p.get("label") not in new_labels] + new_records
                    save_to_file(MOCK_FILE, st.session_state.saved_mock)
                    del st.session_state["mock_draft"]
                    st.success(f"✅ {ekey} 지문 {len(new_records)}개를 저장했습니다!")
                    st.rerun()
            if s_col2.button("✖️ 검수 취소", use_container_width=True, key="mock_draft_cancel"):
                del st.session_state["mock_draft"]
                st.rerun()

    # ---- 3) 년도/월/학년 → 번호 클릭 선택 ----
    saved_mock: List[Dict] = st.session_state.saved_mock
    st.markdown("##### 3️⃣ 모의고사 선택 → 번호 클릭 → 교재 생성")
    if not saved_mock:
        st.info("저장된 모의고사 원문이 없습니다. 위에서 모의고사 PDF를 업로드해 복원해주세요.")
        return

    years = sorted({p.get("year", "") for p in saved_mock if p.get("year")}, reverse=True)
    f_col1, f_col2, f_col3 = st.columns(3)
    sel_year = f_col1.selectbox("📅 연도", ["전체"] + years, key="mock_f_year")
    months = sorted({p.get("month", "") for p in saved_mock if p.get("month") and (sel_year == "전체" or p.get("year") == sel_year)}, key=lambda x: int(x) if x.isdigit() else 99)
    sel_month = f_col2.selectbox("🗓️ 월", ["전체"] + months, key="mock_f_month")
    grades = sorted({p.get("grade", "") for p in saved_mock if p.get("grade")})
    sel_grade = f_col3.selectbox("🎓 학년", ["전체"] + grades, key="mock_f_grade")

    exams: Dict[str, List[Dict]] = {}
    for p in saved_mock:
        if sel_year != "전체" and p.get("year") != sel_year:
            continue
        if sel_month != "전체" and p.get("month") != sel_month:
            continue
        if sel_grade != "전체" and p.get("grade") != sel_grade:
            continue
        exams.setdefault(p.get("exam_key", extract_group_prefix(p["label"])), []).append(p)

    def exam_sort_key(k: str):
        items = exams[k]
        y = int(items[0].get("year", 0) or 0)
        m = int(items[0].get("month", 0) or 0)
        return (-y, -m, items[0].get("grade", ""))

    # 위젯 상태는 화면에 안 그려지면 사라지므로, 시험별 선택 번호를 별도로 보관 (필터를 바꿔도 선택 유지)
    pick_store: Dict[str, List[str]] = st.session_state.setdefault("mock_pick_store", {})

    with st.container(height=420, border=True):
        if not exams:
            st.caption("조건에 맞는 모의고사가 없습니다.")
        for ekey in sorted(exams.keys(), key=exam_sort_key):
            items = sorted(exams[ekey], key=lambda x: mock_number_sort_key(x.get("number", "")))
            num_to_item = {p["number"]: p for p in items}
            pill_key = f"mock_pick_{ekey}"
            if pill_key not in st.session_state:
                st.session_state[pill_key] = [n for n in pick_store.get(ekey, []) if n in num_to_item]
            h_col1, h_col2, h_col3 = st.columns([4, 1, 1])
            h_col1.markdown(f"**{ekey}** · 지문 {len(items)}개")
            if h_col2.button("전체 선택", key=f"mock_all_{ekey}", use_container_width=True):
                st.session_state[pill_key] = list(num_to_item.keys())
            if h_col3.button("선택 해제", key=f"mock_none_{ekey}", use_container_width=True):
                st.session_state[pill_key] = []
            picked = st.pills(
                ekey,
                options=list(num_to_item.keys()),
                selection_mode="multi",
                format_func=lambda n, m=num_to_item: f"{n}번 ⚠️" if m[n].get("needs_review") else f"{n}번",
                key=pill_key,
                label_visibility="collapsed"
            )
            pick_store[ekey] = list(picked or [])

    # 모든 시험의 선택 번호 모으기 (필터로 숨겨진 시험의 선택도 유지)
    by_label = {p["label"]: p for p in saved_mock}
    selected: List[Dict] = []
    for ekey, nums in pick_store.items():
        for num in nums:
            item = by_label.get(f"{ekey} {num}번")
            if item:
                selected.append(item)
    selected.sort(key=natural_sort_key)

    if not selected:
        st.caption("👆 번호를 클릭해 교재로 만들 지문을 선택하세요. (⚠️ = 정답 확인 필요)")
        return

    st.markdown(f"🛒 **선택된 지문 {len(selected)}개:** " + ", ".join(md_safe(p["label"]) for p in selected[:20]) + (f" ... 외 {len(selected) - 20}개" if len(selected) > 20 else ""))

    with st.expander("🔎 선택 지문 원문 확인·수정", expanded=False):
        pick_lbl = st.selectbox("지문 선택", [p["label"] for p in selected], key="mock_view_select")
        target = by_label[pick_lbl]
        st.caption(f"유형: {target.get('q_type', '')} · 정답: {target.get('answer', '')} · 메모: {target.get('note', '')}")
        new_body = st.text_area("원문", value=target.get("body_text", ""), height=260, key=f"mock_body_{target['id']}")
        new_review = st.checkbox("⚠️ 확인필요", value=bool(target.get("needs_review")), key=f"mock_rev_{target['id']}")
        if st.button("💾 이 지문 수정 저장", key=f"mock_save_one_{target['id']}"):
            target["body_text"] = new_body.strip()
            target["needs_review"] = new_review
            save_to_file(MOCK_FILE, st.session_state.saved_mock)
            st.success(f"✅ {pick_lbl} 원문을 저장했습니다.")

    render_mock_image_section(selected)

    a_col1, a_col2 = st.columns(2)
    a_col1.download_button(
        "📊 선택 지문 엑셀 다운로드 (엑셀 원클릭 생성 탭 형식)",
        data=build_mock_excel_bytes(selected),
        file_name="모의고사_원문.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="mock_excel_dl"
    )
    with a_col2.popover("🗑️ 선택 지문 삭제", use_container_width=True):
        st.warning(f"선택한 지문 {len(selected)}개를 원문 라이브러리에서 영구 삭제합니다. (이미 만든 교재는 지워지지 않습니다)")
        if st.button("삭제 확정", type="primary", key="mock_del_confirm"):
            del_ids = {p["id"] for p in selected}
            st.session_state.saved_mock = [p for p in saved_mock if p["id"] not in del_ids]
            for k in [k for k in st.session_state.keys() if str(k).startswith("mock_pick_")]:
                del st.session_state[k]
            st.session_state["mock_pick_store"] = {}
            save_to_file(MOCK_FILE, st.session_state.saved_mock)
            st.rerun()

    with st.container(border=True):
        st.markdown("##### 🎯 선택 지문으로 교재 생성")
        g_col1, g_col2, g_col3, g_col4 = st.columns(4)
        mdo_v = g_col1.checkbox("🏆 Anatomy One-Shot (한판승)", value=True, key="mock_do_v")
        mdo_a = g_col2.checkbox("🦅 Anatomy Self-Note", value=True, key="mock_do_a")
        mdo_c = g_col3.checkbox("💎 Anatomy Core", value=True, key="mock_do_c")
        mdo_w = g_col4.checkbox("📖 Anatomy Voca", value=True, key="mock_do_w")
        if any(p.get("needs_review") for p in selected):
            st.warning("⚠️ 선택한 지문 중 정답 확인이 필요한 지문이 있습니다. 위 '원문 확인·수정'에서 먼저 검토하는 것을 권장합니다.")
        if st.button(f"🚀 선택한 {len(selected)}개 지문 교재 생성 시작", type="primary", use_container_width=True, key="mock_gen_btn"):
            if not (mdo_v or mdo_a or mdo_c or mdo_w):
                st.warning("생성할 교재 종류를 최소 1개 이상 체크해주세요.")
            else:
                rows = [{"label": p["label"], "text": p["body_text"], "school": "", "img_n": "", "img_path": p.get("image_path", "")} for p in selected]
                run_all_in_one_generation(rows, {}, mdo_v, mdo_a, mdo_c, mdo_w, text_model, api_delay, academy_name, school_name)
                st.info("👉 생성된 교재는 각 교재 탭(한판승·Self-Note·Core·Voca)에서 같은 이름으로 확인·렌더링할 수 있습니다.")


# -----------------------------
# 6. 메인 UI 화면 및 사이드바 구성
# -----------------------------
st.title("🦅 4pm 올인원 교재 통합 시스템")

with st.sidebar:
    st.markdown("### 🖼️ 로고 및 정보 설정")
    st.caption("💾 한 번 올린 로고와 입력한 정보는 자동 저장되어, 바꾸기 전까지 계속 유지됩니다.")
    app_settings = st.session_state.app_settings
    settings_changed = False

    logo_urls = {}
    for slot, label in [("h", "머리말 로고"), ("f", "꼬리말 로고"), ("e", "영쌤 로고")]:
        up = st.file_uploader(label, type=["png", "jpg", "jpeg"], key=slot)
        if up is not None:
            # 같은 파일이 화면을 다시 그릴 때마다 저장되지 않도록 내용 해시로 비교
            up_hash = hashlib.md5(up.getvalue()).hexdigest()
            if app_settings.get(f"logo_{slot}_hash") != up_hash:
                app_settings[f"logo_{slot}"] = logo_upload_to_data_url(up)
                app_settings[f"logo_{slot}_hash"] = up_hash
                settings_changed = True
        logo_urls[slot] = app_settings.get(f"logo_{slot}") or DEFAULT_LOGO_URL
        if app_settings.get(f"logo_{slot}") and up is None:
            st.image(app_settings[f"logo_{slot}"], width=48, caption=f"저장된 {label}")

    if any(app_settings.get(f"logo_{s}") for s in "hfe"):
        if st.button("🔄 저장된 로고 모두 기본값으로", use_container_width=True, key="reset_logos"):
            for s in "hfe":
                app_settings.pop(f"logo_{s}", None)
                app_settings.pop(f"logo_{s}_hash", None)
            save_app_settings(app_settings)
            st.rerun()

    academy_name = st.text_input("🏫 학원명", app_settings.get("academy_name", "포피엠 입시 영어"))
    school_name = st.text_input("🏫 학교명 (표지 및 ANATOMY용)", app_settings.get("school_name", "영어는 영쌤!"))

    model_options = [
        "gemini-2.5-flash",
        "gemini-3.5-flash",
        "gemini-3.1-pro",
        "gemini-3.1-flash-lite",
        "gemini-1.5-flash"
    ]
    saved_model = app_settings.get("text_model", model_options[0])
    text_model = st.selectbox("📦 구글 AI 모델", model_options, index=model_options.index(saved_model) if saved_model in model_options else 0)
    api_delay = st.slider(
        "⏱️ 지문 간 API 대기 시간(초)",
        min_value=1, max_value=10, value=int(app_settings.get("api_delay", 3)),
        help="유료 Tier 1 이상이면 1~2초로 줄여 생성 속도를 높일 수 있습니다."
    )

    for k, v in [("academy_name", academy_name), ("school_name", school_name), ("text_model", text_model), ("api_delay", api_delay)]:
        if app_settings.get(k) != v:
            app_settings[k] = v
            settings_changed = True
    if settings_changed:
        save_app_settings(app_settings)

    st.markdown("---")
    st.markdown(f"### ☁️ 저장소 상태: {'**클라우드 DB 연결됨 🟢**' if IS_CLOUD_DB else '**로컬 PC 모드 🟡**'}")
    with st.expander("📦 기존 PC 데이터 클라우드 이전 / 백업"):
        st.caption("💡 이미지는 Supabase의 [Storage -> image_assets 버킷]에 직접 드래그해 올리거나, 구글 드라이브에서 받은 ZIP 파일 1개로 한 번에 올릴 수 있습니다.")
        up_jsons = st.file_uploader(
            "1) 기존 JSON 파일 업로드 (victory, anatomy, core, voca, past_voca)",
            type=["json"],
            accept_multiple_files=True,
            key="cloud_mig_json"
        )
        up_zip = st.file_uploader(
            "2) [선택] image_assets.zip 압축 파일 1개로 통째 업로드",
            type=["zip"],
            key="cloud_mig_zip"
        )
        if st.button("🚀 클라우드 DB로 데이터 이식 실행", type="primary", use_container_width=True):
            with st.spinner("데이터를 클라우드 DB에 동기화 중입니다..."):
                if up_zip is not None:
                    with zipfile.ZipFile(up_zip) as zf:
                        img_entries = [info for info in zf.infolist() if not info.is_dir() and info.filename.lower().endswith((".png", ".jpg", ".jpeg"))]
                        prog = st.progress(0)
                        for idx, info in enumerate(img_entries):
                            fname = os.path.basename(info.filename)
                            if not fname or fname.startswith("."):
                                continue
                            raw_bytes = zf.read(info.filename)
                            with open(os.path.join(IMAGE_DIR, fname), "wb") as f:
                                f.write(raw_bytes)
                            if IS_CLOUD_DB:
                                try:
                                    supabase.storage.from_("image_assets").upload(
                                        path=fname,
                                        file=raw_bytes,
                                        file_options={"content-type": "image/png", "upsert": "true"}
                                    )
                                except Exception:
                                    pass
                            prog.progress((idx + 1) / max(1, len(img_entries)))
                        prog.empty()

                for uj in (up_jsons or []):
                    j_data = json.load(uj)
                    if isinstance(j_data, list):
                        j_data, _ = normalize_passage_records(j_data)
                        fname = uj.name.lower()
                        if "victory" in fname:
                            st.session_state.saved_v = j_data
                            save_to_file(VICTORY_FILE, j_data)
                        elif "anatomy" in fname:
                            st.session_state.saved_a = j_data
                            save_to_file(ANATOMY_FILE, j_data)
                        elif "core" in fname:
                            st.session_state.saved_c = j_data
                            save_to_file(CORE_FILE, j_data)
                        elif "mock" in fname:
                            st.session_state.saved_mock = j_data
                            save_to_file(MOCK_FILE, j_data)
                        elif "past" in fname:
                            j_data, _ = consolidate_past_voca_records(j_data)
                            st.session_state.saved_past_voca = j_data
                            save_to_file(PAST_VOCA_FILE, j_data)
                        elif "voca" in fname:
                            st.session_state.saved_w = j_data
                            save_to_file(VOCA_FILE, j_data)
                st.cache_data.clear()
                st.success("✅ 클라우드 DB로 모든 데이터 이식이 완료되었습니다!")
                st.rerun()

h_url, f_url, e_url = logo_urls["h"], logo_urls["f"], logo_urls["e"]
CSS_VICTORY = build_css_victory(h_url, f_url, e_url)
CSS_ANATOMY = build_css_anatomy(h_url, f_url, e_url)
CSS_CORE = build_css_core(h_url, f_url, e_url)
CSS_VOCA = build_css_voca(h_url, f_url)

tab1, tab_mock, tab2, tab3, tab4, tab5, tab6, tab7 = st.tabs([
    "🚀 엑셀 원클릭 통합 생성",
    "📝 모의고사 원문 라이브러리",
    "🏆 Anatomy One-Shot (한판승)",
    "🦅 Anatomy Self-Note",
    "💎 Anatomy Core",
    "📖 Anatomy Voca",
    "🎯 기출 PDF 출제DNA 학습기",
    "📚 기출 어휘 독립 라이브러리"
])

with tab1:
    st.info("💡 엑셀과 이미지를 업로드하고 생성 버튼을 누르면, **Anatomy One-Shot(한판승)**, **Anatomy Self-Note**, **Anatomy Core**, **Anatomy Voca**가 동시에 생성되며, **Core에 쓰인 패러프레이징 어휘들이 Voca 단어장에 자동으로 포함**됩니다.")
    xlsx = st.file_uploader("📂 엑셀 업로드", type=["xlsx"])
    imgs = st.file_uploader("🖼️ 이미지 대량 업로드", type=["png", "jpg", "jpeg"], accept_multiple_files=True)
    
    st.markdown("##### 🎯 생성할 교재 종류 선택 (원하는 항목만 체크 가능)")
    g_col1, g_col2, g_col3, g_col4 = st.columns(4)
    do_v = g_col1.checkbox("🏆 Anatomy One-Shot (한판승)", value=True)
    do_a = g_col2.checkbox("🦅 Anatomy Self-Note", value=True)
    do_c = g_col3.checkbox("💎 Anatomy Core", value=True)
    do_w = g_col4.checkbox("📖 Anatomy Voca", value=True)

    if xlsx and st.button("🚀 올인원 통합 생성 시작", type="primary"):
        if not (do_v or do_a or do_c or do_w):
            st.warning("생성할 교재 종류를 최소 1개 이상 체크해주세요.")
        else:
            df = pd.read_excel(xlsx)
            img_dict = {
                i.name: f"data:{i.type if hasattr(i, 'type') and i.type else 'image/png'};base64,{base64.b64encode(i.getvalue()).decode()}"
                for i in imgs
            } if imgs else {}

            rows = []
            for i, r in df.iterrows():
                raw_text = r.get('passage', r.get('지문', ''))
                if pd.isna(raw_text) or not str(raw_text).strip():
                    continue
                rows.append({
                    "label": clean_label(r.get('label', r.get('라벨', f'지문_{i+1}'))),
                    "text": str(raw_text).strip(),
                    "school": str(r.get('school', r.get('학교명', ''))),
                    "img_n": str(r.get('image_name', r.get('이미지명', ''))).strip()
                })

            run_all_in_one_generation(rows, img_dict, do_v, do_a, do_c, do_w, text_model, api_delay, academy_name, school_name)

with tab_mock:
    render_mock_exam_tab(text_model, api_delay, academy_name, school_name)

with tab2:
    render_library_manager("v", "Anatomy One-Shot", VICTORY_FILE, CSS_VICTORY, "victory", text_model, api_delay, academy_name, school_name, h_url, f_url)

with tab3:
    render_library_manager("a", "Anatomy Self-Note", ANATOMY_FILE, CSS_ANATOMY, "anatomy", text_model, api_delay, academy_name, school_name, h_url, f_url)

with tab4:
    render_library_manager("c", "Anatomy Core", CORE_FILE, CSS_CORE, "core", text_model, api_delay, academy_name, school_name, h_url, f_url)

with tab5:
    render_library_manager("w", "Anatomy Voca", VOCA_FILE, CSS_VOCA, "voca", text_model, api_delay, academy_name, school_name, h_url, f_url)

with tab6:
    st.markdown("### 🎯 학교별 기출 시험지 PDF 출제DNA 분석 & 누적 학습기")
    st.info(
        "💡 **운양고(고1·2·3) 7개 시험지 및 제일고(고2) 1개 시험지의 출제 DNA가 기본 탑재**되어 있습니다. "
        "앞으로 새로운 시험지 PDF를 이곳에 업로드하면 AI가 해당 시험지의 **주제·요약문 패러프레이징 방식, 빈칸 타겟 어휘 수준, 어휘 추론 함정 패턴**을 분석해 클라우드 DB에 영구 누적시킵니다."
    )

    col_dna1, col_dna2 = st.columns([1.2, 1.8])
    with col_dna1:
        uploaded_pdfs = st.file_uploader(
            "📄 분석할 학교 기출 시험지 PDF 업로드 (여러 개 동시 가능)",
            type=["pdf"],
            accept_multiple_files=True,
            key="dna_pdf_uploader"
        )
        school_tag = st.text_input("🏫 시험지 출처 태그 (예: 26년 2학기 기말 제일고1)", placeholder="예: 26년 2학기 기말 운양고2")

        if uploaded_pdfs and st.button("🧠 업로드한 기출 PDF 출제DNA 추출 및 누적 적용", type="primary", use_container_width=True):
            with st.spinner("기출 시험지 PDF의 패러프레이징 패턴과 어휘 함정을 정밀 분석 중입니다..."):
                try:
                    contents_payload = [types.Part.from_bytes(data=up_pdf.getvalue(), mime_type="application/pdf") for up_pdf in uploaded_pdfs]
                    dna_prompt = f"""너는 대한민국 고등학교 영어 내신 기출 분석 전문가다.
업로드된 영어 시험지 PDF({', '.join([u.name for u in uploaded_pdfs])})를 정밀 분석하여 다음 4가지를 추출하라.
1. 주제/제목/요지/주장 선지에서 원문 표현을 어떤 수준의 상위 개념어(Meta-vocabulary)와 구문으로 패러프레이징하는지 (실제 시험지 속 구체적 영어 예시 포함)
2. 요약문 빈칸 (A), (B) 문항에서 어떤 품사와 성격의 단어를 빈칸 타겟으로 삼는지 (실제 시험지 속 영어 예시 포함)
3. '문맥상 낱말의 쓰임' 및 어휘 변형 문제에서 어떤 반의어/혼동어 함정을 사용하는지 (실제 영어 예시 포함)
4. <보기> 영문 일치/불일치 문항에서 오답과 정답 선지를 어떻게 변형하는지

기존 출제 DNA 규칙:
{st.session_state.exam_dna.get('rules_text', '')}

위 기존 규칙을 유지하면서 새로 발견된 출제 패턴과 구체적 어휘 예시를 통합하여, 앞으로 'ANATOMY CORE'와 'ANATOMY VOCA' 생성 프롬프트에 바로 넣을 수 있는 완성된 [출제 DNA 핵심 규칙] 텍스트(updated_rules_text)와 이번 시험지 핵심 요약(analysis_summary)을 JSON으로 반환하라."""

                    dna_schema = {
                        "type": "object",
                        "properties": {"updated_rules_text": {"type": "string"}, "analysis_summary": {"type": "string"}},
                        "required": ["updated_rules_text", "analysis_summary"]
                    }
                    contents_payload.append(dna_prompt)
                    res = client.models.generate_content(
                        model=text_model,
                        contents=contents_payload,
                        config=types.GenerateContentConfig(temperature=0.2, response_mime_type="application/json", response_json_schema=dna_schema)
                    )
                    parsed_dna = json.loads(res.text.strip())
                    st.session_state.exam_dna["rules_text"] = parsed_dna["updated_rules_text"]
                    st.session_state.exam_dna["updated_at"] = time.strftime("%Y-%m-%d %H:%M")
                    st.session_state.exam_dna.setdefault("history", []).append({
                        "title": school_tag.strip() if school_tag.strip() else ", ".join([u.name for u in uploaded_pdfs]),
                        "summary": parsed_dna["analysis_summary"]
                    })
                    save_exam_dna(st.session_state.exam_dna)
                    st.success("✅ 기출 PDF 출제DNA 분석 및 클라우드 DB 반영 완료!")
                    st.rerun()
                except Exception as e:
                    st.error(f"❌ PDF 분석 중 오류 발생: {e}")

        st.markdown("#### 📚 누적 학습된 기출 히스토리")
        for h_idx, h_item in enumerate(reversed(st.session_state.exam_dna.get("history", [])), start=1):
            with st.container(border=True):
                st.markdown(f"**{h_idx}. {h_item.get('title', '기출 분석')}**")
                st.caption(h_item.get("summary", ""))

    with col_dna2:
        st.markdown(f"#### 🧬 현재 활성화된 출제 DNA 프롬프트 규칙 (최근 업데이트: `{st.session_state.exam_dna.get('updated_at', '')}`)")
        st.caption("이 규칙이 `Anatomy Core`와 `Anatomy Voca`를 생성할 때 AI에게 자동으로 주입됩니다. 직접 텍스트를 다듬어 저장하실 수도 있습니다.")
        edited_rules = st.text_area(
            "출제 DNA 프롬프트 규칙 편집",
            value=st.session_state.exam_dna.get("rules_text", ""),
            height=420,
            label_visibility="collapsed"
        )
        dc1, dc2 = st.columns(2)
        if dc1.button("💾 출제 DNA 규칙 수동 저장", type="primary", use_container_width=True):
            st.session_state.exam_dna["rules_text"] = edited_rules
            save_exam_dna(st.session_state.exam_dna)
            st.success("✅ 출제 DNA 규칙이 저장되었습니다!")
        if dc2.button("🔄 운양고·제일고 기본값으로 초기화", use_container_width=True):
            st.session_state.exam_dna = dict(DEFAULT_EXAM_DNA)
            save_exam_dna(st.session_state.exam_dna)
            st.rerun()

# ==========================================
# 📚 탭 7: 기출 어휘 독립 라이브러리 (시험지 단위 통합 & 시험지별 페이지 구분)
# ==========================================
with tab7:
    st.markdown("### 📚 학교별 기출 시험지 어휘 독립 라이브러리 (시험지 단위 연속 출력 & 시험지 간 페이지 분리)")
    st.info(
        "💡 **한 시험지(예: `[운양고 26년 1학기 중간 고1]`)의 모든 문항 단어들이 여백 없이 연달아서 이어서 출력**되며, "
        "**다음 시험지(예: `[운양고 26년 1학기 기말 고1]`)로 바뀔 때만 다음 장(새 페이지)부터 시작**합니다. "
        "각 단어 아래에는 몇 년도 몇 학년 몇 번 문제(본문/선지)에서 나왔는지 상세 출처가 자동 표기됩니다."
    )

    with st.container(border=True):
        up_past_pdfs = st.file_uploader(
            "📂 어휘를 추출할 기출 시험지 PDF 여러 개 동시 업로드 (수십 개 또는 분할 PDF 동시 선택 가능)",
            type=["pdf"],
            accept_multiple_files=True,
            key="past_pdf_multi_uploader"
        )
        p_col1, p_col2 = st.columns([2, 4])
        override_school = p_col1.text_input("🏫 학교명 강제 지정 (비워두면 파일명/시험지에서 자동 인식)", placeholder="예: 운양고 또는 제일고", key="past_override_school")

        if up_past_pdfs and p_col2.button(f"🚀 업로드된 {len(up_past_pdfs)}개 기출 시험지 어휘 일괄 추출 및 시험지별 통합 저장", type="primary", use_container_width=True):
            prog_bar = st.progress(0)
            status_box = st.empty()
            total_added_words = 0

            for idx, pdf_file in enumerate(up_past_pdfs):
                status_box.info(f"⏳ ({idx + 1}/{len(up_past_pdfs)}) '{pdf_file.name}' 시험지 본문 및 선지 어휘 정밀 추출 중...")
                if idx > 0:
                    time.sleep(api_delay)

                prompt_past = f"""너는 대한민국 고등학교 영어 내신 기출 시험지 분석 전문가다.
업로드된 기출 시험지 PDF(파일명: {pdf_file.name})의 모든 문항(1번~끝번, 서술형/논술형 포함)을 전수 조사하여 다음을 JSON으로 반환하라.

1. 출처 메타데이터 자동 파악:
   - school_name: 파일명이나 시험지 상단/하단에서 학교명을 파악해 '운양고', '제일고' 등으로 기재 (파일명에 '[운양1]'이 있으면 '운양고', '[제일2]'가 있으면 '제일고').
   - exam_year_term: 연도 및 학기/시험 종류 (예: '26년 1학기 중간', '26년 1학기 기말', '26년 2학기 중간').
   - grade: 학년 (예: '고1', '고2', '고3').
2. 문항별(questions) 중고등 핵심 어휘 및 숙어 추출:
   - question_number: 문항 번호 (예: '1번', '2번', ..., '24번', '논술형1').
   - words: 해당 문항의 **지문 본문**과 **객관식 선지(1~5번 및 <보기>)**에 등장한 중고등 필수 어휘·숙어·패러프레이징 어휘를 빠짐없이 추출하라.
   - 초등 기초 단어(be동사, a, the, beautiful, happy, good 등)는 절대 제외하라.
   - 각 단어의 source_part에는 본문에서 나왔으면 '기출본문', 문제 선지나 <보기>/요약문에서 나왔으면 '기출선지'라고 정확히 구분해 적어라.
   - pronunciation은 반드시 자연스러운 '한글 발음'(IPA 금지)으로 표기하고, 문맥에 맞는 유의어(synonyms 1~2개)와 반의어(antonyms 1~2개)도 각각 한글발음과 뜻을 포함해 작성하라."""

                for attempt in range(4):
                    try:
                        res = client.models.generate_content(
                            model=text_model,
                            contents=[
                                types.Part.from_bytes(data=pdf_file.getvalue(), mime_type="application/pdf"),
                                prompt_past
                            ],
                            config=types.GenerateContentConfig(
                                temperature=0.2,
                                response_mime_type="application/json",
                                response_json_schema=SCHEMA_PAST_PDF_VOCA
                            )
                        )
                        parsed_pdf = json.loads(res.text.strip())
                        sch_n = override_school.strip() if override_school.strip() else parsed_pdf.get("school_name", "기출").strip()
                        yr_term = parsed_pdf.get("exam_year_term", "").strip()
                        grd = parsed_pdf.get("grade", "").strip()

                        # 시험지 단위로 통합 라벨 생성 (예: [운양고 26년 1학기 중간 고1])
                        exam_title_clean = re.sub(r'\s+', ' ', f"{sch_n} {yr_term} {grd}").strip()
                        exam_full_label = clean_label(f"[{exam_title_clean}]")

                        exam_words = []
                        for q_obj in parsed_pdf.get("questions", []):
                            q_num = str(q_obj.get("question_number", "")).strip()
                            if not q_num:
                                continue
                            for w in q_obj.get("words", []):
                                src_part = str(w.get("source_part", "기출본문")).strip()
                                exam_words.append({
                                    "word": str(w.get("word", "")).strip(),
                                    "pos": str(w.get("pos", "")).strip(),
                                    "pronunciation": str(w.get("pronunciation", "")).strip(),
                                    "meaning": str(w.get("meaning", "")).strip(),
                                    "category": f"{exam_title_clean} {q_num} [{src_part}]",
                                    "synonyms": w.get("synonyms", []),
                                    "antonyms": w.get("antonyms", [])
                                })

                        if exam_words:
                            st.session_state.saved_past_voca.append({
                                "id": str(uuid.uuid4()),
                                "label": exam_full_label,
                                "data": {"voca_list": exam_words},
                                "body_text": f"{exam_full_label} 기출 본문 및 선지 통합 어휘",
                                "academy_name": academy_name,
                                "school_name": sch_n
                            })
                            total_added_words += len(exam_words)
                            # 동일 시험지(예: 전반부/후반부 분할 PDF)가 있으면 자동으로 하나의 시험지로 합본 및 문항순 정렬
                            st.session_state.saved_past_voca, _ = consolidate_past_voca_records(st.session_state.saved_past_voca)
                            save_to_file(PAST_VOCA_FILE, st.session_state.saved_past_voca)
                        break
                    except Exception as e:
                        err_s = str(e).lower()
                        if attempt < 3 and any(k in err_s for k in ["503", "429", "overloaded", "unavailable", "quota"]):
                            wait_t = 10 + attempt * 10
                            st.toast(f"⏳ 서버 대기 감지 ({pdf_file.name}). {wait_t}초 후 재시도...")
                            time.sleep(wait_t)
                            continue
                        st.error(f"❌ '{pdf_file.name}' 처리 중 오류: {e}")
                        break

                prog_bar.progress((idx + 1) / len(up_past_pdfs))

            status_box.success(f"✅ 총 {len(up_past_pdfs)}개 파일에서 {total_added_words}개의 기출 어휘가 시험지 단위로 통합 저장되었습니다!")
            prog_bar.empty()
            st.rerun()

    st.markdown("---")
    render_library_manager("past_v", "기출 어휘 단어장", PAST_VOCA_FILE, CSS_VOCA, "past_voca", text_model, api_delay, academy_name, school_name, h_url, f_url)
